"""
Bank statement reader (v2.2): turns the statement file an owner downloads from net banking into a clean list of
transactions, then sorts them into the Expenses categories and finds the payments that repeat every month.

Reads (decided from the file's content, not its name):
    CSV / semicolon / tab-separated text, Excel .xlsx, Excel .xls (real BIFF via xlrd, plus the HTML tables, tab-separated
    text and Excel-2003 XML that many Indian banks save with an ".xls" name), HTML tables, and PDF (tables first, then
    text lines; password-protected PDFs open with the password the owner types).
    Header rows, preamble rows (account details), footers and totals, repeated page headers, wrapped narrations,
    Indian number formats (1,23,456.78 / 1,234.00 Dr / (1,234.00) / -1234) and the usual date formats are handled.

Never: no network, no AI, no database or file storage, no logging. The PDF password is only handed to the PDF library to
open the file — it is not stored, returned or put in any message. Narrations are masked (long numbers, e-mails, UPI
handles, PAN) before they leave this module; the unmasked text is not kept.

Limits: 25 MB per file, 60 PDF pages, 60 MB unzipped for .xlsx (zip-bomb guard), 5,000 transactions per statement.
Every cell is cut to 400 characters (amounts and dates to 40) before any pattern matching, so a hostile file can't make
the patterns slow. Scanned (image-only) PDFs and photos can't be read: the owner is asked for the Excel/CSV download.

Use parse_isolated() from the web server: it runs parse() in a separate, memory- and CPU-limited process (at most 2 at
a time), so a decompression bomb or a pathological file can't freeze or exhaust the server.
"""
import csv
import hashlib
import io
import json
import os
import re
import statistics
import struct
import subprocess
import sys
import threading
import zipfile
from bisect import bisect_left, bisect_right
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from html.parser import HTMLParser

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_PAGES = 60
MAX_UNZIPPED = 60 * 1024 * 1024
MAX_TXNS = 5000
MAX_ROWS = 200_000          # rows read per sheet/table before giving up on the rest
MAX_COLS = 60
CELL_CAP = 400              # characters kept from any cell or narration before pattern matching
NUM_CAP = 40                # an amount or date longer than this is not an amount or date
LINE_CAP = 500              # characters kept from a PDF/text line
NARR_CAP = 600              # a narration joined from wrapped rows stops growing here
MAX_HTML_DEPTH = 6          # tables inside tables inside tables...
MAX_HTML_TABLES = 2000
MAX_HTML_CELLS = 300_000
PAGE_CONTENT_CAP = 1024 * 1024            # decoded drawing commands per PDF page (a real statement page: ~10-100 KB)
PDF_CONTENT_CAP = 24 * 1024 * 1024        # ... and per PDF
PAGE_BREAK = "\f"

DEBIT_CATEGORIES = ("rent", "salaries", "pf_esi", "electricity", "phone_internet", "software", "travel", "fuel", "freight",
                    "marketing", "professional", "bank_charges", "loan_emi", "insurance", "repairs", "office", "food_staff",
                    "taxes", "purchases", "equipment", "other", "cash_withdrawal", "transfer")
NOT_EXPENSES = ("cash_withdrawal", "transfer")
CREDIT_CATEGORIES = ("sales_receipt", "refund_in", "interest_in", "loan_in", "cash_deposit", "transfer_in", "other_in")

MSG_NO_TXNS = ("We couldn't find the transactions in this file. Download the statement as Excel or CSV from net banking "
               "and try again.")
MSG_SCAN = ("This PDF is a scanned image, so there's no text in it we can read. Download the statement as Excel or CSV "
            "from net banking and upload that instead.")
MSG_TOO_COMPLEX = ("This file is too big or too complex to read here. Download the statement as Excel or CSV and try "
                   "again.")
ERROR_CODES = ("password_needed", "wrong_password", "unsupported", "unreadable", "no_transactions", "too_big")


class StatementError(Exception):
    """A problem the owner can act on. `message` is plain English and is shown to them as is."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


# ----------------------------------------------------------------------------------------------------------------------
# Dates and amounts

_MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10,
           "nov": 11, "dec": 12}
_MON = (r"(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sept?(?:ember)?|oct(?:ober)?|"
        r"nov(?:ember)?|dec(?:ember)?)")
_D_NUM = re.compile(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4}|\d{2})(?!\d)")
_D_ISO = re.compile(r"(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?!\d)")
_D_MON = re.compile(r"(\d{1,2})(?:st|nd|rd|th)?[\s/\-.,]*" + _MON + r"\.?[\s/\-.,]*(\d{4}|\d{2})(?!\d)", re.I)
_D_MON_US = re.compile(_MON + r"\.?\s+(\d{1,2}),?\s+(\d{4})(?!\d)", re.I)
_TIME_REST = re.compile(r"(?:[ T]*\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?\s*(?:[ap]\.?m\.?)?(?:\s*\(?ist\)?)?)?\s*$", re.I)
_EXCEL_EPOCH = date(1899, 12, 30)


def _year(y: str) -> int:
    n = int(y)
    if len(y) == 2:
        n += 2000 if n < 70 else 1900
    return n


def _mk_date(y: int, m: int, d: int) -> date | None:
    if not 1990 <= y <= 2100:
        return None
    try:
        return date(y, m, d)
    except ValueError:
        return None


def _match_date(s: str, monthfirst: bool = False) -> tuple[date, int] | None:
    """A date at the very start of `s` -> (date, length of the matched text)."""
    m = _D_ISO.match(s)
    if m:
        dt = _mk_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        return (dt, m.end()) if dt else None
    m = _D_NUM.match(s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), _year(m.group(3))
        day, mon = (b, a) if monthfirst else (a, b)
        dt = _mk_date(y, mon, day) or (_mk_date(y, day, mon) if day <= 12 < mon else None)
        return (dt, m.end()) if dt else None
    m = _D_MON.match(s)
    if m:
        dt = _mk_date(_year(m.group(3)), _MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
        return (dt, m.end()) if dt else None
    m = _D_MON_US.match(s)
    if m:
        dt = _mk_date(int(m.group(3)), _MONTHS[m.group(1)[:3].lower()], int(m.group(2)))
        return (dt, m.end()) if dt else None
    return None


def _to_date(v, monthfirst: bool = False) -> date | None:
    """A whole cell as a date: date/datetime objects, Excel serial numbers, or text with an optional time after it."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, datetime):
        return v.date() if 1990 <= v.year <= 2100 else None
    if isinstance(v, date):
        return v if 1990 <= v.year <= 2100 else None
    if isinstance(v, (int, float)):
        if 32874 <= v <= 73050:            # Excel serials for 1990-01-01 .. 2099-12-31
            return _EXCEL_EPOCH + timedelta(days=int(v))
        return None
    s = str(v).strip()
    if len(s) > NUM_CAP:
        return None
    hit = _match_date(s, monthfirst)
    if hit:
        rest = s[hit[1]:].strip()
        if _TIME_REST.fullmatch(rest):
            return hit[0]
        again = _match_date(rest, monthfirst)       # "01/04/2026 01/04/2026": transaction and value date in one cell
        if again and _TIME_REST.fullmatch(rest[again[1]:]):
            return hit[0]
    return None


_AMT_PLAIN = re.compile(r"\d{1,3}(?:,\d{2,3})+(?:\.\d{1,4})?|\d+(?:\.\d{1,4})?|\.\d{1,4}")
_AMT_MARK = re.compile(r"(?<!\s)\s*\(?(?<![A-Za-z])(dr|cr|db|debit|credit)\.?\)?\s*$", re.I)


def _amount(v) -> tuple[float | None, str | None]:
    """A cell as (signed amount, 'dr' | 'cr' | None). Blank, '-', '0.00' and anything not a number give (None, marker)."""
    if v is None or isinstance(v, bool):
        return None, None
    if isinstance(v, (int, float)):
        if v != v or v == 0 or abs(v) >= 1e10:
            return None, None
        return float(v), None
    s = str(v).replace("\xa0", " ").strip()
    if len(s) > NUM_CAP:
        return None, None
    marker = None
    m = _AMT_MARK.search(s)
    if m and m.start() > 0:
        marker = "dr" if m.group(1).lower() in ("dr", "db", "debit") else "cr"
        s = s[:m.start()]
    s = re.sub(r"(?i)₹|\brs\.?|\binr\b|\s", "", s)
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s.startswith("-"):
        neg, s = True, s[1:]
    elif s.endswith("-"):
        neg, s = True, s[:-1]
    elif s.startswith("+"):
        s = s[1:]
    if not _AMT_PLAIN.fullmatch(s):
        return None, marker
    val = float(s.replace(",", ""))
    if val == 0 or val >= 1e10:
        return None, marker
    return (-val if neg else val), marker


def _signed_balance(v, marker_cell=None) -> float | None:
    val, mk = _amount(v)
    if val is None:
        return None
    mk = mk or _drcr(marker_cell)
    return -abs(val) if mk == "dr" else val


def _drcr(v) -> str | None:
    s = re.sub(r"[^a-z]", "", str(v or "").lower())
    if s in ("dr", "d", "debit", "db", "withdrawal", "wdl"):
        return "dr"
    if s in ("cr", "c", "credit", "deposit", "dep"):
        return "cr"
    return None


def _inr(x: float) -> str:
    """₹ with Indian digit grouping, e.g. ₹1,23,456.50."""
    neg, x = x < 0, abs(round(x, 2))
    whole, frac = f"{x:.2f}".split(".")
    head, tail = whole[:-3], whole[-3:]
    while len(head) > 2:
        tail = head[-2:] + "," + tail
        head = head[:-2]
    s = (head + "," + tail) if head else tail
    return ("-" if neg else "") + "₹" + s + ("" if frac == "00" else "." + frac)


def _plural(n: int, one: str, many: str) -> str:
    return f"{n} {one if n == 1 else many}"


# ----------------------------------------------------------------------------------------------------------------------
# Masking and payee

# Every pattern below starts only where a run starts (lookbehind), so a long run is scanned once, not once per character.
_EMAIL = re.compile(r"(?<![A-Za-z0-9._%+])([A-Za-z0-9._%+]+)@([A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]+)*)")
_HANDLE = re.compile(r"(?<![A-Za-z0-9._%+*])[A-Za-z0-9._%+*]*@[A-Za-z0-9.]*")   # e-mail / UPI handle; stops at / - space
_PAN = re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b")
_GROUPED_DIGITS = re.compile(r"(?<!\d)\d+(?:[ \-]\d+)+")   # 4111-1111-1111-1234, 98200 12345, 98200-12345
_LONG_DIGITS = re.compile(r"\d{6,}")


def _keep_last4(s: str) -> str:
    digits = sum(ch.isdigit() for ch in s)
    out, seen = [], 0
    for ch in s:
        if ch.isdigit():
            seen += 1
            out.append(ch if seen > digits - 4 else "X")
        else:
            out.append(ch)
    return "".join(out)


def mask(text: str) -> str:
    """Hide what shouldn't be stored: 6+ digit runs keep their last 4 (account, card, phone, reference numbers), and so
    do digit groups split by spaces or hyphens that add up to 10+ digits (4111-1111-1111-1234, 98200 12345); e-mails and
    UPI handles keep the first character before the @ ("r***@okaxis"), PAN keeps its last 4.
    Names stay — the owner needs to see who they paid."""
    s = str(text or "")
    s = _EMAIL.sub(lambda m: m.group(1)[0] + "***@" + m.group(2), s)
    s = _PAN.sub(lambda m: "XXXXXX" + m.group(0)[-4:], s)
    s = _GROUPED_DIGITS.sub(lambda m: _keep_last4(m.group(0)) if sum(c.isdigit() for c in m.group(0)) >= 10 else m.group(0), s)
    s = _LONG_DIGITS.sub(lambda m: "X" * (len(m.group(0)) - 4) + m.group(0)[-4:], s)
    return s


_WRAP_PREFIX = re.compile(r"^(?:BY|TO)\s+(?:TRANSFER|TRF|CLG|CLEARING|INST)\b\s*[-:/]?\s*")
_POS = re.compile(r"^(?:POS|PCD|VPS|IPS|ECOM(?: PUR)?)\b\s*[/\-:]?\s*(?:[0-9X*]{4,}(?:\s?[0-9X*]{4})*)?\s*[/\-:]?\s*(.+)$")
_CHANNEL = re.compile(r"^(?:UPI|NEFT|RTGS|IMPS|MMT|ACH|NACH|ECS|CHQ|CHEQUE|CLG|INB|IB|MB|BIL|BPAY|FT|INF|INFT|IFT|TRF|"
                      r"FUND TRANSFER|SI|CMS)\b")
_IFSC = re.compile(r"[A-Z]{4}[0X][A-Z0-9X]{6}")
_REFLIKE = re.compile(r"[A-Z]{0,6}[0-9X*]{6,}[A-Z0-9X]*")
_GENERIC = set("""UPI NEFT RTGS IMPS MMT ACH NACH ECS DR CR D C P2A P2M P2P IN OUT INB IB MB NETBANK NET BANKING MOBILE
PAYMENT PAYMENTS PAY PAID PMT FOR TO FROM BY THE OF AND BILL SENT USING VIA NO REMARKS REMARK NA NIL TRANSFER TRF FT FUND
FUNDS ONLINE COLLECT REQUEST PHONE PH CHQ CHEQUE DEP DEPOSIT CLG CLEARING MICR CTS INWARD OUTWARD INST OTHERS OTHER INTENT
UPIINTENT QR BIL ONL BPAY PURCHASE POS ECOM PUR INF INFT IFT OK GOODS TXN TRANSACTION REF UTR MANDATE AUTOPAY DEBIT CREDIT
SI CMS MUM MUMBAI DEL DELHI SENTUSING AT ATM NEFTINW RTGSINW ON YOUR ACCOUNT AC A CHG CHGS CHRG CHRGS CHARGE CHARGES FEE
FEES GST INCL""".split())
_BANKS = set("""HDFC ICIC ICICI SBIN SBI UTIB AXIS KKBK KOTAK YESB YES IDFB IDFC BARB BOB PUNB PNB CNRB CANARA INDB
INDUSIND PYTM PAYTM YBL IBL AXL APL FBL FDRL FEDERAL IDIB UBIN UNION CBIN BKID IOBA MAHB AIRP JIOP AUBL RATN SCBL CITI
HSBC DBSS BANK LTD LIMITED OF INDIA""".split())
_KEEP_UPPER = {"LLP", "HUF", "OPC", "LLC", "CA", "GST", "TDS", "LIC", "HP", "II", "III", "IV", "IT", "AC", "HDFC", "ICICI",
               "SBI", "IDFC", "PNB", "MSEDCL", "BESCOM", "TPDDL", "BSES", "TNEB", "CESC", "UPPCL", "KSEB", "BSNL", "MTNL",
               "IRCTC", "HPCL", "IOCL", "BPCL", "DTDC", "AWS", "EPFO", "ESIC", "CBIC", "GSTN", "CBDT", "UPI", "ATM"}


def _junk(tok: str) -> bool:
    tok = tok[:80]                           # names are cut to 60 anyway; keeps the checks below short
    if len(re.sub(r"[^A-Z]", "", tok)) < 2 or "@" in tok:
        return True
    if _IFSC.fullmatch(tok) or _REFLIKE.fullmatch(tok.replace(" ", "")):
        return True
    words = [w for w in re.findall(r"[A-Z0-9]+", tok) if not re.fullmatch(r"X*\d[X\d]*|X{3,}", w)]   # numbers, masked refs
    return all(w in _GENERIC or w in _BANKS for w in words)


def _title(s: str) -> str:
    out = re.sub(r"[A-Za-z]+", lambda m: m.group(0).capitalize(), s.lower())
    return " ".join(w.upper() if w.upper() in _KEEP_UPPER else w for w in out.split())


def _clean_name(tok: str) -> str:
    words = re.sub(r"\s+", " ", tok).strip(" .,_#:;'\"()").split(" ")
    while len(words) > 1 and re.fullmatch(r"[0-9X*]*\d[0-9X*]*", words[-1]):   # a reference number after the name
        words.pop()
    return _title(" ".join(words))[:60].strip()


def payee(narration: str) -> str:
    """Best-effort counterparty name from common Indian narration patterns (UPI, NEFT, IMPS, RTGS, POS, ACH/NACH,
    "BY/TO TRANSFER", cheques), title-cased, without refs/IFSCs/handles. Empty string when unsure."""
    s = re.sub(r"\s+", " ", str(narration or "")[:CELL_CAP]).strip().upper()
    if not s:
        return ""
    s = re.sub(r"\bM/S\.?\s*", "", s)
    wrapped = False
    m = _WRAP_PREFIX.match(s)
    if m:
        s, wrapped = s[m.end():], True
    m = _POS.match(s)
    if m:
        for tok in re.split(r"\s*[/\-*|]\s*|\s{2,}", m.group(1)):
            tok = tok.strip()
            if tok and not _junk(tok):
                return _clean_name(tok[:80])
        return ""
    if not (_CHANNEL.match(s) or wrapped or re.match(r"^CHQ (?:PAID|DEP)", s)):
        return ""
    for tok in re.split(r"\s*[/\-*:|]\s*|\s{2,}", s):
        tok = tok.strip()
        if tok and not _junk(tok):
            return _clean_name(tok[:80])
    return ""


# ----------------------------------------------------------------------------------------------------------------------
# Categories

def _rx(p: str) -> re.Pattern:
    return re.compile(p, re.I)


_STRUCTURED = _rx(r"\b(?:UPI|NEFT|IMPS|RTGS|MMT)\b[^/\-*]{0,8}[/\-*][^/\-*]*[/\-*]")
_CHANNEL_WORD = r"\b(?:UPI|NEFT|IMPS|RTGS|MMT)\b"

# (category, pattern, only when the narration names no counterparty: not a structured UPI/NEFT/IMPS/RTGS transfer and no
# payee found — so "SMS CHG" is a bank charge but "UPI/.../RAHUL TRANSPORT/FREIGHT CHARGES" isn't). First match wins.
# A pattern given as a tuple means "all of these appear" (kept as separate searches so none of them needs a slow ".*").
_DEBIT_RULES = [
    ("transfer", r"\bTRF TO OWN\b|\bOWN (?:A/?C|ACCOUNT|ACCT)\b|\bSELF TRANSFER\b|\bTRANSFER TO SELF\b|\b(?:AUTO ?)?SWEEP\b|"
                 r"\bTO FD\b|\bFD (?:BOOK|CREAT|OPEN)|\bFIXED DEPOSIT\b|\bTD BOOK|\bMUTUAL FUND\b|\bSIP\b|\bRD INSTAL", False),
    ("transfer", (_CHANNEL_WORD, r"\bSELF\b"), False),          # UPI/NEFT to the owner's own account elsewhere
    ("bank_charges", r"\bCHRGS?\b|\bCHGS?\b|\bCHARGES?\b|\bSMS ALERT|\bAMC\b|\bMIN(?:IMUM)?\.? BAL|\bNON[- ]?MAINT|\bGST ON\b|"
                     r"\bDEBIT CARD (?:FEE|ANNUAL|ISSUANCE)|\bDC (?:ANNUAL )?FEE\b|\bCARD FEE\b|\bANNUAL FEE\b|"
                     r"\bINTEREST DEBIT\b|\bINT(?:EREST)?\.? DR\b|\bINT\.? ON OD\b|\bOD INT|\bCHE?Q(?:UE)? BOOK\b|"
                     r"\bSERVICE CHARGE|\bPROCESSING FEE\b|\bSTAMP DUTY\b|\bFOLIO\b|\bINSUFFICIENT|\bRTN CHG|\bRETURN CHG", True),
    ("cash_withdrawal", r"\bATM\b|\bNWD\b|\bCASH ?WDL\b|\bCASH WITHDRAWAL\b|\bCASH W/?D\b|\bCASH PAID\b|\bSELF\b", False),
    ("pf_esi", r"\bEPFO\b|\bEPF\b|\bPF CHALLAN\b|\bPF (?:PAYMENT|CONTRIBUTION|DUES|DEPOSIT)\b|\bPROVIDENT FUND\b|\bESIC\b|"
               r"\bESI\b|\bEMPLOYEES? STATE INSURANCE\b|\bPROF(?:ESSIONAL)?\.? ?TAX\b|\bPTAX\b|\bP\.? TAX\b|"
               r"\bLABOU?R WELFARE\b|\bLWF\b", False),
    ("taxes", r"\bGSTN\b|\bCBIC\b|\bCPIN\b|\bGST ?(?:CHALLAN|PAYMENT|PMT|PAID|DEPOSIT)\b|\b[CSI]GST\b|\bTDS\b|"
              r"\bINCOME ?TAX\b|\bITNS ?\d*|\bCBDT\b|\bADVANCE TAX\b|\bSELF ASSESSMENT TAX\b|\bTAX PAYMENT\b|\bOLTAS\b|"
              r"\bE-?TAX\b|\bTIN-?NSDL\b|\bPROPERTY TAX\b|\bMUNICIPAL TAX\b|\bCHALLAN\b", False),
    ("taxes", r"\bGST\b", True),
    ("salaries", r"\bSALARY\b|\bSALARIES\b|\bSAL\b|\bPAYROLL\b|\bWAGES?\b|\bSTIPEND\b|\bBONUS\b|\bINCENTIVES?\b", False),
    ("loan_emi", r"\bEMI\b|\bLOAN\b|\bLN REPAY|\bHOUSING FIN", False),
    ("insurance", r"\bLIC\b|\bLIFE INS|\bINSURANCE\b|\bERGO\b|\bLOMBARD\b|\bALLIANZ\b|\bSTAR HEALTH\b|"
                  r"\bNIVA BUPA\b|\bCARE HEALTH\b|\bTATA AIG\b|\bASSURANCE\b|\bPOLICY\b", False),
    ("insurance", r"\bPREMIUM\b", True),     # "PREMIUM PLYWOOD" paid by UPI is a supplier, not a policy
    ("loan_emi", r"\bACH ?D\b|\bACH[- ]?DR\b", False),
    ("loan_emi", (r"\b(?:NACH|ECS|ACH|MANDATE)\b", r"\b(?:FIN|FINANCE|FINANCIAL|FINCORP|FINSERV|FINVEST|CAPITAL|CREDIT|"
                                                     r"LOANS?|LENDING|NBFC|HOUSING|LEASING|HFC|HFL|BANK)\b"), False),
    ("rent", r"\bRENT\b|\bRENTAL\b|\bRENTS\b|\bLEASE\b|\bLEAVE (?:AND|&) LICEN[CS]E\b", False),
    ("electricity", r"\bMSEDCL\b|\bMAHAVITARAN\b|\bMSEB\b|\bBESCOM\b|\bTPDDL\b|\bTATA POWER\b|\bBSES\b|\bTORRENT POWER\b|"
                    r"\bADANI ELEC\w*|\bTNEB\b|\bTANGEDCO\b|\bCESC\b|\bUPPCL\b|\bKSEB\b|\bPSPCL\b|\b[AJ]?[AJ]?VVNL\b|"
                    r"\b[PMD]VVNL\b|\b[DU]HBVN\w*|\bWBSEDCL\b|\bAP[SEC]PDCL\b|\bTS[SN]PDCL\b|\bCSPDCL\b|\bMP\w{1,3}VVCL\b|"
                    r"\b[DMPU]GVCL\b|\b(?:CHE|HE|GE|ME)SCOM\b|\b[NS]BPDCL\b|\bTP[CSWN]ODL\b|\bAPDCL\b|\bNPCL\b|\bKESCO\b|"
                    r"\bHPSEB\w*|\bJBVNL\b|\bBEST UNDERTAKING\b|\bELECTRICITY\b|\bELECTRIC(?:AL)? BILL\b|\bPOWER BILL\b|"
                    r"\bEB BILL\b|\bLIGHT BILL\b", False),
    ("fuel", r"\bHPCL\b|\bHINDUSTAN PETROLEUM\b|\bIOCL\b|\bINDIAN ?OIL\b|\bIOC\b|\bBPCL\b|\bBHARAT PETROLEUM\b|\bPETROL\b|"
             r"\bPETROLEUM\b|\bDIESEL\b|\bFUELS?\b|\bFILLING STATION\b|\bSERVICE STATION\b|\bFAS ?TAG\b|\bNETC\b|\bCNG\b|"
             r"\bNAYARA\b|\bSHELL\b|\bHP ?PAY\b|\bPETRO\b|\bTOLL\b", False),
    ("phone_internet", r"\bAIRTEL\b|\bJIO\b|\bVODAFONE\b|\bIDEA CELLULAR\b|\bVI (?:POSTPAID|PREPAID|BILL|RECHARGE|MOBILE)\b|"
                       r"[/\-]VI[/\-]|^VI\b|\bBSNL\b|\bMTNL\b|\bACT ?FIBERNET\b|\bATRIA CONVERGENCE\b|\bHATHWAY\b|"
                       r"\bBROADBAND\b|\bFIBERNET\b|\bFIB(?:ER|RE)\b|\bEXCITEL\b|\bSPECTRA\b|\bTIKONA\b|\bRECHARGE\b|"
                       r"\bPOSTPAID\b|\bPREPAID\b|\bMOBILE BILL\b|\bTELEPHONE\b|\bINTERNET\b|\bLEASED LINE\b|\bWI-?FI\b|"
                       r"\bTATA ?TELE\w*", False),
    ("software", r"\bAMAZON WEB SERVICES\b|\bAWS\b|\bGOOGLE ?CLOUD\b|\bGCP\b|\bG ?SUITE\b|\bGOOGLE ?WORKSPACE\b|"
                 r"\bMICROSOFT\b|\bMSFT\b|\bOFFICE ?365\b|\bAZURE\b|\bZOHO\b|\bTALLY\b|\bCANVA\b|\bADOBE\b|\bOPENAI\b|"
                 r"\bCHATGPT\b|\bANTHROPIC\b|\bGODADDY\b|\bHOSTINGER\b|\bDIGITAL ?OCEAN\b|\bNOTION\b|\bSLACK\b|\bZOOM\b|"
                 r"\bDROPBOX\b|\bGITHUB\b|\bATLASSIAN\b|\bFRESHWORKS\b|\bWIX\b|\bSHOPIFY\b|\bBIGROCK\b|\bNAMECHEAP\b|"
                 r"\bAPPLE\.COM\b|\bICLOUD\b|\bCLOUDFLARE\b|\bBUSY INFOTECH\b|\bMARG ERP\b|\bSOFTWARE\b|"
                 r"\bSUBSCRIPTION\b|\bSAAS\b|\bHOSTING\b|\bDOMAIN\b|\bLICEN[CS]E\b", False),
    ("marketing", r"\bFACEBOOK\b|\bFACEBK\b|\bFB ADS\b|\bMETA\b|\bINSTAGRAM\b|\bGOOGLE ?ADS\b|\bADWORDS\b|\bJUST ?DIAL\b|"
                  r"\bINDIA ?MART\b|\bTRADE ?INDIA\b|\bSULEKHA\b|\bLINKEDIN\b|\bADVERTIS\w*|\bADS\b|\bHOARDINGS?\b|"
                  r"\bPAMPHLETS?\b|\bBANNERS?\b|\bFLEX\b|\bBRANDING\b|\bMARKETING\b|\bSEO\b|\bBULK SMS\b|\bAISENSY\b|"
                  r"\bINTERAKT\b|\bWATI\b|\bMSG91\b|\bTEXTLOCAL\b|\bGUPSHUP\b|\bMAILCHIMP\b|\bPROMOTIONS?\b|"
                  r"\bEXHIBITION\b|\bEXPO\b", False),
    ("travel", r"\bIRCTC\b|\bMAKE ?MY ?TRIP\b|\bGOIBIBO\b|\bYATRA\b|\bCLEARTRIP\b|\bEASE ?MY ?TRIP\b|\bIXIGO\b|\bINDIGO\b|"
               r"\bINTERGLOBE\b|\bAIR ?INDIA\b|\bVISTARA\b|\bSPICE ?JET\b|\bAKASA\b|\bAIR ?ASIA\b|\bUBER\b|\bOLA\b|"
               r"\bOLA ?CABS\b|\bANI TECHNOLOGIES\b|\bRAPIDO\b|\bRED ?BUS\b|\bABHIBUS\b|\bHOTELS?\b|\bOYO\b|\bTREEBO\b|"
               r"\bTAXI\b|\bCABS?\b|\bAIRPORT\b|\bAIRLINES?\b|\bTRAVELS?\b|\bTOURS?\b|\bRAILWAYS?\b|\bFLIGHTS?\b", False),
    ("freight", r"\bBLUE ?DART\b|\bDTDC\b|\bDELHIVERY\b|\bSHIPROCKET\b|\bECOM EXPRESS\b|\bXPRESSBEES\b|\bEKART\b|"
                r"\bSHADOWFAX\b|\bCOURIERS?\b|\bTRANSPORTS?\b|\bTRANSPORTERS?\b|\bROADWAYS?\b|\bROADLINES?\b|"
                r"\bLOGISTICS?\b|\bCARGO\b|\bFREIGHT\b|\bPORTER\b|\bLORRY\b|\bTEMPO\b|\bGATI\b|\bVRL\b|\bTCI\b|"
                r"\bSAFEXPRESS\b|\bPACKERS\b|\bMOVERS\b|\bSHIPPING\b|\bCARRIERS?\b|\bTRUCK\w*|\bDELIVERY\b", False),
    ("professional", r"\bCA FEES?\b|\bCA CHARGES\b|\bCHARTERED ACCOUNTANTS?\b|\bADVOCATES?\b|\bLEGAL\b|\bLAWYERS?\b|"
                     r"\bSOLICITORS?\b|\bCONSULT\w*|\bAUDIT\b|\bAUDIT FEES?\b|\bACCOUNTING\b|\bBOOK ?KEEPING\b|"
                     r"\bCOMPANY SECRETARY\b|\bCS FEES?\b|\bROC\b|\bMCA\b|\bPROFESSIONAL FEES?\b|\bPROF\.? FEES?\b|"
                     r"\bRETAINER\b|\bNOTARY\b|\bTRADEMARK\b|\bARCHITECTS?\b|\bVALUATION\b|\bASSOCIATES\b", False),
    ("food_staff", r"\bSWIGGY\b|\bZOMATO\b|\bCAFE\b|\bCAFÉ\b|\bRESTAURANTS?\b|\bDHABA\b|\bCANTEEN\b|\bTIFFIN\b|"
                   r"\bCATERERS?\b|\bCATERING\b|\bBAKERY\b|\bBAKERS\b|\bSWEETS\b|\bTEA\b|\bCHAI\b|\bDOMINOS\b|"
                   r"\bSTARBUCKS\b|\bMCDONALDS?\b|\bKFC\b|\bPIZZA\b|\bEATSURE\b|\bMESS\b|\bFOOD\b|\bSNACKS?\b|\bLUNCH\b", False),
    ("office", r"\bAMAZON\b(?! ?WEB SERVICES)|\bAMZN\b|\bFLIPKART\b|\bSTATION[AE]RY\b|\bXEROX\b|\bPRINTOUTS?\b|\bD ?MART\b|"
               r"\bBIG ?BASKET\b|\bBLINKIT\b|\bZEPTO\b|\bJIOMART\b|\bINSTAMART\b|\bHOUSEKEEPING\b|\bCLEANING\b|"
               r"\bWATER (?:CAN|JAR|SUPPLY)\b|\bDRINKING WATER\b|\bSECURITY (?:SERVICES?|AGENCY|GUARDS?)\b|\bOFFICE\b|"
               r"\bPANTRY\b|\bGROCER(?:Y|IES)\b", False),
    ("repairs", r"\bREPAIRS?\b|\bMAINTENANCE\b|\bSERVICING\b|\bSERVICE CENT(?:RE|ER)\b|\bPLUMB\w*|\bELECTRICIAN\b|"
                r"\bCARPENT\w*|\bPEST CONTROL\b|\bPAINT(?:ING|ER)\b|\bSPARES?\b|\bSPARE PARTS\b|\bWELDING\b|\bAMC\b", False),
    ("equipment", r"\bMACHINE(?:RY|S)?\b|\bEQUIPMENTS?\b|\bLAPTOPS?\b|\bCOMPUTERS?\b|\bPRINTERS?\b|\bTOOLS\b|"
                  r"\bGENERATOR\b|\bDG SET\b|\bAIR CONDITIONERS?\b|\bFURNITURE\b|\bCROMA\b|\bRELIANCE DIGITAL\b|"
                  r"\bVIJAY SALES\b|\bDELL\b|\bLENOVO\b|\bAPPLE\b|\bINVERTER\b|\bUPS\b", False),
    ("purchases", r"\bPURCHASES?\b|\bRAW MATERIALS?\b|\bMATERIALS?\b|\bSTEEL\b|\bCEMENT\b|\bHARDWARES?\b|\bTRADERS?\b|"
                  r"\bTRADING\b|\bENTERPRISES?\b|\bSUPPLIERS?\b|\bSUPPLY\b|\bSUPPLIES\b|\bINDUSTRIES\b|\bMETALS?\b|"
                  r"\bPLYWOODS?\b|\bPLY\b|\bTIMBERS?\b|\bCHEMICALS?\b|\bPLASTICS?\b|\bPACKAGING\b|\bWHOLESALE\b|"
                  r"\bDISTRIBUTORS?\b|\bAGENC(?:Y|IES)\b|\bSTORES?\b|\bINV(?:OICE)?\b|\bBILL NO\b|\bPO NO\b|"
                  r"\bFABRICS?\b|\bTEXTILES?\b|\bPAINTS\b|\bTILES\b|\bSANITARY\b|\bELECTRICALS\b|\bGLASS\b|"
                  r"\bALUMINIUM\b|\bMARBLES?\b|\bGRANITE\b|\bBRICKS?\b", False),
]
_CREDIT_RULES = [
    ("transfer_in", r"\b(?:AUTO ?)?SWEEP\b|\bOWN (?:A/?C|ACCOUNT|ACCT)\b|\bTRF FROM OWN\b|\bSELF TRANSFER\b|"
                    r"\bFD (?:PREMAT|CLOS|MATUR|REDEEM|PROCEEDS)|\bFIXED DEPOSIT\b|\bMUTUAL FUND\b|\bREDEMPTION\b|"
                    r"\bTD (?:CLOS|MATUR|PREMAT)", False),
    ("interest_in", r"\bINT\.? ?PD\b|\bINT PAID\b|\bINTEREST\b|\bSB INT\b|\bINT\.? ?CR\b|\bINT\.? ?COLL\b", False),
    ("refund_in", r"\bREFUND\b|\bRFND\b|\bREVERSAL\b|\bREVERSED\b|\bREV\b|\bCASH ?BACK\b|\bCHARGEBACK\b|\bRETURNED\b|"
                  r"\bRETURN\b|\bRTN\b|\bFAILED\b", False),
    ("loan_in", r"\bLOAN\b|\bDISB(?:URSAL|URSEMENT|URSED)?\b", False),
    ("cash_deposit", r"\bCASH ?DEP(?:OSIT)?\b|\bCDM\b|\bBNA\b|\bBY CASH\b|\bCASH CR\b|\bCSH DEP\b|"
                     r"\bDEPOSIT BY CASH\b|^CASH\b", False),
    ("other_in", r"\bDIVIDEND\b|\bDIV\b|\bSUBSIDY\b|\bSALARY\b|\bSAL\b|\bINCENTIVE\b|\bCLAIM\b", False),
    ("sales_receipt", r"\bNEFT\b|\bIMPS\b|\bRTGS\b|\bUPI\b|\bMMT\b|\bCHQ\b|\bCHEQUE\b|\bCLG\b|\bCLEARING\b|\bMICR\b|\bCTS\b|"
                      r"\bINWARD\b|\bBY TRANSFER\b|\bTRANSFER FROM\b|\bSETTL\w*|\bRAZORPAY\b|\bCASHFREE\b|\bPAYU\b|"
                      r"\bCCAVENUE\b|\bBILLDESK\b|\bINSTAMOJO\b|\bPHONEPE\b|\bPAYTM\b|\bBHARATPE\b|\bPINE ?LABS\b|"
                      r"\bSTRIPE\b|\bMERCHANT\b|\bECS CR\b|\bNACH CR\b|\bREMITTANCE\b|\bSWIFT\b|\bFIRC\b", False),
]
def _compile_rules(rules) -> list:
    return [(c, tuple(_rx(x) for x in (p if isinstance(p, tuple) else (p,))), g) for c, p, g in rules]


_RULES = {"debit": _compile_rules(_DEBIT_RULES), "credit": _compile_rules(_CREDIT_RULES)}


def _cat_text(t: dict) -> str:
    narr = str(t.get("narration") or "")[:CELL_CAP]
    p = str(t.get("payee") or "")[:60] if t.get("payee") is not None else payee(narr)
    text = _HANDLE.sub(" ", narr)          # UPI handles carry PSP names (@airtel, @jio) that aren't the payee
    return re.sub(r"\s+", " ", f"{text} {p}").strip().upper()


def _category(t: dict, extra: list) -> str:
    kind = "credit" if t.get("kind") == "credit" else "debit"
    text = _cat_text(t)
    for k, rx, cat in extra:
        if k == kind and rx.search(text):
            return cat
    p = t.get("payee")
    names_someone = bool(_STRUCTURED.search(text)) or bool(p if p is not None else payee(str(t.get("narration") or "")))
    for cat, rxs, no_counterparty_only in _RULES[kind]:
        if no_counterparty_only and names_someone:
            continue
        if all(rx.search(text) for rx in rxs):
            return cat
    return "other" if kind == "debit" else "other_in"


def _compile_extra(extra_rules) -> list:
    out = []
    for rule in extra_rules or []:
        try:
            kind, pattern, cat = rule
            if kind not in ("debit", "credit") or cat not in (DEBIT_CATEGORIES if kind == "debit" else CREDIT_CATEGORIES):
                continue
            out.append((kind, re.compile(pattern, re.I), cat))
        except (ValueError, TypeError, re.error):
            continue                                   # a broken rule is skipped, never fatal
    return out


def categorize(txns: list[dict], extra_rules: list[tuple[str, str, str]] | None = None) -> list[dict]:
    """Adds "category" to each transaction (in place) and returns them. extra_rules are (kind "debit"/"credit", regex,
    category), checked before the built-in rules; a rule with an unknown kind/category or a broken regex is skipped."""
    extra = _compile_extra(extra_rules)
    for t in txns:
        t["category"] = _category(t, extra)
    return txns


# ----------------------------------------------------------------------------------------------------------------------
# Recurring payments and summary

_MONTH_WORDS = re.compile(r"\b(?:" + _MON + r"|20\d\d|19\d\d)\b", re.I)


def _label(narration: str) -> str:
    s = _HANDLE.sub(" ", str(narration or "")[:CELL_CAP]).upper()
    s = re.sub(r"(?<![A-Z0-9])[A-Z]*\d[A-Z0-9]*|\bX{2,}\w*", " ", s)
    s = _MONTH_WORDS.sub(" ", s)
    s = re.sub(r"[^A-Z&]+", " ", s)
    return " ".join(s.split()[:6])


_LABEL_NOISE = {"CIN", "GSTIN", "MOB", "MICR", "TO", "BY", "DR", "CR", "REF", "NO", "TXN", "TRF", "TRANSFER"}


def payee_label(narration: str) -> str:
    """A short name for a payment with no payee ("Billpay MSEDCL Electricity", "GST Payment"): the narration without
    numbers, IDs, handles, month names or banking filler words. "" when nothing is left."""
    words = [w for w in _label(narration).split() if w not in _LABEL_NOISE][:4]
    return _title(" ".join(words)).replace("Sms ", "SMS ")


def _cat_of(t: dict) -> str:
    return t.get("category") or _category(t, [])


def _best_cluster(items: list[dict]) -> list[dict]:
    """The payments whose amounts sit within ±15% of a typical amount, choosing the typical amount that covers the most
    calendar months (then the most payments). Sorted amounts + a sliding window: O(n log n), not O(n²)."""
    items = sorted(items, key=lambda i: float(i["amount"]))
    amts = [float(i["amount"]) for i in items]
    months = [str(i["date"])[:7] for i in items]
    in_window, lo, hi, best = Counter(), 0, 0, None
    for k in range(len(items)):
        if k and amts[k] == amts[k - 1]:
            continue
        centre = amts[k]
        while hi < len(items) and amts[hi] <= centre * 1.15:
            in_window[months[hi]] += 1
            hi += 1
        while amts[lo] < centre * 0.85:
            in_window[months[lo]] -= 1
            if not in_window[months[lo]]:
                del in_window[months[lo]]
            lo += 1
        score = (len(in_window), hi - lo)
        if best is None or score > best[0]:
            best = (score, lo, hi)
    _, lo, hi = best
    typical = statistics.median(amts[lo:hi])
    return items[bisect_left(amts, typical * 0.85):bisect_right(amts, typical * 1.15)]


def recurring(txns: list[dict]) -> list[dict]:
    """Debits to the same payee (or the same narration, numbers and month names removed) in the same category, at an
    amount within ±15% of the typical one, in at least 2 different calendar months. Cash withdrawals and transfers to
    own accounts are left out. Sorted by typical amount, biggest first."""
    groups = defaultdict(list)
    for t in txns:
        if t.get("kind") != "debit" or not float(t.get("amount") or 0) > 0:
            continue
        cat = _cat_of(t)
        if cat in NOT_EXPENSES:
            continue
        name = (t.get("payee") if t.get("payee") is not None else payee(t.get("narration", ""))) or ""
        key = name.strip().lower() or "~" + _label(t.get("narration", ""))
        if key in ("", "~"):
            continue
        groups[(key, cat)].append(t)
    out = []
    for (key, cat), items in groups.items():
        if len(items) < 2:
            continue
        members = _best_cluster(items)
        months = {str(i["date"])[:7] for i in members}
        if len(months) < 2:
            continue
        days = sorted(int(str(i["date"])[8:10]) for i in members)
        name = next((i.get("payee") for i in members if i.get("payee")), "") or _title(key.lstrip("~"))[:60]
        out.append({"payee": name, "category": cat, "amount": round(statistics.median(float(i["amount"]) for i in members), 2),
                    "day_of_month": statistics.median_low(days), "months": len(months),
                    "last_date": max(str(i["date"]) for i in members)})
    out.sort(key=lambda r: (-r["amount"], r["payee"]))
    return out


def summary(txns: list[dict]) -> dict:
    """Month-by-month money in/out, debits by category, top 10 payees by money paid (cash and own-account transfers
    excluded), recurring payments, and totals. Transactions without a "category" are categorised on the fly (not saved)."""
    months = defaultdict(lambda: {"in": 0.0, "out": 0.0})
    by_cat = defaultdict(float)
    payees = defaultdict(lambda: [0.0, 0])
    money_in = money_out = cash = transfers = 0.0
    enriched = []
    for t in txns:
        amt = float(t.get("amount") or 0)
        cat = _cat_of(t)
        enriched.append({**t, "category": cat})
        m = str(t.get("date", ""))[:7]
        if t.get("kind") == "credit":
            months[m]["in"] += amt
            money_in += amt
            continue
        months[m]["out"] += amt
        money_out += amt
        by_cat[cat] += amt
        if cat == "cash_withdrawal":
            cash += amt
        elif cat == "transfer":
            transfers += amt
        name = (t.get("payee") if t.get("payee") is not None else payee(t.get("narration", ""))) or ""
        if name and cat not in NOT_EXPENSES:
            payees[name][0] += amt
            payees[name][1] += 1
    top = sorted(payees.items(), key=lambda kv: (-kv[1][0], kv[0]))[:10]
    return {
        "months": [{"month": k, "in": round(v["in"], 2), "out": round(v["out"], 2)} for k, v in sorted(months.items())],
        "by_category": {k: round(v, 2) for k, v in sorted(by_cat.items(), key=lambda kv: -kv[1])},
        "top_payees": [{"payee": k, "amount": round(v[0], 2), "count": v[1]} for k, v in top],
        "recurring": recurring(enriched),
        "cash_withdrawn": round(cash, 2), "transfers_out": round(transfers, 2),
        "money_in": round(money_in, 2), "money_out": round(money_out, 2),
    }


# ----------------------------------------------------------------------------------------------------------------------
# Reading tables (CSV, Excel, HTML, PDF tables)

_OPEN_RE = _rx(r"\bopening\s+bal(?:ance)?\b|\bop\.?\s*bal(?:ance)?\b|\bbalance\s+b\s*/\s*f\b|^\s*b\s*/\s*f\b|"
               r"\bbrought\s+forward\b|\bbalance\s+brought\b|\bbal\.?\s+b\s*/\s*f\b")
_CLOSE_RE = _rx(r"\bclosing\s+bal(?:ance)?\b|\bcl\.?\s*bal(?:ance)?\b|\bbalance\s+c\s*/\s*f\b|^\s*c\s*/\s*f\b|"
                r"\bcarried\s+forward\b|\bbalance\s+carried\b")
_ASON_RE = _rx(r"\bbalance\s+as\s+(?:on|of|at)\b")
_TOTAL_RE = _rx(r"^\s*(?:grand\s+|transaction\s+|sub\s*-?\s*|page\s+)?totals?\s*:?\s*$|^\s*total\s+(?:debits?|credits?|"
                r"withdrawals?|deposits?|amount)\b")
_END_RE = _rx(r"\bend\s+of\s+(?:the\s+)?statement\b|^\s*statement\s+summary\b")
# An amount written out in text: needs paise or Indian comma grouping, and must not be part of a date like 01.04.2026.
_STRICT_AMT = re.compile(r"(?<![\w./\-])\(?-?(?:₹|Rs\.?|INR)?\s?(?:\d+(?:,\d{2,3})*\.\d{2}|\d{1,3}(?:,\d{2,3})+)"
                         r"(?!\d|[./\-]\d)\)?(?:\s?(?:Cr|Dr|CR|DR|cr|dr)\b\.?)?")


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)[:CELL_CAP].replace("\xa0", " ")).strip()


def _role(text: str) -> str | None:
    h = re.sub(r"[^a-z0-9/]+", " ", text.lower()).strip()
    h = re.sub(r"\s*/\s*", "/", h)
    h = re.sub(r"\b(?:inr|rs)\b", "", h).strip()
    if not h or len(h) > 40:
        return None
    if re.fullmatch(r"dr/cr|cr/dr|d/c|c/d|debit/credit|credit/debit|dr cr|cr dr|drcr|crdr|type|txn type|tran type|"
                    r"transaction type|(?:cr/dr|dr/cr) indicator|indicator", h):
        return "drcr"
    if re.search(r"\bbal(?:ance)?\b", h):
        return "balance"
    if re.search(r"\bvalue\b", h) and re.search(r"\b(?:date|dt)\b", h):
        return "value_date"
    if re.search(r"\b(?:date|dt)\b", h):
        return "date"
    if re.search(r"\b(?:narration|naration|description|particulars|remarks|details)\b", h):
        return "narration"
    if re.search(r"\b(?:withdrawals?|debits?|dr|paid out|money out)\b", h):
        return "debit"
    if re.search(r"\b(?:deposits?|credits?|cr|paid in|money in)\b", h):
        return "credit"
    if re.search(r"\b(?:amount|amt)\b", h):
        return "amount"
    if re.search(r"\b(?:chq|cheque|check|ref|reference|utr|instrument)\b", h):
        return "ref"
    return None


def _header_map(cells: list[str]) -> dict | None:
    roles, drcr_cols = {}, []
    for i, t in enumerate(cells):
        r = _role(t) if t and len(t) <= 60 else None
        if r == "drcr":
            drcr_cols.append(i)
        elif r and r not in roles:
            roles[r] = i
    if not ("date" in roles or "value_date" in roles) or "narration" not in roles:
        return None
    if not (("debit" in roles and "credit" in roles) or "amount" in roles):
        return None
    for i in drcr_cols:   # a Dr/Cr column belongs to the amount or balance column just before it
        before = [(roles[k], k) for k in ("amount", "balance", "debit", "credit") if k in roles and roles[k] < i]
        owner = max(before)[1] if before else "amount"
        key = "balance_drcr" if owner == "balance" else "drcr"
        roles.setdefault(key, i)
    if "debit" in roles and "credit" in roles:
        roles.pop("amount", None)
    return roles


class _Collected:
    def __init__(self):
        self.txns: list[dict] = []
        self.opening = None
        self.closing = None
        self.pre: list[list[str]] = []      # rows/lines before the header: bank name, account number
        self.post: list[list[str]] = []     # rows/lines after the transactions: summary, closing balance
        self.skipped = 0
        self.both_sides = 0
        self.signed = False                 # a single amount column used minus signs for debits
        self.soft = set()                   # "open"/"close" read from "Balance as on <date>" (could be today's balance)

    def add(self, t: dict):
        self.txns.append(t)
        if len(self.txns) > MAX_TXNS:
            raise StatementError("too_big", "This statement has more than 5,000 transactions. Upload fewer months at a "
                                            "time (for example 3 months per file).")

    def drop(self, pending: dict):
        """A dated row that never got an amount: before the first transaction it's preamble, after it a skipped row."""
        if self.txns:
            self.skipped += 1
        else:
            self.pre.append([pending["narr"]])

    def set_open(self, v):
        if v is not None and self.opening is None:
            self.opening = v

    def set_close(self, v):
        if v is not None:
            self.closing = v


def _monthfirst(tables: list[list[list]]) -> bool:
    """Indian statements are day-first; switch only when a date can't be day-first (e.g. 04/25/2026) and none is
    clearly day-first."""
    first_big = second_big = False
    for table in tables:
        for row in table:
            for v in row:
                if isinstance(v, str) and len(v) <= 24:
                    m = _D_NUM.match(v.strip())
                    if m and _TIME_REST.fullmatch(v.strip()[m.end():]):
                        first_big |= int(m.group(1)) > 12
                        second_big |= int(m.group(2)) > 12
    return second_big and not first_big


def _join(a: str, b: str) -> str:
    if not a:
        return b[:NARR_CAP]
    if len(a) >= NARR_CAP:
        return a
    if a.endswith(("/", "-")) or b.startswith(("/", "-")):
        return (a + b)[:NARR_CAP]
    return (a + " " + b)[:NARR_CAP]


def _read_tables(tables: list[list[list]]) -> _Collected:
    """Rows from one statement (one or more tables in order). Finds the header row, then reads transactions; a header
    found again later (next page) is skipped or replaces the column map."""
    c = _Collected()
    mf = _monthfirst(tables)
    hm = None
    last = None          # the transaction a wrapped narration row belongs to
    pending = None       # a dated row with no amount yet (its amounts may be on the next row)
    ended = False
    joins = 0

    def cell(row, role):
        i = hm.get(role)
        return row[i] if i is not None and i < len(row) else None

    def txt(row, role):
        return _cell_text(cell(row, role))

    def amounts(row):
        """[(amount, kind or None, sign)] for a row, plus the balance."""
        out = []
        if "debit" in hm:
            d, _ = _amount(cell(row, "debit"))
            cr, _ = _amount(cell(row, "credit"))
            if d:
                out.append((abs(d), "debit", 0))
            if cr:
                out.append((abs(cr), "credit", 0))
        else:
            a, mk = _amount(cell(row, "amount"))
            if a:
                mk = mk or _drcr(cell(row, "drcr"))
                kind = {"dr": "debit", "cr": "credit"}.get(mk)
                out.append((abs(a), kind, -1 if a < 0 else 1))
        bal = _signed_balance(cell(row, "balance"), cell(row, "balance_drcr"))
        return out, bal

    def emit(d, narr, amts, bal):
        nonlocal last, joins
        if len(amts) == 2:
            c.both_sides += 1
        for a, kind, sign in amts:
            if sign < 0 and kind is None:
                c.signed = True
            t = {"date": d, "narr": narr, "amount": round(a, 2), "kind": kind, "sign": sign, "balance": bal}
            c.add(t)
            last, joins = t, 0

    for table in tables:
        for raw in table[:MAX_ROWS]:
            row = list(raw[:MAX_COLS]) if raw else []
            cells = [_cell_text(v) for v in row]
            if not any(cells):
                last = None
                continue
            if ended:
                c.post.append(cells)
                continue
            found = _header_map(cells)
            if found:
                hm, last, pending = found, None, None
                continue
            if hm is None:
                c.pre.append(cells)
                continue
            line = " ".join(x for x in cells if x)
            if _END_RE.search(line):
                ended = True
                c.post.append(cells)
                continue
            d = _to_date(cell(row, "date"), mf)
            if d is None and not txt(row, "date"):
                d = _to_date(cell(row, "value_date"), mf)
            narr = txt(row, "narration")
            label = narr or line
            amts, bal = amounts(row)
            ason = bool(_ASON_RE.search(label))
            if _OPEN_RE.search(label) or (ason and not c.txns):
                if ason and not _OPEN_RE.search(label):
                    c.soft.add("open")
                c.set_open(bal if bal is not None else (amts[0][0] if amts else None))
                last = pending = None
                continue
            if _CLOSE_RE.search(label) or (ason and c.txns):
                if ason and not _CLOSE_RE.search(label):
                    c.soft.add("close")
                c.set_close(bal if bal is not None else (amts[0][0] if amts else None))
                last = pending = None
                continue
            if d is None:
                if _TOTAL_RE.search(label) or _TOTAL_RE.search(line):
                    c.post.append(cells)
                    last = pending = None
                    continue
                if pending and amts:
                    emit(pending["date"], _join(pending["narr"], narr), amts, bal)
                    pending = None
                    continue
                if amts:
                    c.post.append(cells)
                    last = None
                    continue
                others = [x for i, x in enumerate(cells) if x and i not in (hm.get("narration"), hm.get("ref"))]
                if narr and not others and pending:
                    pending["narr"] = _join(pending["narr"], narr)
                    continue
                if narr and not others and last is not None and joins < 5:
                    last["narr"] = _join(last["narr"], narr)
                    joins += 1
                    continue
                c.post.append(cells)
                last = None
                continue
            if _TOTAL_RE.search(narr):
                last = pending = None
                continue
            if pending:
                c.drop(pending)
            pending = None
            if not amts:
                pending = {"date": d, "narr": narr}
                last = None
                continue
            emit(d, narr, amts, bal)
    if pending:
        c.drop(pending)
    return c


# ----------------------------------------------------------------------------------------------------------------------
# Reading text lines (PDFs without table lines, fixed-width text)

_TXT_AMT = r"\(?-?(?:₹|Rs\.?|INR)?\s?\d+(?:,\d{2,3})*\.\d{2}\)?(?:\s?(?:Cr|Dr|CR|DR|cr|dr)\b\.?)?"
_TAIL = re.compile(r"(?:^|\s)(" + _TXT_AMT + r"|-)\s*$")
_BOILER = _rx(r"^page\s*\d+|\bpage\s+\d+\s+(?:of|/)\s*\d+|statement\s+of\s+account|account\s+statement|computer[- ]generated|"
              r"does\s+not\s+require|\bcontinued\b|^\*{3,}|customer\s+id|\bcust\s*id\b|\bifsc\b|\bmicr\b|\bbranch\b|"
              r"account\s+(?:no|number)|a/c\s+no|\bperiod\b|\bnominee\b|registered\s+office|toll[- ]free|call\s+us|"
              r"www\.|https?://|generated\s+on|printed\s+on|statement\s+(?:for|from|date)")


def _peel(s: str) -> tuple[list[str], str]:
    """Amount tokens at the end of a line (up to 3; '-' as an empty column only before a real amount)."""
    toks = []
    while len(toks) < 3:
        m = _TAIL.search(s)
        if not m:
            break
        tok = m.group(1)
        if tok == "-" and not toks:
            break
        toks.insert(0, tok)
        s = s[:m.start()].rstrip()
    return toks, s


def _text_header(line: str) -> tuple[list[str], bool] | None:
    """A column-header line -> (the amount columns in order, e.g. ['debit', 'credit', 'balance'], whether a Value Date
    column comes before the transaction date)."""
    low = line.lower()
    if not re.search(r"\b(?:date|dt)\b", low):
        return None
    if not re.search(r"\b(?:narration|description|particulars|remarks|details)\b", low):
        return None
    low = re.sub(r"\b(?:dr\s*/\s*cr|cr\s*/\s*dr)\b", " drcr ", low)
    roles = []
    for w in re.findall(r"[a-z]+", low):
        r = {"withdrawal": "debit", "withdrawals": "debit", "debit": "debit", "debits": "debit", "dr": "debit",
             "deposit": "credit", "deposits": "credit", "credit": "credit", "credits": "credit", "cr": "credit",
             "amount": "amount", "balance": "balance", "bal": "balance"}.get(w)
        if r and (not roles or roles[-1] != r) and r not in roles:
            roles.append(r)
    if "debit" in roles and "credit" in roles:
        roles = [r for r in roles if r != "amount"]
    elif "amount" in roles:
        roles = [r for r in roles if r in ("amount", "balance")]
    else:
        return None
    dates = re.findall(r"(value\s+)?\b(?:date|dt)\b", line.lower())
    return roles, len(dates) >= 2 and bool(dates[0])


def _strip_trailing_date(s: str, txn_date: date) -> str:
    """Text PDFs print the Value Date column between the narration and the amounts; drop it from the narration (only a
    date within a week of the transaction, and not the end of a range like "01-04-2026 to 30-04-2026")."""
    parts = s.split(" ")
    for n in (3, 2, 1):
        if len(parts) > n:
            d = _to_date(" ".join(parts[-n:]))
            if d and abs((d - txn_date).days) <= 7 and parts[-n - 1].lower() not in ("to", "till", "-", "from", "upto"):
                return " ".join(parts[:-n])
    return s


def _line_date(line: str, mf: bool, value_first: bool = False) -> tuple[date, str] | None:
    """A line that starts with a date (optionally after a serial number) -> (transaction date, rest of the line)."""
    hit = _match_date(line, mf)
    start = 0
    if not hit:
        m = re.match(r"\d{1,4}\s+", line)           # a serial-number column before the date
        if not m:
            return None
        hit = _match_date(line[m.end():], mf)
        if not hit:
            return None
        start = m.end()
    d, rest = hit[0], line[start + hit[1]:].strip()
    again = _match_date(rest, mf)                   # a second date: the value date (or, if it comes first, the txn date)
    if again:
        if value_first:
            d = again[0]
        rest = rest[again[1]:].strip()
    return d, rest


def _tok_value(toks: list[str]) -> float | None:
    """The last amount on an opening/closing balance line, negative when marked Dr (an overdraft)."""
    v, mk = _amount(toks[-1]) if toks else (None, None)
    return (-abs(v) if mk == "dr" else v) if v is not None else None


def _assign(toks: list[str], roles: list[str] | None):
    """Amount tokens -> ([(amount, kind or None, sign)], balance)."""
    vals = []
    for tok in toks:
        v, mk = (None, None) if tok == "-" else _amount(tok)
        vals.append((v, mk))
    if roles is None or len(roles) != len(vals):
        if roles and "balance" in roles and len(vals) == 2:
            roles = ["amount", "balance"]
        elif roles and "balance" not in roles and len(vals) == 1:
            roles = ["amount"]
        else:
            roles = {3: ["debit", "credit", "balance"], 2: ["amount", "balance"], 1: ["amount"]}.get(len(vals), [])
    out, bal = [], None
    for (v, mk), r in zip(vals, roles):
        if r == "balance":
            bal = (-abs(v) if mk == "dr" else v) if v is not None else None
        elif v:
            kind = "debit" if r == "debit" else "credit" if r == "credit" else {"dr": "debit", "cr": "credit"}.get(mk)
            out.append((abs(v), kind, -1 if v < 0 else 1))
    return out, bal


def _read_lines(lines: list[str]) -> _Collected:
    c = _Collected()
    mf = _monthfirst([[[ln.split(" ")[0]] for ln in lines if ln.strip()]])
    roles, value_first = None, False
    last = None
    pending = None
    joins = 0
    seen = False
    ended = False

    def emit(d, narr, amts, bal):
        nonlocal last, joins, seen
        if len(amts) == 2:
            c.both_sides += 1
        for a, kind, sign in amts:
            if sign < 0 and kind is None:
                c.signed = True
            t = {"date": d, "narr": narr, "amount": round(a, 2), "kind": kind, "sign": sign, "balance": bal}
            c.add(t)
            last, joins, seen = t, 0, True

    for raw in lines[:MAX_ROWS]:
        if raw == PAGE_BREAK:            # a page header/footer must never be glued onto the last row of the page
            if pending:
                c.drop(pending)
            last = pending = None
            continue
        line = re.sub(r"\s+", " ", raw[:LINE_CAP * 2].replace("\xa0", " ")).strip()[:LINE_CAP]
        if not line:
            continue
        if ended:
            c.post.append([line])
            continue
        hdr = _text_header(line) if not _line_date(line, mf) else None
        if hdr:
            (roles, value_first), last, pending = hdr, None, None
            continue
        if _END_RE.search(line):
            ended = True
            c.post.append([line])
            continue
        md = _line_date(line, mf, value_first)
        if md is None:
            toks, head = _peel(line)
            amts, bal = _assign(toks, roles) if toks else ([], None)
            if _OPEN_RE.search(head) or (_ASON_RE.search(head) and not seen):
                if _ASON_RE.search(head) and not _OPEN_RE.search(head) and toks:
                    c.soft.add("open")
                c.set_open(_tok_value(toks))
                if not toks:
                    (c.post if seen else c.pre).append([line])
                last = pending = None
                continue
            if _CLOSE_RE.search(head) or (_ASON_RE.search(head) and seen):
                if _ASON_RE.search(head) and not _CLOSE_RE.search(head) and toks:
                    c.soft.add("close")
                c.set_close(_tok_value(toks))
                if not toks:
                    c.post.append([line])
                last = pending = None
                continue
            if toks and pending and not _TOTAL_RE.search(head) and not _BOILER.search(head):
                emit(pending["date"], _join(pending["narr"], head), amts, bal)
                pending = None
                continue
            if toks:
                (c.post if seen else c.pre).append([line])
                last = pending = None
                continue
            if _BOILER.search(line) or len(line) > 90:
                (c.post if seen else c.pre).append([line])
                last = pending = None
                continue
            if pending and joins < 3:
                pending["narr"] = _join(pending["narr"], line)
                joins += 1
                continue
            if last is not None and joins < 3:
                last["narr"] = _join(last["narr"], line)
                joins += 1
                continue
            (c.post if seen else c.pre).append([line])
            continue
        d, rest = md
        toks, narr = _peel(rest)
        narr = _strip_trailing_date(narr, d)
        amts, bal = _assign(toks, roles) if toks else ([], None)
        if _OPEN_RE.search(narr):
            c.set_open(_tok_value(toks))
            last = pending = None
            continue
        if _CLOSE_RE.search(narr):
            c.set_close(_tok_value(toks))
            last = pending = None
            continue
        if _TOTAL_RE.search(narr):
            last = pending = None
            continue
        if pending:
            c.drop(pending)
        pending = None
        if not amts:
            pending, last, joins = {"date": d, "narr": narr}, None, 0
            continue
        emit(d, narr, amts, bal)
    if pending:
        c.drop(pending)
    return c


# ----------------------------------------------------------------------------------------------------------------------
# File formats

def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _read_csv(text: str) -> tuple[list[list[list]], list[str]]:
    text = text.replace("\x00", "")
    lines = text.splitlines()[:MAX_ROWS]
    best, best_score = ",", -1
    for delim in (",", ";", "\t", "|"):
        try:
            rows = list(csv.reader(lines[:400], delimiter=delim))
        except csv.Error:
            continue
        score = sum(1 for r in rows if sum(1 for x in r if x.strip()) >= 4)
        if score > best_score:
            best, best_score = delim, score
    rows = []
    for row in csv.reader(io.StringIO(text, newline=""), delimiter=best):   # quoted narrations may span lines
        rows.append([re.sub(r"\s+", " ", x[:CELL_CAP]) for x in row[:MAX_COLS]])
        if len(rows) >= MAX_ROWS:
            break
    return [rows], lines


class _TooComplex(Exception):
    pass


class _TableReader(HTMLParser):
    """Rows of every <table> in an HTML page (or every <Worksheet> of Excel-2003 XML), in one pass with no document
    tree: memory and time grow with the file, never faster. A cell keeps its own text, not that of tables nested in it.
    Entities are only the standard HTML ones (no XML entity expansion). Too deep, too many tables or cells -> _TooComplex."""

    def __init__(self, xml: bool):
        super().__init__(convert_charrefs=True)
        self.xml = xml
        self.table_tag, self.row_tag = ("worksheet", "row") if xml else ("table", "tr")
        self.cell_tags = ("cell",) if xml else ("td", "th")
        self.sheets: list[list[list]] = []      # xml: one list of rows per worksheet; html: a single list
        self.stack: list[dict] = []             # open tables: {"row": [...] | None, "cell": [...] | None, ...}
        self.skip = 0
        self.tables = self.cells = 0
        self.rows: list[list] = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
            return
        if tag == self.table_tag:
            self.tables += 1
            if self.tables > MAX_HTML_TABLES or len(self.stack) >= MAX_HTML_DEPTH:
                raise _TooComplex()
            if self.xml:
                self.rows = []
                self.sheets.append(self.rows)
            self.stack.append({"row": None, "cell": None, "len": 0, "type": ""})
            return
        if not self.stack:
            return
        t = self.stack[-1]
        if tag == self.row_tag:
            self._end_row(t)
            t["row"] = []
        elif tag in self.cell_tags:
            self._end_cell(t)
            if t["row"] is None:
                t["row"] = []
            a = dict(attrs)
            if self.xml and (a.get("ss:index") or "").isdigit():
                t["row"].extend([None] * max(0, min(int(a["ss:index"]), MAX_COLS) - 1 - len(t["row"])))
            t["cell"], t["len"], t["type"] = [], 0, ""
            t["span"] = 1
            if not self.xml:
                try:
                    t["span"] = max(1, min(int(a.get("colspan") or 1), 20))
                except ValueError:
                    pass
            self.cells += 1
            if self.cells > MAX_HTML_CELLS:
                raise _TooComplex()
        elif tag == "data" and self.xml:
            t["type"] = dict(attrs).get("ss:type") or ""
        elif tag == "br" and t["cell"] is not None:
            t["cell"].append(" ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
            return
        if not self.stack:
            return
        t = self.stack[-1]
        if tag in self.cell_tags:
            self._end_cell(t)
        elif tag == self.row_tag:
            self._end_row(t)
        elif tag == self.table_tag:
            self._end_row(t)
            self.stack.pop()

    def handle_data(self, data):
        if self.skip or not self.stack:
            return
        t = self.stack[-1]
        if t["cell"] is not None and t["len"] < CELL_CAP:
            t["cell"].append(data[:CELL_CAP])
            t["len"] += len(data)

    def _end_cell(self, t):
        if t["cell"] is None:
            return
        val = re.sub(r"\s+", " ", "".join(t["cell"])[:CELL_CAP]).strip()
        if self.xml and t["type"] == "Number":
            try:
                val = float(val)
            except ValueError:
                pass
        elif self.xml and t["type"] == "DateTime":
            hit = _match_date(val)
            val = hit[0] if hit else val
        t["row"].append(val)
        if not self.xml:
            t["row"].extend([""] * (t.get("span", 1) - 1))
        t["cell"] = None

    def _end_row(self, t):
        self._end_cell(t)
        if t["row"] is not None:
            if len(self.rows) < MAX_ROWS:
                self.rows.append(t["row"][:MAX_COLS])
            t["row"] = None


def _read_markup(text: str, xml: bool) -> list[list[list]]:
    reader = _TableReader(xml)
    try:
        reader.feed(text)
        reader.close()
    except _TooComplex:
        raise StatementError("unreadable", MSG_TOO_COMPLEX) from None
    for t in list(reader.stack):           # tables left open at the end of a damaged file
        reader._end_row(t)
    return [s for s in reader.sheets if s] if xml else [reader.rows]


def _read_html(text: str) -> list[list[list]]:
    return _read_markup(text, xml=False)


def _read_spreadsheet_xml(text: str) -> list[list[list]]:
    """Excel 2003 XML ("SpreadsheetML"), which some banks save as .xls."""
    return _read_markup(text, xml=True)


def _read_xlsx(data: bytes) -> list[list[list]]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except (zipfile.BadZipFile, OSError, ValueError):
        raise StatementError("unreadable", "We couldn't open this Excel file — it may be damaged. Download the statement "
                                           "again from net banking and try again.") from None
    if sum(i.file_size for i in infos) > MAX_UNZIPPED or len(infos) > 5000:
        raise StatementError("too_big", "This Excel file is too large to open safely. Download fewer months at a time, "
                                        "or the CSV version, and upload again.")
    names = {i.filename for i in infos}
    if "xl/workbook.xml" not in names:
        raise _unsupported()
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets[:10]:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    break
                rows.append(list(r[:MAX_COLS]))
            sheets.append(rows)
        return sheets
    finally:
        wb.close()


def _read_xls(data: bytes) -> list[list[list]]:
    if b"E\x00n\x00c\x00r\x00y\x00p\x00t\x00e\x00d\x00P\x00a\x00c\x00k\x00a\x00g\x00e" in data:
        raise _locked_excel()
    try:
        import xlrd
    except ImportError:
        raise StatementError("unsupported", "This older Excel (.xls) file can't be read here yet. Download the statement "
                                            "as CSV, Excel (.xlsx) or PDF and try again.") from None
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as e:  # noqa: BLE001  xlrd raises several error types for damaged or locked files
        if "encrypt" in str(e).lower():
            raise _locked_excel() from None
        raise StatementError("unreadable", "We couldn't open this Excel file — it may be damaged. Download the statement "
                                           "again from net banking and try again.") from None
    try:
        sheets = []
        for si in range(min(book.nsheets, 10)):
            sh = book.sheet_by_index(si)
            rows = []
            for r in range(min(sh.nrows, MAX_ROWS)):
                row = []
                for ci in range(min(sh.ncols, MAX_COLS)):
                    cl = sh.cell(r, ci)
                    if cl.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate_as_datetime(cl.value, book.datemode))
                        except Exception:  # noqa: BLE001
                            row.append(cl.value)
                    elif cl.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        row.append(None)
                    else:
                        row.append(cl.value)
                rows.append(row)
            sheets.append(rows)
            book.unload_sheet(si)
        return sheets
    finally:
        book.release_resources()


def _read_pdf(data: bytes, password: str | None) -> tuple[_Collected, bool, list[str]]:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect
    from pdfplumber.utils.exceptions import PdfminerException

    def locked(e) -> bool:
        return isinstance(e, PDFPasswordIncorrect) or any(isinstance(a, PDFPasswordIncorrect) for a in getattr(e, "args", ()))

    try:
        pdf = pdfplumber.open(io.BytesIO(data), password=password or "")
    except (PdfminerException, PDFPasswordIncorrect) as e:
        if _out_of_memory(e):
            raise StatementError("unreadable", MSG_TOO_COMPLEX) from None
        if locked(e):
            if password:
                raise StatementError("wrong_password", "That password didn't open the PDF. Check it and try again — the "
                                                       "bank's e-mail with the statement usually explains the password "
                                                       "format.") from None
            raise StatementError("password_needed", "This PDF is locked with a password. Type the statement password "
                                                    "(the bank's e-mail usually explains it) and try again.") from None
        raise StatementError("unreadable", "We couldn't open this PDF — it may be damaged. Download the statement again "
                                           "from net banking (Excel or CSV works best) and try again.") from None
    with pdf:
        pages = pdf.pages
        if len(pages) > MAX_PAGES:
            raise StatementError("too_big", "This PDF has more than 60 pages. Download fewer months at a time (or the "
                                            "Excel/CSV version) and upload again.")
        tables, lines, chars, content = [], [], 0, 0
        for p in pages:
            size = _content_size(p, PAGE_CONTENT_CAP)
            content += size
            if size > PAGE_CONTENT_CAP or content > PDF_CONTENT_CAP:   # a decompression bomb: stop before laying it out
                raise StatementError("unreadable", MSG_TOO_COMPLEX)
            text = p.extract_text() or ""
            chars += len(text.strip())
            lines.extend(text.splitlines())
            lines.append(PAGE_BREAK)
            try:
                for t in p.extract_tables():
                    tables.append([[(x or "").replace("\n", " ") for x in row] for row in t])
            except Exception as e:  # noqa: BLE001  a page whose drawing confuses table detection still has its text lines
                if _out_of_memory(e):
                    raise
            p.close()
    by_table = _read_tables([[r for t in tables for r in t]]) if tables else _Collected()
    by_text = _read_lines(lines)
    if not by_table.txns and not by_text.txns:
        garbled = sum(ln.count("(cid:") for ln in lines) > 20       # fonts with no text mapping: looks like text, isn't
        if chars < 30 * max(1, len(pages)) or garbled:
            raise StatementError("unreadable", MSG_SCAN)
        raise StatementError("no_transactions", MSG_NO_TXNS)
    if by_table.txns and len(by_table.txns) >= len(by_text.txns):
        by_table.pre = by_text.pre + by_table.pre    # the bank name and account number are usually above the table
        by_table.post = by_table.post + by_text.post
        if by_table.opening is None:
            by_table.opening = by_text.opening
        if by_table.closing is None:
            by_table.closing = by_text.closing
        return by_table, False, lines
    return by_text, True, lines


def _content_size(page, cap: int) -> int:
    """How many bytes a page's drawing commands decode to, counting no further than cap + 1, so a decompression bomb is
    measured without being inflated. Best effort: anything unusual counts as 0 and the child-process limits still apply."""
    try:
        import zlib
        from pdfminer.ascii85 import ascii85decode, asciihexdecode
        from pdfminer.pdftypes import PDFStream, resolve1
        total = 0
        for ref in page.page_obj.contents:
            st = resolve1(ref)
            if not isinstance(st, PDFStream):
                continue
            if st.rawdata is None:                 # already decoded by the library
                total += len(st.data or b"")
                continue
            data = st.rawdata
            if st.decipher:
                data = st.decipher(st.objid, st.genno, data, st.attrs)
            for f, _ in st.get_filters():
                name = getattr(f, "name", f)
                if name in ("FlateDecode", "Fl"):
                    data = zlib.decompressobj().decompress(data, max(1, cap - total + 1))
                elif name in ("ASCII85Decode", "A85"):
                    data = ascii85decode(data)
                elif name in ("ASCIIHexDecode", "AHx"):
                    data = asciihexdecode(data)
                else:
                    break
            total += len(data)
            if total > cap:
                break
        return total
    except Exception:  # noqa: BLE001
        return 0


def _unsupported() -> StatementError:
    return StatementError("unsupported", "We can't read this kind of file. Upload the bank statement as PDF, Excel "
                                         "(.xls or .xlsx) or CSV.")


def _locked_excel() -> StatementError:
    return StatementError("unsupported", "This Excel file is locked with a password, and only locked PDFs can be opened "
                                         "here. Open it in Excel and save a copy without the password, or download the "
                                         "CSV version, and upload that.")


_IMAGE_SIGS = (b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"II*\x00", b"MM\x00*")


def _sniff(data: bytes) -> str:
    head = data[:2048]
    if b"%PDF" in data[:1024]:
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    webp_or_heic = (head[:4] == b"RIFF" and head[8:12] == b"WEBP") or head[4:12] in (b"ftypheic", b"ftypmif1")
    if head.startswith(_IMAGE_SIGS) or webp_or_heic:
        return "image"
    if b"\x00" in head and not head.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "binary"
    return "text"


# ----------------------------------------------------------------------------------------------------------------------
# Bank, account, and putting it together

_BANK_PATTERNS = [
    ("HDFC Bank", r"\bHDFC\s*BANK\b|\bHDFC0|\bHDFC\b"),
    ("ICICI Bank", r"\bICICI\b|\bICIC0"),
    ("State Bank of India", r"STATE\s+BANK\s+OF\s+INDIA|\bSBIN0|\bSBI\b"),
    ("Axis Bank", r"\bAXIS\s*BANK\b|\bUTIB0"),
    ("Kotak Mahindra Bank", r"\bKOTAK\b|\bKKBK0"),
    ("Yes Bank", r"\bYES\s*BANK\b|\bYESB0"),
    ("IDFC First Bank", r"\bIDFC\s*FIRST\b|\bIDFB0"),
    ("Bank of Baroda", r"BANK\s+OF\s+BARODA|\bBARB0"),
    ("Punjab National Bank", r"PUNJAB\s+NATIONAL\s+BANK|\bPUNB0|\bPNB\b"),
    ("Canara Bank", r"\bCANARA\s+BANK\b|\bCNRB0"),
    ("IndusInd Bank", r"\bINDUSIND\b|\bINDB0"),
]
_ACCOUNT_RE = _rx(r"\b(?:A/?C|ACCOUNT|ACCT)\.?\s*(?:NO|NUMBER|NUM|#)?\.?\s*[:\-]?\s*([X*\d][X*\d\-]{2,24}\d{4})(?!\d)")


def _guess_bank(text: str, filename: str) -> str | None:
    for src in (text[:20000].upper(), re.sub(r"[_\-.]+", " ", filename[:300].upper())):
        hits = []
        for name, p in _BANK_PATTERNS:
            m = re.search(p, src)
            if m:
                hits.append((m.start(), name))
        if hits:
            return min(hits)[1]
    return None


def _guess_last4(text: str, filename: str) -> str | None:
    text, filename = text[:20000], filename[:300]
    for m in _ACCOUNT_RE.finditer(text):
        digits = re.sub(r"\D", "", m.group(1))
        if len(digits) >= 4:
            return digits[-4:]
    m = re.search(r"[X*]{2,}(\d{4})(?!\d)", filename) or re.search(r"(?<!\d)\d{5,14}(\d{4})(?!\d)", filename)
    return m.group(1) if m else None


def _strict_value(text: str, last: bool = False) -> float | None:
    """The first (or last) amount written in a piece of text, signed by a Dr suffix."""
    found = list(_STRICT_AMT.finditer(text or ""))
    if not found:
        return None
    v, mk = _amount(found[-1 if last else 0].group(0))
    if v is None:
        return None
    return -abs(v) if mk == "dr" else v


def _scan_balances(c: _Collected):
    """Opening/closing balance written outside the transaction rows: "Opening Balance: 1,00,000.00" in one row, or a
    summary with the labels in one row and the figures in the next (HDFC style)."""
    for where, rows in (("pre", c.pre), ("post", c.post)):
        for ri, cells in enumerate(rows):
            nxt = rows[ri + 1] if ri + 1 < len(rows) else []
            for ci, cl in enumerate(cells):
                if not cl:
                    continue
                for kind, rx in (("open", _OPEN_RE), ("close", _CLOSE_RE), ("pre" if where == "pre" else "post", _ASON_RE)):
                    m = rx.search(cl)
                    if not m:
                        continue
                    kind = {"pre": "open", "post": "close"}.get(kind, kind)
                    if (kind == "open" and c.opening is not None) or (kind == "close" and c.closing is not None):
                        continue
                    val = None
                    for x in [cl[m.end():]] + [x for x in cells[ci + 1:] if x]:
                        val = _strict_value(x)
                        if val is not None or _OPEN_RE.search(x) or _CLOSE_RE.search(x):
                            break                      # stop at the next label: its figure isn't ours
                    if val is None and len(cells) == 1 and len(nxt) == 1:
                        val = _strict_value(nxt[0], last=(kind == "close"))   # text: labels on one line, figures below
                    elif val is None and ci < len(nxt) and _STRICT_AMT.fullmatch((nxt[ci] or "").strip()):
                        val = _strict_value(nxt[ci])
                    if val is not None and rx is _ASON_RE:
                        c.soft.add(kind)
                    if kind == "open":
                        c.set_open(val)
                    elif val is not None:
                        c.closing = val


_HINT_DR = _rx(r"\bUPI/DR\b|\bUPI-DR\b|/DR/|\bDR\b|^TO\b|\bATM\b|\bPOS\b|\bWDL\b|\bWITHDRAWAL\b|\bCHQ PAID\b|\bACH D\b|"
               r"\bNACH\b|\bEMI\b|\bDEBIT\b|\bCHRGS?\b|\bCHARGES\b|\bPURCHASE\b")
_HINT_CR = _rx(r"\bUPI/CR\b|\bUPI-CR\b|/CR/|\bCR\b|^BY\b|\bINT\.?\s?PD\b|\bINTEREST CREDIT\b|\bCASH DEP|\bREFUND\b|"
               r"\bREVERSAL\b|\bCASHBACK\b|\bCREDIT\b|\bDEPOSIT\b|\bCHQ DEP\b")


def _hint(narr: str) -> str | None:
    dr, cr = bool(_HINT_DR.search(narr)), bool(_HINT_CR.search(narr))
    if dr != cr:
        return "debit" if dr else "credit"
    return None


def _signed(t: dict) -> float:
    return t["amount"] if t["kind"] == "credit" else -t["amount"]


def _chain_score(seq: list[dict]) -> int:
    ok, prev = 0, None
    for t in seq:
        if prev is not None and prev["balance"] is not None and t["balance"] is not None:
            d = round(t["balance"] - prev["balance"], 2)
            if t["kind"] is None:
                ok += abs(abs(d) - t["amount"]) <= 0.011
            else:
                ok += abs(d - _signed(t)) <= 0.011
        prev = t
    return ok


def _finish(c: _Collected, fmt: str, filename: str, extra_warnings: list[str]) -> dict:
    if not c.txns:
        raise StatementError("no_transactions", MSG_NO_TXNS)
    _scan_balances(c)
    raws = c.txns
    stated_open, stated_close = c.opening, c.closing
    # Which way does the running balance go? Most statements are oldest first; some exports are newest first.
    asc, desc = _chain_score(raws), _chain_score(raws[::-1])
    newest_first = desc > asc or (desc == asc and raws[0]["date"] > raws[-1]["date"])
    chron = raws[::-1] if newest_first else raws
    # A "Balance as on <date>" line may be the balance on the day of download, not at the start/end of the statement:
    # keep it only if it agrees with the first/last row.
    f, l_ = chron[0], chron[-1]
    if "open" in c.soft and f["kind"] and f["balance"] is not None and abs(stated_open + _signed(f) - f["balance"]) > 0.011:
        stated_open = None
    if "close" in c.soft and l_["balance"] is not None and abs(stated_close - l_["balance"]) > 0.011:
        stated_close = None
    prev = stated_open
    for t in chron:
        if t["kind"] is None and t["balance"] is not None and prev is not None:
            d = round(t["balance"] - prev, 2)
            if abs(d + t["amount"]) <= 0.011:
                t["kind"] = "debit"
            elif abs(d - t["amount"]) <= 0.011:
                t["kind"] = "credit"
        if t["balance"] is not None:
            prev = t["balance"]
        elif prev is not None and t["kind"]:
            prev = round(prev + _signed(t), 2)
        else:
            prev = None
    by_words = defaulted = 0
    for t in raws:
        t["guessed"] = False
        if t["kind"] is None:
            if c.signed:
                t["kind"] = "debit" if t["sign"] < 0 else "credit"
                continue
            t["guessed"] = True                    # from the description, or not at all: the owner must check it
            t["kind"] = _hint(t["narr"])
            if t["kind"] is None:
                t["kind"] = "debit"
                defaulted += 1
            else:
                by_words += 1
    breaks, prev = 0, stated_open
    for t in chron:
        if t["balance"] is not None and prev is not None and abs(prev + _signed(t) - t["balance"]) > 0.011:
            breaks += 1
        prev = t["balance"] if t["balance"] is not None else (prev + _signed(t) if prev is not None else None)
    first, lastt = chron[0], chron[-1]
    opening = stated_open
    if opening is None and first["balance"] is not None:
        opening = round(first["balance"] - _signed(first), 2)
    closing = stated_close if stated_close is not None else lastt["balance"]

    warnings = list(extra_warnings)
    if defaulted:
        warnings.append(f"{defaulted} of {_plural(len(raws), 'transaction', 'transactions')}: the statement doesn't say "
                        f"whether {'it was' if defaulted == 1 else 'they were'} money in or money out, so we marked "
                        f"{'it' if defaulted == 1 else 'them'} as money out. {'It is' if defaulted == 1 else 'They are'} "
                        f"flagged — please check before adding to Expenses.")
    if by_words:
        warnings.append(f"{by_words} of {_plural(len(raws), 'transaction', 'transactions')}: we went by the description "
                        f"to tell money in from money out. {'It is' if by_words == 1 else 'They are'} flagged — please "
                        f"check before adding to Expenses.")
    if breaks:
        warnings.append(f"On {_plural(breaks, 'row', 'rows')} the running balance doesn't match the amount. A row may be "
                        f"missing or misread — please compare with your statement.")
    if stated_open is not None and stated_close is not None:
        diff = round(stated_open + sum(_signed(t) for t in raws) - stated_close, 2)
        if abs(diff) > 0.011:
            warnings.append(f"The transactions we read don't add up to the closing balance in the statement (off by "
                            f"{_inr(abs(diff))}). Some rows may be missing — please check.")
    if c.both_sides:
        warnings.append(f"{_plural(c.both_sides, 'row had', 'rows had')} both a withdrawal and a deposit; we added each "
                        f"as a separate transaction.")
    if c.skipped:
        warnings.append(f"{_plural(c.skipped, 'row had', 'rows had')} a date but no amount, so we left "
                        f"{'it' if c.skipped == 1 else 'them'} out.")

    pre_text = "\n".join(" ".join(x for x in r if x) for r in c.pre)
    seen = Counter()
    out = []
    for t in raws:
        narr = mask(re.sub(r"\s+", " ", t["narr"][:CELL_CAP]).strip())[:200].strip()
        paise = int(round(t["amount"] * 100))
        key = f"{t['date'].isoformat()}|{t['kind']}|{paise}|{re.sub(r'[^A-Z0-9]', '', narr.upper())}"
        n = seen[key]
        seen[key] += 1
        out.append({"date": t["date"].isoformat(), "narration": narr, "payee": payee(narr),
                    "amount": round(t["amount"], 2), "kind": t["kind"],
                    "balance": round(t["balance"], 2) if t["balance"] is not None else None,
                    "ref": hashlib.sha1(f"{key}|{n}".encode()).hexdigest()[:20], "guessed": t["guessed"]})
    dates = [t["date"] for t in raws]
    return {"format": fmt, "bank": _guess_bank(pre_text, filename), "account_last4": _guess_last4(pre_text, filename),
            "from": min(dates).isoformat(), "to": max(dates).isoformat(),
            "opening_balance": round(opening, 2) if opening is not None else None,
            "closing_balance": round(closing, 2) if closing is not None else None,
            "transactions": out, "warnings": warnings}


def _best(cands: list[_Collected]) -> _Collected:
    return max(cands, key=lambda c: len(c.txns)) if cands else _Collected()


def _parse(data: bytes, filename: str, password: str | None) -> dict:
    if not data:
        raise StatementError("unreadable", "This file is empty. Download the statement again from net banking and try again.")
    if len(data) > MAX_FILE_BYTES:
        raise StatementError("too_big", "This file is bigger than 25 MB. Download fewer months at a time and upload again.")
    kind = _sniff(data)
    if kind == "pdf":
        c, by_lines, _ = _read_pdf(data, password)
        notes = ["This PDF has no table lines, so we read it line by line. Please check a few rows — the Excel or CSV "
                 "download from net banking is more reliable."] if by_lines and c.txns else []
        return _finish(c, "pdf", filename, notes)
    if kind == "xlsx":
        return _finish(_best([_read_tables([s]) for s in _read_xlsx(data)]), "xlsx", filename, [])
    if kind == "xls":
        return _finish(_best([_read_tables([s]) for s in _read_xls(data)]), "xls", filename, [])
    if kind == "image":
        raise StatementError("unreadable", "This looks like a photo or image of the statement, which we can't read. "
                                           "Download the statement as Excel or CSV from net banking and upload that instead.")
    if kind == "binary":
        raise _unsupported()
    text = _decode(data)
    head = text[:65536]
    excel_xml = r"urn:schemas-microsoft-com:office:spreadsheet|<\s*(?:ss:)?Workbook\b"
    if re.search(r"<\?xml", head[:512]) and re.search(excel_xml, head, re.I):
        return _finish(_best([_read_tables([s]) for s in _read_spreadsheet_xml(text)]), "xls", filename, [])
    if re.search(r"<\s*(?:table|html)\b", head, re.I):
        return _finish(_read_tables(_read_html(text)), "html", filename, [])
    grids, lines = _read_csv(text)
    c = _read_tables(grids)
    if not c.txns:
        c = _read_lines(lines)
    return _finish(c, "csv", filename, [])


def _out_of_memory(e: BaseException) -> bool:
    """MemoryError, also when a library wraps it (pdfplumber raises PdfminerException(MemoryError(...)))."""
    seen = 0
    while e is not None and seen < 10:
        if isinstance(e, MemoryError) or any(isinstance(a, MemoryError) for a in getattr(e, "args", ())):
            return True
        e, seen = e.__cause__ or e.__context__, seen + 1
    return False


def parse(data: bytes, filename: str, password: str | None = None) -> dict:
    """Read a bank statement. Returns {"format", "bank", "account_last4", "from", "to", "opening_balance",
    "closing_balance", "transactions": [{"date", "narration" (masked), "payee", "amount", "kind", "balance", "ref",
    "guessed" (True when money in/out came from the description or a default, not the statement)}], "warnings"}.
    Raises StatementError (codes: password_needed, wrong_password, unsupported, unreadable, no_transactions, too_big)
    with a message to show the owner as is. Runs in this process: from the web server use parse_isolated()."""
    try:
        return _parse(bytes(data or b""), filename or "", password)
    except StatementError:
        raise
    except Exception as e:  # noqa: BLE001  a damaged file is the owner's to replace; never surface library internals or row text
        if _out_of_memory(e):
            raise StatementError("unreadable", MSG_TOO_COMPLEX) from None
        raise StatementError("unreadable", "We couldn't read this file. Download the statement again from net banking "
                                           "(Excel or CSV works best) and try again.") from None


# ----------------------------------------------------------------------------------------------------------------------
# Running parse() in a separate process (what the web server should call)

ISOLATED_AT_ONCE = 2                     # isolated parses running at once in one server process
CHILD_MEMORY = 768 * 1024 * 1024         # address-space limit of the child process
_SLOTS = threading.BoundedSemaphore(ISOLATED_AT_ONCE)
MSG_BUSY = "Several statements are being read right now. Please try again in a minute."
MSG_CANT_START = "We couldn't read this file right now. Please try again in a minute."
_RESULT_KEYS = {"format", "bank", "account_last4", "from", "to", "opening_balance", "closing_balance", "transactions",
                "warnings"}


def _frame(*fields: bytes) -> bytes:
    """Length-prefixed fields (8-byte big-endian length + bytes): the child's whole input, sent over stdin."""
    return b"".join(struct.pack(">Q", len(f)) + f for f in fields)


def _unframe(raw: bytes, n: int) -> list[bytes]:
    out, pos = [], 0
    for _ in range(n):
        (size,) = struct.unpack_from(">Q", raw, pos)
        pos += 8
        out.append(raw[pos:pos + size])
        pos += size
    return out


def parse_isolated(data: bytes, filename: str, password: str | None = None, timeout: float = 60) -> dict:
    """parse() in a separate Python process with a memory limit (768 MB), a CPU limit (about `timeout` seconds) and a
    wall-clock timeout after which it is killed. At most 2 run at once per server process; a third call waits up to
    `timeout` seconds for a slot. Blocking — from async code call it with `await asyncio.to_thread(parse_isolated, ...)`.

    The file and the password go to the child over stdin, never in its command line or environment (so they can't show
    up in `ps`); the child gets an empty environment (no server secrets) and writes no files or core dumps.
    Returns exactly what parse() returns, or raises StatementError with parse()'s code and message. A child that runs
    out of time or memory, crashes or answers nonsense -> StatementError("unreadable", MSG_TOO_COMPLEX)."""
    data = bytes(data or b"")
    if not data:
        raise StatementError("unreadable", "This file is empty. Download the statement again from net banking and try again.")
    if len(data) > MAX_FILE_BYTES:
        raise StatementError("too_big", "This file is bigger than 25 MB. Download fewer months at a time and upload again.")
    timeout = max(1.0, float(timeout))
    if not _SLOTS.acquire(timeout=timeout):
        raise StatementError("unreadable", MSG_BUSY)
    try:
        return _run_child(data, filename or "", password, timeout)
    finally:
        _SLOTS.release()


def _child_argv(timeout: float) -> list[str]:
    # -I: ignore PYTHON* variables and don't put the script's folder or the working folder on sys.path; -B: write no .pyc.
    # The module is run by its file path (not -m app.bankstatement) because -I keeps the backend folder off sys.path.
    # Only the limits travel in argv — never the file, the file name or the password.
    return [sys.executable, "-I", "-B", os.path.abspath(__file__), "--child", f"{timeout:g}", str(int(CHILD_MEMORY))]


def _run_child(data: bytes, filename: str, password: str | None, timeout: float) -> dict:
    paths = "\0".join(p for p in sys.path if p and os.path.isabs(p))     # so the child finds the same libraries
    frame = _frame(b"1" if password is not None else b"0", (password or "").encode("utf-8", "surrogatepass"),
                   filename[:300].encode("utf-8", "replace"), paths.encode("utf-8", "surrogateescape"), data)
    try:
        proc = subprocess.Popen(_child_argv(timeout), stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                stderr=subprocess.DEVNULL, env={"PATH": os.environ.get("PATH", os.defpath)},
                                cwd=os.path.dirname(os.path.abspath(__file__)), close_fds=True, start_new_session=True)
    except OSError:
        raise StatementError("unreadable", MSG_CANT_START) from None
    try:
        out, _ = proc.communicate(frame, timeout=timeout)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.communicate()
        raise StatementError("unreadable", MSG_TOO_COMPLEX) from None
    finally:
        frame = None
    if proc.returncode != 0:                 # killed for CPU or memory, or crashed
        raise StatementError("unreadable", MSG_TOO_COMPLEX)
    try:
        msg = json.loads(out)
    except ValueError:
        raise StatementError("unreadable", MSG_TOO_COMPLEX) from None
    return _checked(msg)


def _checked(msg) -> dict:
    """The child's answer is treated as untrusted: only a well-formed result or a known error code gets through."""
    if isinstance(msg, dict) and msg.get("ok") is True:
        r = msg.get("result")
        if (isinstance(r, dict) and set(r) == _RESULT_KEYS and isinstance(r["transactions"], list)
                and len(r["transactions"]) <= MAX_TXNS and isinstance(r["warnings"], list)):
            return r
    elif isinstance(msg, dict):
        code, message = msg.get("code"), msg.get("message")
        if code in ERROR_CODES and isinstance(message, str) and 0 < len(message) <= 500:
            raise StatementError(code, message)
    raise StatementError("unreadable", MSG_TOO_COMPLEX)


def _limit_child(timeout: float, memory: int):
    try:
        import resource
    except ImportError:                      # not on Linux/macOS: the wall-clock timeout still applies
        return
    cpu = int(timeout) + 1
    for lim, soft, hard in ((resource.RLIMIT_AS, memory, memory), (resource.RLIMIT_CPU, cpu, cpu + 1),
                            (resource.RLIMIT_CORE, 0, 0), (resource.RLIMIT_FSIZE, 0, 0)):
        try:
            _, cur_hard = resource.getrlimit(lim)
            if cur_hard != resource.RLIM_INFINITY:
                soft, hard = min(soft, cur_hard), min(hard, cur_hard)
            resource.setrlimit(lim, (soft, hard))
        except (ValueError, OSError):
            pass


def _preload(kind: str):
    try:
        if kind == "pdf":
            __import__("pdfplumber.utils.exceptions")
            __import__("pdfminer.pdfdocument")
        elif kind == "xlsx":
            __import__("openpyxl")
        elif kind == "xls":
            __import__("xlrd")
    except ImportError:
        pass                                 # parse() reports it as an unreadable file


def _child_main(args: list[str]):
    """Entry point of the child process: limits first, then read the framed input from stdin, parse, answer on stdout
    as one JSON object. Nothing is logged; stderr goes nowhere."""
    try:
        timeout = float(args[0]) if args else 60.0
        memory = int(args[1]) if len(args) > 1 else CHILD_MEMORY
    except ValueError:
        timeout, memory = 60.0, CHILD_MEMORY
    try:
        flag, pw, name, paths, data = _unframe(sys.stdin.buffer.read(MAX_FILE_BYTES + 1_000_000), 5)
        for path in paths.decode("utf-8", "surrogateescape").split("\0"):
            if path and os.path.isabs(path) and path not in sys.path:
                sys.path.append(path)
        _preload(_sniff(data))               # the reading library is trusted code: load it before the limits go on,
        _limit_child(timeout, memory)        # so the limits are spent on the file, not on importing
        password = pw.decode("utf-8", "surrogatepass") if flag == b"1" else None
        del pw
        out = {"ok": True, "result": parse(data, name.decode("utf-8", "replace"), password)}
    except StatementError as e:
        out = {"ok": False, "code": e.code, "message": e.message}
    except BaseException:  # noqa: BLE001  out of memory or anything else: one plain answer, no details
        out = {"ok": False, "code": "unreadable", "message": MSG_TOO_COMPLEX}
    password = None  # noqa: F841
    sys.stdout.buffer.write(json.dumps(out).encode("utf-8"))
    sys.stdout.buffer.flush()


if __name__ == "__main__" and sys.argv[1:2] == ["--child"]:
    _child_main(sys.argv[2:])
