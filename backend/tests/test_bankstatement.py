"""Bank statement reader (app/bankstatement.py). Pure tests: every statement here is invented and built in the test
(one real BIFF .xls lives in tests/fixtures/bank). No database, network or AI."""
import io
import os
import struct
import zipfile
from datetime import date, datetime

import openpyxl
import pytest

from app import bankstatement as bs
from app.bankstatement import StatementError

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures", "bank")


def by_date(result, d):
    return [t for t in result["transactions"] if t["date"] == d]


def kinds(result):
    return [t["kind"] for t in result["transactions"]]


def err(data, name, password=None) -> StatementError:
    with pytest.raises(StatementError) as e:
        bs.parse(data, name, password)
    return e.value


# ---------------------------------------------------------------------------------------------------------------------
# Builders (invented names, numbers and handles only)

def hdfc_html_xls() -> bytes:
    """An HDFC-style ".xls" download that is really an HTML table, with a wrapped narration and a summary block."""
    rows = [
        ("********", "********", "", "", "", "", ""),
        ("01/04/26", "UPI-SHARMA STEEL TRADERS-zz.vendor01@okhdfcbank-HDFC0ZZ0123-", "0000412345678901", "01/04/26", "12,500.00", "", "1,37,500.00"),
        ("", "412345678901-PAYMENT FROM PHONE", "", "", "", "", ""),
        ("02/04/26", "NEFT CR-ICIC0ZZ0456-PATEL AGRO EXPORTS-SHARMA STEEL TRADERS-ICICN52026040212345", "ICICN52026040212345", "02/04/26", "", "50,000.00", "1,87,500.00"),
        ("05/04/26", "POS 416021XXXXXX1234 SWIGGY BANGALORE", "0000000000000001", "05/04/26", "850.00", "", "1,86,650.00"),
        ("10/04/26", "ATM WDL-ATM CASH 4567 ANDHERI EAST", "0000000000004567", "10/04/26", "10,000.00", "", "1,76,650.00"),
        ("30/04/26", "CREDIT INTEREST CAPITALISED", "0000000000000000", "30/04/26", "", "312.00", "1,76,962.00"),
    ]
    tr = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>\n" for r in rows)
    return f"""<html><head><meta http-equiv="Content-Type" content="text/html; charset=utf-8"></head><body>
<table><tr><td>HDFC BANK Ltd.</td><td></td><td>Page No .: 1</td></tr>
<tr><td>M/S. SHARMA STEEL TRADERS</td></tr>
<tr><td>Account No :000123456789&nbsp;&nbsp;OTHER</td></tr>
<tr><td>Statement From : 01/04/26 To : 30/04/26</td></tr></table>
<table border="1"><tr><th>Date</th><th>Narration</th><th>Chq./Ref.No.</th><th>Value Dt</th><th>Withdrawal Amt.</th>
<th>Deposit Amt.</th><th>Closing Balance</th></tr>
{tr}</table>
<table><tr><td colspan="6">STATEMENT SUMMARY :-</td></tr>
<tr><td>Opening Balance</td><td>Dr Count</td><td>Cr Count</td><td>Debits</td><td>Credits</td><td>Closing Bal</td></tr>
<tr><td>1,50,000.00</td><td>3</td><td>2</td><td>23,350.00</td><td>50,312.00</td><td>1,76,962.00</td></tr></table>
<p>Generated On: 21-Oct-2026. This is a computer generated statement.</p></body></html>""".encode()


def excel_serial(d: date) -> int:
    return (d - date(1899, 12, 30)).days


def icici_xlsx() -> bytes:
    """ICICI-style XLSX: preamble rows, Value Date before Transaction Date, one Amount column + a Dr/Cr column."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, "DETAILED STATEMENT"])
    ws.append([None, "Transactions List - SHARMA STEEL TRADERS (INR) - 000123456789"])
    ws.append([None, "Account Number", "000123456789 ( INR ) - SHARMA STEEL TRADERS", None, "IFSC", "ICIC0ZZ0001"])
    ws.append([None, "Transaction Date from", "01/04/2026", "to", "30/04/2026"])
    ws.append([])
    ws.append([None, "S No.", "Value Date", "Transaction Date", "Cheque Number", "Transaction Remarks", "Amount (INR)",
               "Dr/Cr", "Balance (INR)"])
    ws.append([None, 1, "02/04/2026", datetime(2026, 4, 1), "-", "UPI/412345678901/SHARMA STEEL/zz.vendor03@okaxis/AXIS BANK/AXI1234",
               12500.0, "DR", 137500.0])
    ws.append([None, 2, "04/04/2026", excel_serial(date(2026, 4, 3)), "-", "NEFT-ICICN52026040312345-PATEL AGRO EXPORTS-INV 991",
               "50,000.00", "CR", "1,87,500.00"])
    ws.append([None, 3, "06/04/2026", "06/04/2026", "-", "POS/AMAZON WEB SERVICES/SEATTLE/416021XXXXXX1234", 3450.0, "Dr", 184050.0])
    ws.append([None, 4, "09/04/2026", "09-Apr-2026", "-", "BIL/ONL/000123456789/MSEDCL/LT CONSUMER", 8760.5, "DR", 175289.5])
    ws.append([None, 5, "30/04/2026", "30 Apr 2026", "-", "INT.PD:01-04-2026 to 30-04-2026", 210.0, "CR", 175499.5])
    ws.append([])
    ws.append([None, "Legends Used in Account Statement"])
    ws.append([None, "INF - Internet Fund Transfer", "BIL - Bill payment"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SBI_CSV = """Account Name,:,SHREE GANESH FABRICATORS
Address,:,"PLOT 12, MIDC, THANE"
Date,:,21 Oct 2026
Account Number,:,000123456789
Account Description,:,CURRENT ACCOUNT
Branch,:,THANE INDUSTRIAL AREA
IFS Code,:,SBIN0ZZ0003
Balance as on 1 Apr 2026,:,"2,00,000.00"
Start Date,:,1 Apr 2026
End Date,:,30 Apr 2026

Txn Date,Value Date,Description,Ref No./Cheque No.,Debit,Credit,Balance
1 Apr 2026,1 Apr 2026,TO TRANSFER-UPI/DR/412345678901/PATEL HARDWARE/,TRANSFER TO 4897000000001,"12,500.00",,"1,87,500.00"
,,HDFC/zz.vendor04@okhdfcbank/Payment,,,,
3 Apr 2026,3 Apr 2026,BY TRANSFER-NEFT*HDFC0ZZ0001*N123456789012*KULKARNI AUTO PARTS*,TRANSFER FROM 4897000000002,,"75,000.00","2,62,500.00"
5 Apr 2026,5 Apr 2026,TO TRANSFER-INB NEFT GUPTA ESTATES OFFICE RENT APR,TRANSFER TO 4897000000003,"45,000.00",,"2,17,500.00"
12 Apr 2026,12 Apr 2026,TO TRANSFER-EPFO CHALLAN TRRN 1234567890123,TRANSFER TO 4897000000004,"14,400.00",,"2,03,100.00"
,,FOR MAR 2026,,,,
20 Apr 2026,20 Apr 2026,BY CLG-CHQ DEP-PATIL INTERIORS,000456,,"1,20,000.00","3,23,100.00"
,,,,,,
**This is a computer generated statement and does not require a signature.
"""


def axis_table_pdf(extra_rows=0) -> bytes:
    """An Axis-style PDF with a ruled table (wrapped particulars inside the cells), opening/total/closing rows."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()["BodyText"]
    st.fontSize, st.leading = 7, 8
    P = lambda s: Paragraph(s, st)  # noqa: E731
    data = [["Tran Date", "Chq No", "Particulars", "Debit", "Credit", "Balance", "Init. Br"],
            ["01-04-2026", "", "OPENING BALANCE", "", "", "1,50,000.00", ""],
            ["02-04-2026", "", P("UPI/P2A/412345678901/KULKARNI AUTO PARTS/HDFC BANK/INV 2231 PAYMENT FOR APRIL"), "",
             "60,000.00", "2,10,000.00", "1234"],
            ["04-04-2026", "", P("UPI/P2M/412345678902/AMAZON PAY INDIA/Payment for order"), "2,399.00", "", "2,07,601.00", "1234"],
            ["06-04-2026", "000123", P("CLG/SHARMA STEEL TRADERS/CTS INWARD"), "40,000.00", "", "1,67,601.00", "1234"],
            ["07-04-2026", "", "GST ON CHG", "18.00", "", "1,67,583.00", "1234"]]
    bal = 167583.00
    for i in range(extra_rows):
        bal -= 100
        data.append([f"{8 + i % 20:02d}-04-2026", "", f"UPI/P2M/41234567{i:04d}/RAMU TEA STALL/Pay", "100.00", "",
                     f"{bal:,.2f}", "1234"])
    data += [["", "", "TRANSACTION TOTAL", f"{42417 + 100 * extra_rows:,.2f}", "60,000.00", "", ""],
             ["", "", "CLOSING BALANCE", "", "", f"{bal:,.2f}", ""]]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4))
    t = Table(data, colWidths=[60, 45, 260, 60, 60, 70, 40], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 7),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    doc.build([Paragraph("Statement of Account No :000123456789 for the period (From : 01-04-2026 To : 30-04-2026)", st),
               Paragraph("IFSC Code : UTIB0ZZ0001", st), Spacer(1, 12), t])
    return buf.getvalue()


KOTAK_PAGES = [
    ["Kotak Mahindra Bank", "Account Statement", "Account No. XXXXXXXX4321 CRN XXXX5678",
     "Period 01-Apr-2026 - 30-Apr-2026",
     "Date Narration Chq/Ref No Withdrawal (Dr) Deposit (Cr) Balance",
     "Opening Balance 1,50,000.00",
     "01-Apr-2026 UPI/SHARMA STEEL TRADERS/412345678901/Payment UPI-412345678901 12,500.00 1,37,500.00",
     "03-Apr-2026 NEFT INWARD-PATEL AGRO EXPORTS-ICICN52026040312345 NEFTINW-0012 40,000.00 1,77,500.00",
     "05-Apr-2026 NEFT DR-SBIN0ZZ0002-GUPTA ESTATES-OFFICE RENT APR NEFT-0045 45,000.00 1,32,500.00",
     "Page 1 of 2"],
    ["Kotak Mahindra Bank", "Account Statement",
     "Date Narration Chq/Ref No Withdrawal (Dr) Deposit (Cr) Balance",
     "08-Apr-2026 IMPS/P2A/412345678904/RAHUL TRANSPORT CO 12345 8,000.00 1,24,500.00",
     "/HDFC/LR NO 5521",
     "10-Apr-2026 ACH D- SHUBH FINCORP LTD-000987654321 18,500.00 1,06,000.00",
     "15-Apr-2026 CASH DEPOSIT BY SELF 25,000.00 1,31,000.00",
     "Closing Balance 1,31,000.00",
     "Page 2 of 2"],
]


def text_pdf(pages: list[list[str]]) -> bytes:
    """Text drawn line by line, no table lines (pdfplumber finds no tables)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for lines in pages:
        y = 800
        for ln in lines:
            c.setFont("Helvetica", 8)
            c.drawString(30, y, ln)
            y -= 14
        c.showPage()
    c.save()
    return buf.getvalue()


def encrypt(pdf: bytes, password: str) -> bytes:
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    w.append(PdfReader(io.BytesIO(pdf)))
    w.encrypt(user_password=password, owner_password="owner-" + password, algorithm="AES-256")
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def scanned_pdf() -> bytes:
    """A 'scan': one image per page and no text layer."""
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    img = Image.new("RGB", (600, 800), "white")
    d = ImageDraw.Draw(img)
    for y in range(40, 760, 24):
        d.rectangle([40, y, 560, y + 10], fill="black")
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawImage(ImageReader(img), 20, 20, width=550, height=800)
    c.showPage()
    c.save()
    return buf.getvalue()


def simple_csv(rows: list[tuple], header="Date,Narration,Withdrawal,Deposit,Balance") -> bytes:
    lines = [header] + [",".join(f'"{x}"' if "," in str(x) else str(x) for x in r) for r in rows]
    return ("\n".join(lines) + "\n").encode()


# ---------------------------------------------------------------------------------------------------------------------
# Formats

def test_hdfc_xls_that_is_really_html():
    r = bs.parse(hdfc_html_xls(), "Acct_Statement_XXXXXXXX6789_21102026.xls")
    assert r["format"] == "html"
    assert r["bank"] == "HDFC Bank"
    assert r["account_last4"] == "6789"
    assert (r["from"], r["to"]) == ("2026-04-01", "2026-04-30")
    assert r["opening_balance"] == 150000.0 and r["closing_balance"] == 176962.0
    assert len(r["transactions"]) == 5
    assert kinds(r) == ["debit", "credit", "debit", "debit", "credit"]
    assert [t["amount"] for t in r["transactions"]] == [12500.0, 50000.0, 850.0, 10000.0, 312.0]
    assert [t["balance"] for t in r["transactions"]] == [137500.0, 187500.0, 186650.0, 176650.0, 176962.0]
    first = r["transactions"][0]
    # the wrapped second row is joined, then masked: handle keeps its first letter, the 12-digit ref keeps its last 4
    assert first["narration"].endswith("XXXXXXXX8901-PAYMENT FROM PHONE")
    assert "z***@okhdfcbank" in first["narration"] and "zz.vendor01" not in first["narration"]
    assert "412345678901" not in first["narration"]
    assert first["payee"] == "Sharma Steel Traders"
    assert r["warnings"] == []


def test_icici_xlsx_amount_and_drcr_column_with_preamble():
    r = bs.parse(icici_xlsx(), "OpTransactionHistory21-10-2026.xlsx")
    assert r["format"] == "xlsx"
    assert r["bank"] == "ICICI Bank"
    assert r["account_last4"] == "6789"
    # Transaction Date wins over Value Date (02/04 vs 01/04), Excel serial and datetime cells both read
    assert [t["date"] for t in r["transactions"]] == ["2026-04-01", "2026-04-03", "2026-04-06", "2026-04-09", "2026-04-30"]
    assert kinds(r) == ["debit", "credit", "debit", "debit", "credit"]
    assert [t["amount"] for t in r["transactions"]] == [12500.0, 50000.0, 3450.0, 8760.5, 210.0]
    assert r["opening_balance"] == 150000.0 and r["closing_balance"] == 175499.5
    assert r["transactions"][3]["payee"] == "MSEDCL"
    assert r["warnings"] == []


def test_sbi_csv_debit_credit_columns_wrapped_narrations_and_footer():
    r = bs.parse(SBI_CSV.encode(), "statement.csv")
    assert r["format"] == "csv"
    assert r["bank"] == "State Bank of India"
    assert r["account_last4"] == "6789"
    assert r["opening_balance"] == 200000.0 and r["closing_balance"] == 323100.0
    assert len(r["transactions"]) == 5
    assert kinds(r) == ["debit", "credit", "debit", "debit", "credit"]
    t0 = r["transactions"][0]
    assert t0["narration"] == "TO TRANSFER-UPI/DR/XXXXXXXX8901/PATEL HARDWARE/HDFC/z***@okhdfcbank/Payment"
    assert t0["payee"] == "Patel Hardware"
    assert r["transactions"][3]["narration"].endswith("FOR MAR 2026")      # wrapped row joined
    assert r["transactions"][1]["payee"] == "Kulkarni Auto Parts"
    assert not any("computer generated" in t["narration"].lower() for t in r["transactions"])
    assert r["warnings"] == []


@pytest.mark.parametrize("delim,name", [(";", "stmt.csv"), ("\t", "stmt.xls"), ("|", "stmt.txt")])
def test_other_delimiters_and_text_named_xls(delim, name):
    rows = ["Bank statement for A/C 000123456789", "",
            delim.join(["Date", "Particulars", "Debit", "Credit", "Balance"]),
            delim.join(["01-04-2026", "UPI/DR/412345678901/GUPTA ESTATES/HDFC/RENT", "45000.00", "", "55000.00"]),
            delim.join(["02-04-2026", "NEFT CR-KULKARNI AUTO PARTS", "", "20000.00", "75000.00"])]
    r = bs.parse("\n".join(rows).encode(), name)
    assert r["format"] == "csv"            # decided by content: a tab-separated ".xls" is text
    assert kinds(r) == ["debit", "credit"]
    assert r["account_last4"] == "6789"
    assert r["opening_balance"] == 100000.0


def test_axis_like_pdf_with_table():
    r = bs.parse(axis_table_pdf(), "Axis_statement.pdf")
    assert r["format"] == "pdf"
    assert r["bank"] == "Axis Bank"
    assert r["account_last4"] == "6789"
    assert r["opening_balance"] == 150000.0 and r["closing_balance"] == 167583.0
    assert len(r["transactions"]) == 4                     # opening, total and closing rows are not transactions
    assert kinds(r) == ["credit", "debit", "debit", "debit"]
    assert [t["amount"] for t in r["transactions"]] == [60000.0, 2399.0, 40000.0, 18.0]
    t0 = r["transactions"][0]
    assert "INV 2231 PAYMENT FOR APRIL" in t0["narration"]   # text wrapped inside the cell is kept
    assert t0["payee"] == "Kulkarni Auto Parts"
    assert r["warnings"] == []


def test_kotak_like_text_pdf_infers_debit_credit_from_balance():
    r = bs.parse(text_pdf(KOTAK_PAGES), "kotak.pdf")
    assert r["format"] == "pdf"
    assert r["bank"] == "Kotak Mahindra Bank"
    assert r["account_last4"] == "4321"
    assert r["opening_balance"] == 150000.0 and r["closing_balance"] == 131000.0
    assert [t["date"] for t in r["transactions"]] == ["2026-04-01", "2026-04-03", "2026-04-05", "2026-04-08",
                                                       "2026-04-10", "2026-04-15"]
    assert kinds(r) == ["debit", "credit", "debit", "debit", "debit", "credit"]
    assert [t["amount"] for t in r["transactions"]] == [12500.0, 40000.0, 45000.0, 8000.0, 18500.0, 25000.0]
    imps = r["transactions"][3]
    assert imps["narration"].endswith("/HDFC/LR NO 5521")    # continuation line joined
    assert imps["payee"] == "Rahul Transport Co"
    assert not any("Page" in t["narration"] or "Kotak" in t["narration"] for t in r["transactions"])
    assert len(r["warnings"]) == 1 and "line by line" in r["warnings"][0]


def test_text_pdf_page_header_is_not_glued_to_last_row_without_footer():
    pages = [["Date Narration Withdrawal Deposit Balance",
              "01-04-2026 UPI/DR/412345678901/SHARMA STEEL TRADERS/HDFC 1,000.00 9,000.00"],
             ["SOME BANK HEADER LINE", "Date Narration Withdrawal Deposit Balance",
              "02-04-2026 UPI/DR/412345678902/GUPTA ESTATES/HDFC 1,000.00 8,000.00"]]
    r = bs.parse(text_pdf(pages), "x.pdf")
    assert [t["narration"] for t in r["transactions"]][0].endswith("/HDFC")


def test_encrypted_pdf_password_flow():
    locked = encrypt(axis_table_pdf(), "SHAR0104")
    e = err(locked, "stmt.pdf")
    assert e.code == "password_needed" and "password" in e.message.lower()
    e = err(locked, "stmt.pdf", password="WRONG999")
    assert e.code == "wrong_password"
    assert "WRONG999" not in e.message and "WRONG999" not in str(e)
    r = bs.parse(locked, "stmt.pdf", password="SHAR0104")
    assert len(r["transactions"]) == 4 and r["opening_balance"] == 150000.0
    assert "SHAR0104" not in repr(r)


def test_scanned_pdf_is_unreadable_with_a_useful_message():
    e = err(scanned_pdf(), "scan.pdf")
    assert e.code == "unreadable"
    assert "Excel or CSV" in e.message and "net banking" in e.message


def test_real_biff_xls_fixture():
    with open(os.path.join(FIXTURES, "canara_like.xls"), "rb") as f:
        data = f.read()
    assert data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"      # a real BIFF/OLE2 workbook
    pytest.importorskip("xlrd")
    r = bs.parse(data, "statement.xls")
    assert r["format"] == "xls"
    assert r["bank"] == "Canara Bank"
    assert r["account_last4"] == "6789"
    assert [t["date"] for t in r["transactions"]] == ["2026-04-02", "2026-04-05", "2026-04-10", "2026-04-30", "2026-05-05"]
    assert kinds(r) == ["credit", "debit", "debit", "credit", "debit"]
    assert r["transactions"][3]["amount"] == 412.5
    assert r["opening_balance"] == 200000.0 and r["closing_balance"] == 176912.5
    assert r["transactions"][2]["narration"] == "ACH D- SHUBH FINCORP LTD-XXXXXXXX4321"


def test_excel_2003_xml_saved_as_xls():
    xml = """<?xml version="1.0"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
<Worksheet ss:Name="Sheet1"><Table>
<Row><Cell><Data ss:Type="String">Account No: 000123456789</Data></Cell></Row>
<Row><Cell><Data ss:Type="String">Txn Date</Data></Cell><Cell><Data ss:Type="String">Description</Data></Cell>
<Cell><Data ss:Type="String">Debit</Data></Cell><Cell><Data ss:Type="String">Credit</Data></Cell>
<Cell><Data ss:Type="String">Balance</Data></Cell></Row>
<Row><Cell><Data ss:Type="DateTime">2026-04-01T00:00:00.000</Data></Cell><Cell><Data ss:Type="String">NEFT DR-GUPTA ESTATES-RENT</Data></Cell>
<Cell><Data ss:Type="Number">45000</Data></Cell><Cell ss:Index="5"><Data ss:Type="Number">55000</Data></Cell></Row>
<Row><Cell><Data ss:Type="String">02/04/2026</Data></Cell><Cell><Data ss:Type="String">UPI/CR/412345678901/PATIL INTERIORS</Data></Cell>
<Cell ss:Index="4"><Data ss:Type="Number">5000.5</Data></Cell><Cell><Data ss:Type="Number">60000.5</Data></Cell></Row>
</Table></Worksheet></Workbook>"""
    r = bs.parse(xml.encode(), "statement.xls")
    assert r["format"] == "xls"
    assert [(t["date"], t["kind"], t["amount"], t["balance"]) for t in r["transactions"]] == [
        ("2026-04-01", "debit", 45000.0, 55000.0), ("2026-04-02", "credit", 5000.5, 60000.5)]


HEADERS = [  # (bank style, header row, what each column holds) — roles are written out here, not taken from the module
    ("HDFC", ["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"],
     ["date", "narr", "ref", "vdate", "debit", "credit", "bal"]),
    ("ICICI", ["S No.", "Value Date", "Transaction Date", "Cheque Number", "Transaction Remarks", "Withdrawal Amount (INR )",
               "Deposit Amount (INR )", "Balance (INR )"], ["sno", "vdate", "date", "ref", "narr", "debit", "credit", "bal"]),
    ("ICICI corporate", ["Transaction ID", "Value Date", "Txn Posted Date", "ChequeNo.", "Description", "Cr/Dr",
                         "Transaction Amount(INR)", "Available Balance(INR)"],
     ["sno", "vdate", "date", "ref", "narr", "drcr", "amount", "bal"]),
    ("SBI", ["Txn Date", "Value Date", "Description", "Ref No./Cheque No.", "Debit", "Credit", "Balance"],
     ["date", "vdate", "narr", "ref", "debit", "credit", "bal"]),
    ("Axis", ["Tran Date", "CHQNO", "PARTICULARS", "DR", "CR", "BAL", "SOL"],
     ["date", "ref", "narr", "debit", "credit", "bal", "sno"]),
    ("Kotak", ["Sl. No.", "Transaction Date", "Value Date", "Description", "Chq / Ref No.", "Debit", "Credit", "Balance"],
     ["sno", "date", "vdate", "narr", "ref", "debit", "credit", "bal"]),
    ("Yes", ["Transaction Date", "Value Date", "Cheque No/Reference No", "Description", "Withdrawals", "Deposits",
             "Running Balance"], ["date", "vdate", "ref", "narr", "debit", "credit", "bal"]),
    ("IDFC First", ["Transaction Date", "Value Date", "Particulars", "Cheque No.", "Debit", "Credit", "Balance"],
     ["date", "vdate", "narr", "ref", "debit", "credit", "bal"]),
    ("Bank of Baroda", ["TRAN DATE", "VALUE DATE", "NARRATION", "CHQ.NO.", "WITHDRAWAL(DR)", "DEPOSIT(CR)", "BALANCE(INR)"],
     ["date", "vdate", "narr", "ref", "debit", "credit", "bal"]),
    ("PNB", ["Transaction Date", "Cheque Number", "Withdrawal", "Deposit", "Balance", "Narration"],
     ["date", "ref", "debit", "credit", "bal", "narr"]),
    ("Canara", ["Txn Date", "Value Date", "Cheque No.", "Description", "Branch Code", "Debit", "Credit", "Balance"],
     ["date", "vdate", "ref", "narr", "sno", "debit", "credit", "bal"]),
    ("IndusInd", ["Date", "Particulars", "Chq No", "Withdrawal", "Deposit", "Balance"],
     ["date", "narr", "ref", "debit", "credit", "bal"]),
    ("Kotak Amount + two Dr/Cr", ["Date", "Description", "Amount", "Dr / Cr", "Balance", "Dr / Cr"],
     ["date", "narr", "amount", "drcr", "bal", "baldrcr"]),
]


@pytest.mark.parametrize("bank,header,roles", HEADERS, ids=[h[0] for h in HEADERS])
def test_header_variants_of_indian_banks(bank, header, roles):
    txns = [("01/04/2026", "02/04/2026", "UPI/DR/412345678901/SHARMA STEEL TRADERS/HDFC", "debit", "1,000.00", "9,000.00"),
            ("03/04/2026", "03/04/2026", "NEFT CR-PATIL INTERIORS", "credit", "2,500.00", "11,500.00")]
    rows = []
    for i, (d, vd, narr, kind, amt, bal) in enumerate(txns):
        val = {"date": d, "vdate": vd, "narr": narr, "ref": "000123", "sno": str(i + 1), "bal": bal, "baldrcr": "CR",
               "debit": amt if kind == "debit" else "", "credit": amt if kind == "credit" else "",
               "amount": amt, "drcr": "DR" if kind == "debit" else "CR"}
        rows.append(tuple(val[r] for r in roles))
    r = bs.parse(simple_csv(rows, header=",".join(header)), f"{bank}.csv")
    assert [(t["date"], t["kind"], t["amount"], t["balance"]) for t in r["transactions"]] == [
        ("2026-04-01", "debit", 1000.0, 9000.0), ("2026-04-03", "credit", 2500.0, 11500.0)]
    assert r["opening_balance"] == 10000.0 and r["warnings"] == []


def test_csv_narration_spanning_lines_inside_quotes():
    data = ('Date,Narration,Withdrawal,Deposit,Balance\n01/04/2026,"UPI/DR/412345678901/SHARMA STEEL TRADERS\n/HDFC/Payment",'
            '"1,000.00",,"9,000.00"\n02/04/2026,NEFT CR-PATIL INTERIORS,,"500.00","9,500.00"\n').encode()
    r = bs.parse(data, "s.csv")
    assert len(r["transactions"]) == 2
    assert r["transactions"][0]["narration"] == "UPI/DR/XXXXXXXX8901/SHARMA STEEL TRADERS /HDFC/Payment"


def test_text_pdf_with_value_date_column_and_ref_column():
    pages = [["HDFC BANK LIMITED", "Account No : 000123456789",
              "Date Narration Chq./Ref.No. Value Dt Withdrawal Amt. Deposit Amt. Closing Balance",
              "01/04/26 UPI-SHARMA STEEL TRADERS-HDFC0ZZ0123-PAY 0000412345678901 02/04/26 12,500.00 1,37,500.00",
              "02/04/26 INT.PD:01-04-2026 to 30-04-2026 0000000000000000 02/04/26 312.00 1,37,812.00"]]
    r = bs.parse(text_pdf(pages), "x.pdf")
    t0, t1 = r["transactions"]
    assert t0["date"] == "2026-04-01" and t0["narration"] == "UPI-SHARMA STEEL TRADERS-HDFC0ZZ0123-PAY XXXXXXXXXXXX8901"
    # the range's first date stays; the second runs into the 16-digit ref (one space between), so that whole digit group
    # is masked down to its last 4 — the safe side of "digit groups split by spaces/hyphens"
    assert t1["narration"] == "INT.PD:01-04-2026 to XX-XX-XXXX XXXXXXXXXXXX0000"
    assert t1["kind"] == "credit"                                   # from the balance going up
    assert t0["kind"] == "debit"                                    # first row: no opening balance, "UPI-...-PAY" alone
    assert r["bank"] == "HDFC Bank"


def test_text_pdf_where_value_date_comes_first():
    pages = [["Value Date Transaction Date Particulars Amount Balance",
              "Opening Balance 10,000.00",
              "02/04/2026 01/04/2026 NEFT DR-GUPTA ESTATES-RENT 1,000.00 9,000.00"]]
    r = bs.parse(text_pdf(pages), "x.pdf")
    assert [(t["date"], t["kind"]) for t in r["transactions"]] == [("2026-04-01", "debit")]


# ---------------------------------------------------------------------------------------------------------------------
# Numbers and dates

@pytest.mark.parametrize("text,value,marker", [
    ("1,23,456.78", 123456.78, None), ("₹ 1,234", 1234.0, None), ("Rs. 12,500.00", 12500.0, None),
    ("INR 99.5", 99.5, None), ("1,234.00 Dr", 1234.0, "dr"), ("5,000.00 Cr", 5000.0, "cr"), ("12,500.00CR", 12500.0, "cr"),
    ("(1,234.00)", -1234.0, None), ("-1234", -1234.0, None), ("1234-", -1234.0, None), ("+250.00", 250.0, None),
    ("0.00", None, None), ("", None, None), ("-", None, None), ("Nil", None, None), ("1234,56", None, None),
])
def test_indian_amounts(text, value, marker):
    assert bs._amount(text) == (value, marker)


def test_numeric_amount_cells():
    assert bs._amount(1234.5) == (1234.5, None)
    assert bs._amount(0) == (None, None)
    assert bs._amount(None) == (None, None)


@pytest.mark.parametrize("text", ["01/04/26", "01/04/2026", "01-04-2026", "01.04.2026", "01-Apr-2026", "01 Apr 2026",
                                  "01-Apr-26", "1 April, 2026", "2026-04-01", "2026-04-01 00:00:00", "01APR2026",
                                  "01/04/2026 10:23 AM", "Apr 01, 2026"])
def test_date_formats(text):
    assert bs._to_date(text) == date(2026, 4, 1)


def test_date_cells_and_excel_serials():
    assert bs._to_date(datetime(2026, 4, 1, 9, 30)) == date(2026, 4, 1)
    assert bs._to_date(date(2026, 4, 1)) == date(2026, 4, 1)
    assert bs._to_date(excel_serial(date(2026, 4, 1))) == date(2026, 4, 1)
    assert bs._to_date(12500.0) is None            # an amount is not a date
    assert bs._to_date("31/02/2026") is None
    assert bs._to_date("UPI/DR/123") is None


def test_day_first_by_default_and_month_first_only_when_unambiguous():
    day_first = simple_csv([("03/04/2026", "NEFT CR-A ONE", "", "100.00", "1100.00"),
                            ("13/04/2026", "NEFT CR-A ONE", "", "100.00", "1200.00")])
    assert [t["date"] for t in bs.parse(day_first, "a.csv")["transactions"]] == ["2026-04-03", "2026-04-13"]
    us = simple_csv([("04/03/2026", "NEFT CR-A ONE", "", "100.00", "1100.00"),
                     ("04/25/2026", "NEFT CR-A ONE", "", "100.00", "1200.00")])
    assert [t["date"] for t in bs.parse(us, "a.csv")["transactions"]] == ["2026-04-03", "2026-04-25"]


def test_footers_totals_balance_rows_and_repeated_headers_are_skipped():
    data = simple_csv([
        ("", "Opening Balance", "", "", "1,00,000.00"),
        ("01/04/2026", "B/F", "", "", "1,00,000.00"),
        ("02/04/2026", "UPI/DR/412345678901/SHARMA STEEL TRADERS/HDFC", "10,000.00", "", "90,000.00"),
        ("", "Page 1 of 2", "", "", ""),
        ("Date", "Narration", "Withdrawal", "Deposit", "Balance"),
        ("03/04/2026", "NEFT CR-PATIL INTERIORS-INV 77", "", "25,000.00", "1,15,000.00"),
        ("", "Total", "10,000.00", "25,000.00", ""),
        ("", "Grand Total", "10,000.00", "25,000.00", ""),
        ("", "Closing Balance", "", "", "1,15,000.00"),
        ("*** End of Statement ***", "", "", "", ""),
        ("21/10/2026", "Printed by net banking", "", "", ""),
    ])
    r = bs.parse(data, "s.csv")
    assert [(t["date"], t["kind"], t["amount"]) for t in r["transactions"]] == [
        ("2026-04-02", "debit", 10000.0), ("2026-04-03", "credit", 25000.0)]
    assert r["opening_balance"] == 100000.0 and r["closing_balance"] == 115000.0
    assert r["warnings"] == []


def test_single_amount_column_without_marker_uses_balance():
    data = simple_csv([("01/04/2026", "B/F", "", "50,000.00"),
                       ("02/04/2026", "NEFT-PATIL INTERIORS", "20,000.00", "70,000.00"),
                       ("03/04/2026", "UPI-GUPTA ESTATES-RENT", "15,000.00", "55,000.00"),
                       ("04/04/2026", "SOMETHING", "5,000.00", "50,000.00")],
                      header="Date,Description,Amount,Balance")
    r = bs.parse(data, "s.csv")
    assert kinds(r) == ["credit", "debit", "debit"]
    assert r["warnings"] == []


def test_newest_first_statement_is_read_in_its_own_order():
    rows = [("04/04/2026", "SOMETHING", "5,000.00", "50,000.00"),
            ("03/04/2026", "UPI-GUPTA ESTATES-RENT", "15,000.00", "55,000.00"),
            ("02/04/2026", "NEFT-PATIL INTERIORS", "20,000.00", "70,000.00"),
            ("01/04/2026", "UPI/CR/412345678901/A ONE", "1,000.00", "50,000.00")]
    r = bs.parse(simple_csv(rows, header="Date,Description,Amount,Balance"), "s.csv")
    assert [t["date"] for t in r["transactions"]] == ["2026-04-04", "2026-04-03", "2026-04-02", "2026-04-01"]
    assert kinds(r) == ["debit", "debit", "credit", "credit"]     # last one from its "UPI/CR" narration
    assert r["opening_balance"] == 49000.0 and r["closing_balance"] == 50000.0


def test_signed_amount_column_without_balance():
    data = simple_csv([("01/04/2026", "UPI-GUPTA ESTATES", "-1,234.00"), ("02/04/2026", "NEFT-PATIL INTERIORS", "5,000.00")],
                      header="Date,Narration,Amount")
    r = bs.parse(data, "s.csv")
    assert kinds(r) == ["debit", "credit"]
    assert [t["amount"] for t in r["transactions"]] == [1234.0, 5000.0]
    assert r["opening_balance"] is None and r["closing_balance"] is None


def test_unknown_direction_is_marked_money_out_with_a_warning():
    data = simple_csv([("01/04/2026", "SOMETHING", "5,000.00")], header="Date,Narration,Amount")
    r = bs.parse(data, "s.csv")
    assert kinds(r) == ["debit"]
    assert any("money in or money out" in w for w in r["warnings"])


def test_amount_with_dr_cr_suffix_in_one_column():
    data = simple_csv([("01/04/2026", "UPI-GUPTA ESTATES", "1,234.00 Dr", "8,766.00 Cr"),
                       ("02/04/2026", "NEFT-PATIL INTERIORS", "5,000.00 Cr", "13,766.00 Cr")],
                      header="Date,Narration,Amount,Balance")
    r = bs.parse(data, "s.csv")
    assert kinds(r) == ["debit", "credit"]
    assert r["opening_balance"] == 10000.0


def test_overdraft_balance_with_dr_suffix():
    data = simple_csv([("01/04/2026", "UPI-GUPTA ESTATES", "1,000.00", "", "500.00 Dr")])
    r = bs.parse(data, "s.csv")
    assert r["transactions"][0]["balance"] == -500.0
    assert r["opening_balance"] == 500.0


def test_balance_as_on_today_in_the_preamble_is_not_taken_as_opening():
    data = ("Account Number,:,000123456789\nBalance as on 21 Oct 2026,:,\"5,55,555.00\"\n\n"
            "Txn Date,Description,Debit,Credit,Balance\n"
            "1 Apr 2026,NEFT DR-GUPTA ESTATES-RENT,\"1,000.00\",,\"9,000.00\"\n"
            "2 Apr 2026,NEFT CR-PATIL INTERIORS,,\"500.00\",\"9,500.00\"\n").encode()
    r = bs.parse(data, "s.csv")
    assert r["opening_balance"] == 10000.0          # worked out from the first row, not the 5,55,555 of today
    assert r["warnings"] == []


def test_text_overdraft_opening_balance_keeps_its_sign():
    pages = [["Date Particulars Withdrawal Deposit Balance", "Opening Balance 5,000.00 Dr",
              "01-04-2026 NEFT CR-PATIL INTERIORS 7,000.00 2,000.00 Cr"]]
    r = bs.parse(text_pdf(pages), "od.pdf")
    assert r["opening_balance"] == -5000.0
    assert [(t["kind"], t["amount"], t["balance"]) for t in r["transactions"]] == [("credit", 7000.0, 2000.0)]


def test_damaged_or_odd_files_never_crash():
    """Every truncation of every sample either parses or raises StatementError — never anything else."""
    samples = [hdfc_html_xls(), icici_xlsx(), SBI_CSV.encode(), axis_table_pdf(), text_pdf(KOTAK_PAGES)]
    with open(os.path.join(FIXTURES, "canara_like.xls"), "rb") as f:
        samples.append(f.read())
    for data in samples:
        for cut in (7, 50, 300, len(data) // 3, len(data) // 2, len(data) - 40):
            try:
                bs.parse(data[:cut], "x")
            except StatementError as e:
                assert e.code in {"password_needed", "wrong_password", "unsupported", "unreadable", "no_transactions",
                                  "too_big"}
        garbled = bytes(b ^ 0x5A if i % 97 == 0 else b for i, b in enumerate(data))
        try:
            bs.parse(garbled, "x")
        except StatementError:
            pass


def test_balance_that_does_not_add_up_gives_a_warning():
    data = simple_csv([("01/04/2026", "NEFT CR-A ONE", "", "1,000.00", "11,000.00"),
                       ("02/04/2026", "UPI/DR/412345678901/B TWO/HDFC", "500.00", "", "9,500.00")])
    r = bs.parse(data, "s.csv")
    assert any("running balance doesn't match" in w for w in r["warnings"])


def test_both_withdrawal_and_deposit_in_one_row():
    data = simple_csv([("01/04/2026", "ODD ROW", "100.00", "300.00", "10,200.00")])
    r = bs.parse(data, "s.csv")
    assert kinds(r) == ["debit", "credit"]
    assert any("both a withdrawal and a deposit" in w for w in r["warnings"])


# ---------------------------------------------------------------------------------------------------------------------
# Masking and payee

@pytest.mark.parametrize("raw,masked", [
    ("UPI/DR/412345678901/SHARMA S/HDFC", "UPI/DR/XXXXXXXX8901/SHARMA S/HDFC"),
    ("NEFT-N123456789012-PATEL", "NEFT-NXXXXXXXX9012-PATEL"),
    ("TRF TO A/C 000123456789", "TRF TO A/C XXXXXXXX6789"),
    ("POS 4160XXXXXXXX1234 SWIGGY", "POS 4160XXXXXXXX1234 SWIGGY"),          # card number already masked: unchanged
    ("zz.payee06@okaxis", "z***@okaxis"),
    ("UPI/9820000001@ybl/A ONE", "UPI/9***@ybl/A ONE"),
    ("PAY TO a.b@example.in DONE", "PAY TO a***@example.in DONE"),
    ("CALL 98200 00001", "CALL XXXXX X0001"),
    ("CARD 4111 1111 1111 1234", "CARD XXXX XXXX XXXX 1234"),
    ("TDS PAN ABCDE1234F", "TDS PAN XXXXXX234F"),
    ("RENT APR 2026 INV 12345", "RENT APR 2026 INV 12345"),                 # short numbers stay
    ("SHARMA STEEL TRADERS", "SHARMA STEEL TRADERS"),                       # names stay
])
def test_mask(raw, masked):
    assert bs.mask(raw) == masked
    assert bs.mask(bs.mask(raw)) == masked                                   # masking twice changes nothing


@pytest.mark.parametrize("narration,expected", [
    ("UPI/DR/412345678901/SHARMA S/HDFC/zz.vendor02@okhdfc/Payment", "Sharma S"),
    ("UPI/P2A/412345678901/KULKARNI AUTO PARTS/HDFC BANK/INV 2231", "Kulkarni Auto Parts"),
    ("UPI/SHARMA STEEL TRADERS/412345678901/Payment", "Sharma Steel Traders"),
    ("UPI-SHARMA STEEL TRADERS-zz.vendor01@okhdfcbank-HDFC0ZZ0123-412345678901-PAYMENT FROM PHONE", "Sharma Steel Traders"),
    ("NEFT DR-HDFC0ZZ0123-GUPTA ESTATES-NETBANK, MUM-N123456789012-OFFICE RENT", "Gupta Estates"),
    ("NEFT/N123456789012/PATEL AGRO EXPORTS/HDFC BANK", "Patel Agro Exports"),
    ("IMPS/P2A/412345678901/RAHUL TRANSPORT CO/HDFC", "Rahul Transport Co"),
    ("MMT/IMPS/412345678901/PATIL INTERIORS/ICICI Bank", "Patil Interiors"),
    ("RTGS-HDFC0ZZ0123-MEHTA MACHINE TOOLS-HDFCR52026040112345678", "Mehta Machine Tools"),
    ("POS 416021XXXXXX1234 SWIGGY BANGALORE", "Swiggy Bangalore"),
    ("ACH D- SHUBH FINCORP LTD-000987654321", "Shubh Fincorp Ltd"),
    ("NACH-DR-SHUBH FINCORP-XXXX1234", "Shubh Fincorp"),
    ("BY TRANSFER-PATIL INTERIORS", "Patil Interiors"),
    ("TO TRANSFER-INB IMPS/P2A/412345678901/SHARMA STEEL/HDFC", "Sharma Steel"),
    ("BY TRANSFER-NEFT*HDFC0ZZ0001*N123456789012*KULKARNI AUTO PARTS*", "Kulkarni Auto Parts"),
    ("CHQ PAID-MICR INWARD CLEARING-M/S SHARMA STEEL TRADERS", "Sharma Steel Traders"),
    ("BIL/ONL/000123456789/MSEDCL/LT CONSUMER", "MSEDCL"),
    ("UPI/412345678901/PAYMENT/zz.vendor03@okaxis/AXIS BANK", ""),               # no name in it: empty, not a guess
    ("ATM WDL-ATM CASH 4567 ANDHERI EAST", ""),
    ("SALARY APR 2026", ""),
    ("", ""),
])
def test_payee(narration, expected):
    assert bs.payee(narration) == expected
    assert bs.payee(bs.mask(narration)) == expected                         # works on stored (masked) narrations too


def test_narration_and_payee_length_limits():
    long = "NEFT DR-HDFC0ZZ0123-" + "VERY LONG SUPPLIER NAME " * 6 + "-" + "REMARK " * 40
    r = bs.parse(simple_csv([("01/04/2026", long, "1,000.00", "", "9,000.00")]), "s.csv")
    t = r["transactions"][0]
    assert len(t["narration"]) <= 200 and len(t["payee"]) <= 60 and t["payee"].startswith("Very Long Supplier")


# ---------------------------------------------------------------------------------------------------------------------
# Categories

DEBITS = [
    ("SALARY APR 2026", "salaries"),
    ("UPI/DR/412345678901/RAMESH K/HDFC/zz.payee07@okhdfc/SAL ADVANCE", "salaries"),
    ("NEFT DR-SBIN0ZZ0002-GUPTA ESTATES-NETBANK, MUM-N123456789012-OFFICE RENT APR", "rent"),
    ("RENTOKIL PEST CONTROL", "repairs"),                                   # RENT only as a whole word
    ("TORRENT POWER LTD AHMEDABAD BILL", "electricity"),
    ("BIL/ONL/000123456789/MSEDCL/LT CONSUMER", "electricity"),
    ("EPFO CHALLAN TRRN 1234567890123", "pf_esi"),
    ("ESIC CONTRIBUTION MAR 2026", "pf_esi"),
    ("PROF TAX MAHARASHTRA", "pf_esi"),
    ("GST CHALLAN CPIN 26042700012345", "taxes"),
    ("ITNS 280 ADVANCE TAX", "taxes"),
    ("TDS PAYMENT Q1", "taxes"),
    ("GST ON CHG", "bank_charges"),                                         # GST on bank charges is a bank charge
    ("SMS CHG FOR QTR", "bank_charges"),
    ("DEBIT CARD FEE INCL GST", "bank_charges"),
    ("UPI/DR/412345678901/RAHUL TRANSPORT/HDFC/zz.payee07@okhdfc/FREIGHT CHARGES", "freight"),   # "CHARGES" in a UPI remark
    ("UPI/DR/412345678901/AIRTEL/AIRP/zz.biller01@airtel/Postpaid", "phone_internet"),
    ("UPI/DR/412345678901/SHARMA STEEL TRADERS/AIRP/zz.vendor05@airtel/Payment", "purchases"),     # @airtel handle is not Airtel
    ("UPI/DR/412345678901/RELIANCE JIO/HDFC/zz.payee08@okicici/Recharge", "phone_internet"),
    ("UPI/DR/412345678901/JIOMART/HDFC/zz.payee08@okicici/Order", "office"),
    ("POS 416021XXXXXX1234 AMAZON WEB SERVICES", "software"),               # the AWS-vs-Amazon trap
    ("AWS SERVICES BLR", "software"),
    ("POS 416021XXXXXX1234 AMAZON PAY INDIA PRIV", "office"),
    ("ECOM PUR/FLIPKART/BANGALORE", "office"),
    ("GOOGLE WORKSPACE GSUITE", "software"),
    ("GOOGLE ADS 1234", "marketing"),
    ("FACEBK *AB12CD34", "marketing"),
    ("UPI/DR/412345678901/SHARMA METALS/HDFC/zz.payee10@okicici/Payment", "purchases"),   # METALS is not META
    ("IRCTC E TICKET 1234567890", "travel"),
    ("MMT/IMPS/412345678901/MAKEMYTRIP/HDFC", "travel"),
    ("MMT/IMPS/412345678901/PATIL INTERIORS/ICICI Bank", "other"),          # MMT is ICICI's IMPS prefix, not MakeMyTrip
    ("UPI/DR/412345678901/INDIAN OIL PETROL PUMP/HDFC", "fuel"),
    ("NETC FASTAG RECHARGE MH04AB1234", "fuel"),
    ("UPI/DR/412345678901/DELHIVERY LTD/HDFC", "freight"),
    ("NEFT DR-HDFC0ZZ0123-MEHTA AND ASSOCIATES-CA FEES FY 2025-26", "professional"),
    ("ACH D- SHUBH FINCORP LTD-000987654321", "loan_emi"),
    ("NACH DR LIC HOUSING FINANCE LTD", "loan_emi"),
    ("LIC OF INDIA PREMIUM", "insurance"),
    ("HDFC ERGO GENERAL INSURANCE", "insurance"),
    ("UPI/DR/412345678901/PREMIUM PLYWOOD TRADERS/HDFC", "purchases"),      # PREMIUM in a supplier's name
    ("POS 416021XXXXXX1234 SWIGGY", "food_staff"),
    ("ATM WDL-ATM CASH 4567 ANDHERI EAST", "cash_withdrawal"),
    ("CHQ PAID-SELF", "cash_withdrawal"),
    ("TRF TO OWN A/C XXXX1234", "transfer"),
    ("NEFT DR-HDFC0ZZ0123-MEHTA MACHINE TOOLS-NETBANK", "equipment"),
    ("UPI/DR/412345678901/SURESH KUMAR/HDFC", "other"),
]
CREDITS = [
    ("NEFT CR-ICIC0ZZ0456-PATEL AGRO EXPORTS-SHARMA STEEL TRADERS-ICICN52026040212345", "sales_receipt"),
    ("UPI/CR/412345678901/KULKARNI AUTO/HDFC/zz.payee09@okicici/INV 2231", "sales_receipt"),
    ("RAZORPAY SETTLEMENT 1234", "sales_receipt"),
    ("BY CLG-CHQ DEP-PATIL INTERIORS", "sales_receipt"),
    ("INT.PD:01-04-2026 to 30-04-2026", "interest_in"),
    ("CREDIT INTEREST CAPITALISED", "interest_in"),
    ("REFUND AMAZON ORDER 1234", "refund_in"),
    ("REV-UPI/DR/412345678901/SHARMA", "refund_in"),
    ("CASHBACK", "refund_in"),
    ("LOAN DISB SHUBH FINCORP", "loan_in"),
    ("CASH DEP ANDHERI", "cash_deposit"),
    ("CASH DEPOSIT BY SELF", "cash_deposit"),
    ("SWEEP IN FROM FD", "transfer_in"),
    ("DIVIDEND ABC LTD", "other_in"),
    ("MISC CREDIT", "other_in"),
]


@pytest.mark.parametrize("narration,category", DEBITS)
def test_categorize_debits(narration, category):
    t = {"narration": bs.mask(narration), "kind": "debit", "amount": 100.0, "date": "2026-04-01"}
    assert bs.categorize([t])[0]["category"] == category


@pytest.mark.parametrize("narration,category", CREDITS)
def test_categorize_credits(narration, category):
    t = {"narration": bs.mask(narration), "kind": "credit", "amount": 100.0, "date": "2026-04-01"}
    assert bs.categorize([t])[0]["category"] == category


def test_category_keys_are_the_fixed_ones():
    assert {c for _, c in DEBITS} <= set(bs.DEBIT_CATEGORIES)
    assert {c for _, c in CREDITS} <= set(bs.CREDIT_CATEGORIES)
    assert len(bs.DEBIT_CATEGORIES) == 23 and len(bs.CREDIT_CATEGORIES) == 7


@pytest.mark.parametrize("narration,category", [
    ("IMPS CHG 412345", "bank_charges"),                                       # no counterparty: a bank charge
    ("CHQ PAID-MICR CTS-SHARMA FREIGHT CHARGES", "freight"),                   # names someone: not a bank charge
    ("UPI/DR/412345678901/SHARMA STEEL TRADERS/HDFC/zz.payee10@okicici/GST INV 221", "purchases"),   # GST on a supplier's invoice
    ("GST PAYMENT", "taxes"),
    ("UPI/DR/412345678901/SELF/HDFC/zz.payee10@okicici/Transfer", "transfer"),                # UPI to own account at another bank
    ("CASH-SELF", "cash_withdrawal"),
])
def test_categorize_counterparty_aware_rules(narration, category):
    t = {"narration": bs.mask(narration), "kind": "debit", "amount": 100.0, "date": "2026-04-01"}
    assert bs.categorize([t])[0]["category"] == category


def test_categorize_adds_to_the_same_dicts_and_extra_rules_come_first():
    a = {"narration": "UPI/DR/XXXXXXXX8901/SURESH KUMAR/HDFC", "kind": "debit", "amount": 100.0, "date": "2026-04-01"}
    b = {"narration": "SALARY APR 2026", "kind": "debit", "amount": 100.0, "date": "2026-04-01"}
    c = {"narration": "NEFT CR-PATIL INTERIORS", "kind": "credit", "amount": 100.0, "date": "2026-04-01"}
    out = bs.categorize([a, b, c], extra_rules=[("debit", r"SURESH KUMAR", "repairs"), ("debit", r"SALARY", "office"),
                                                ("credit", r"PATIL", "loan_in"), ("debit", r"(", "rent"),
                                                ("debit", r"SURESH", "not_a_category"), ("both", r".", "rent")])
    assert out[0] is a and a["category"] == "repairs"
    assert b["category"] == "office"            # the owner's own rule wins over the built-in one
    assert c["category"] == "loan_in"


# ---------------------------------------------------------------------------------------------------------------------
# Recurring and summary

def txn(d, narration, amount, kind="debit"):
    n = bs.mask(narration)
    return {"date": d, "narration": n, "payee": bs.payee(n), "amount": amount, "kind": kind, "balance": None, "ref": d + n}


def sample_txns():
    return [
        txn("2026-04-01", "SALARY APR 2026", 180000.0),
        txn("2026-04-05", "NEFT DR-SBIN0ZZ0002-GUPTA ESTATES-NETBANK, MUM-N123456789012-OFFICE RENT APR", 45000.0),
        txn("2026-04-07", "UPI/DR/412345678901/SHARMA STEEL TRADERS/HDFC", 12500.0),
        txn("2026-04-08", "ATM WDL-ATM CASH 4567 ANDHERI EAST", 10000.0),
        txn("2026-04-09", "UPI/DR/412345678902/RAMU TEA STALL/HDFC", 200.0),
        txn("2026-04-10", "ACH D- SHUBH FINCORP LTD-000987654321", 18500.0),
        txn("2026-04-12", "NEFT CR-ICIC0ZZ0456-PATEL AGRO EXPORTS-INV 991", 250000.0, "credit"),
        txn("2026-04-20", "UPI/DR/412345678903/RAMU TEA STALL/HDFC", 220.0),
        txn("2026-05-01", "SALARY MAY 2026", 185000.0),
        txn("2026-05-05", "NEFT DR-SBIN0ZZ0002-GUPTA ESTATES-NETBANK, MUM-N123456789013-OFFICE RENT MAY", 45000.0),
        txn("2026-05-07", "UPI/DR/412345678904/SHARMA STEEL TRADERS/HDFC", 60000.0),
        txn("2026-05-08", "ATM WDL-ATM CASH 4567 ANDHERI EAST", 10000.0),
        txn("2026-05-10", "ACH D- SHUBH FINCORP LTD-000987654321", 18500.0),
        txn("2026-05-15", "RTGS-HDFC0ZZ0123-MEHTA MACHINE TOOLS-HDFCR52026051512345678", 250000.0),
        txn("2026-05-20", "TRF TO OWN A/C 000987654321", 50000.0),
        txn("2026-06-01", "SALARY JUN 2026", 182000.0),
        txn("2026-06-06", "NEFT DR-SBIN0ZZ0002-GUPTA ESTATES-NETBANK, MUM-N123456789014-OFFICE RENT JUN", 45000.0),
        txn("2026-06-12", "INT.PD:01-04-2026 to 30-06-2026", 600.0, "credit"),
    ]


def test_recurring_finds_rent_salary_and_emi_but_not_one_offs():
    rec = bs.recurring(bs.categorize(sample_txns()))
    got = {(r["payee"], r["category"]): r for r in rec}
    assert set(got) == {("Salary", "salaries"), ("Gupta Estates", "rent"), ("Shubh Fincorp Ltd", "loan_emi")}
    rent = got[("Gupta Estates", "rent")]
    assert (rent["amount"], rent["day_of_month"], rent["months"], rent["last_date"]) == (45000.0, 5, 3, "2026-06-06")
    sal = got[("Salary", "salaries")]
    assert (sal["amount"], sal["day_of_month"], sal["months"], sal["last_date"]) == (182000.0, 1, 3, "2026-06-01")
    emi = got[("Shubh Fincorp Ltd", "loan_emi")]
    assert (emi["amount"], emi["day_of_month"], emi["months"]) == (18500.0, 10, 2)
    # not recurring: steel (12,500 then 60,000 — outside ±15%), the machine (once), chai (one month only), ATM (cash),
    # the own-account transfer
    assert [r["payee"] for r in rec] == ["Salary", "Gupta Estates", "Shubh Fincorp Ltd"]


def test_recurring_works_without_categories_and_outlier_payment_is_ignored():
    ts = [txn("2026-04-05", "NEFT DR-GUPTA ESTATES-RENT", 45000.0), txn("2026-05-05", "NEFT DR-GUPTA ESTATES-RENT", 46000.0),
          txn("2026-05-20", "NEFT DR-GUPTA ESTATES-RENT", 5000.0)]
    rec = bs.recurring(ts)
    assert len(rec) == 1 and rec[0]["amount"] == 45500.0 and rec[0]["months"] == 2 and rec[0]["last_date"] == "2026-05-05"


def test_summary():
    s = bs.summary(bs.categorize(sample_txns()))
    assert s["months"] == [{"month": "2026-04", "in": 250000.0, "out": 266420.0},
                           {"month": "2026-05", "in": 0.0, "out": 618500.0},
                           {"month": "2026-06", "in": 600.0, "out": 227000.0}]
    assert s["money_in"] == 250600.0 and s["money_out"] == 1111920.0
    assert s["cash_withdrawn"] == 20000.0 and s["transfers_out"] == 50000.0
    assert s["by_category"]["salaries"] == 547000.0 and s["by_category"]["rent"] == 135000.0
    assert s["by_category"]["cash_withdrawal"] == 20000.0 and s["by_category"]["transfer"] == 50000.0
    assert sum(s["by_category"].values()) == s["money_out"]
    top = s["top_payees"]
    assert top[0] == {"payee": "Mehta Machine Tools", "amount": 250000.0, "count": 1}
    assert {"payee": "Gupta Estates", "amount": 135000.0, "count": 3} in top
    assert {"payee": "Sharma Steel Traders", "amount": 72500.0, "count": 2} in top
    assert all(p["payee"] for p in top) and len(top) <= 10
    assert [r["payee"] for r in s["recurring"]] == ["Salary", "Gupta Estates", "Shubh Fincorp Ltd"]


def test_summary_categorises_on_the_fly_without_changing_input():
    ts = sample_txns()
    s = bs.summary(ts)
    assert "category" not in ts[0] and s["by_category"]["salaries"] == 547000.0


# ---------------------------------------------------------------------------------------------------------------------
# Refs (de-duplication)

def test_refs_are_stable_across_reparse_and_unique_within_a_file():
    a = bs.parse(SBI_CSV.encode(), "a.csv")
    b = bs.parse(SBI_CSV.encode(), "renamed.csv")
    assert [t["ref"] for t in a["transactions"]] == [t["ref"] for t in b["transactions"]]
    refs = [t["ref"] for t in a["transactions"]]
    assert len(set(refs)) == len(refs)


def test_two_identical_transactions_on_one_day_get_different_refs():
    data = simple_csv([("01/04/2026", "UPI/DR/RAMU TEA STALL", "20.00", "", "980.00"),
                       ("01/04/2026", "UPI/DR/RAMU TEA STALL", "20.00", "", "960.00")])
    r = bs.parse(data, "s.csv")
    assert len({t["ref"] for t in r["transactions"]}) == 2


def test_refs_match_across_overlapping_statements_and_formats():
    april = [("01/04/2026", "UPI/DR/412345678901/SHARMA STEEL TRADERS/HDFC", "12,500.00", "", "87,500.00"),
             ("15/04/2026", "NEFT CR-PATIL INTERIORS-INV 77", "", "25,000.00", "1,12,500.00")]
    may = [("05/05/2026", "NEFT DR-GUPTA ESTATES-OFFICE RENT MAY", "45,000.00", "", "67,500.00"),
           ("20/05/2026", "UPI/CR/412345678902/KULKARNI AUTO/HDFC", "", "10,000.00", "77,500.00")]
    june = [("05/06/2026", "NEFT DR-GUPTA ESTATES-OFFICE RENT JUN", "45,000.00", "", "32,500.00")]
    apr_may = bs.parse(simple_csv(april + may), "apr-may.csv")
    # the second download spaces the narration differently and is an Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date", "Narration", "Withdrawal", "Deposit", "Balance"])
    for d, n, w, dep, bal in may + june:
        ws.append([datetime.strptime(d, "%d/%m/%Y"), n.replace("-", " - "), bs._amount(w)[0], bs._amount(dep)[0],
                   bs._amount(bal)[0]])
    buf = io.BytesIO()
    wb.save(buf)
    may_jun = bs.parse(buf.getvalue(), "may-jun.xlsx")
    refs1 = {t["date"]: t["ref"] for t in apr_may["transactions"]}
    refs2 = {t["date"]: t["ref"] for t in may_jun["transactions"]}
    assert refs1["2026-05-05"] == refs2["2026-05-05"] and refs1["2026-05-20"] == refs2["2026-05-20"]
    assert refs1["2026-04-01"] not in refs2.values() and refs2["2026-06-05"] not in refs1.values()


# ---------------------------------------------------------------------------------------------------------------------
# Errors and limits

def test_too_many_transactions():
    rows = [(f"{1 + i % 28:02d}/04/2026", f"UPI/DR/CHAI {i}", "1.00", "", "") for i in range(5001)]
    e = err(simple_csv(rows), "big.csv")
    assert e.code == "too_big" and "fewer months" in e.message


def test_exactly_5000_is_fine():
    rows = [(f"{1 + i % 28:02d}/04/2026", f"UPI/DR/CHAI {i}", "1.00", "", "") for i in range(5000)]
    assert len(bs.parse(simple_csv(rows), "ok.csv")["transactions"]) == 5000


def test_too_many_pdf_pages():
    e = err(text_pdf([["Page"]] * 61), "long.pdf")
    assert e.code == "too_big" and "60 pages" in e.message


def test_file_size_limit():
    e = err(b"x" * (bs.MAX_FILE_BYTES + 1), "huge.csv")
    assert e.code == "too_big"


@pytest.mark.parametrize("data", [
    b"Date,Narration,Withdrawal,Deposit,Balance\n",
    b"Hello, this is not a bank statement.\nJust some text.\n",
    "<html><body><p>No table here</p></body></html>".encode(),
])
def test_no_transactions(data):
    e = err(data, "x.csv")
    assert e.code == "no_transactions"
    assert e.message == ("We couldn't find the transactions in this file. Download the statement as Excel or CSV from "
                         "net banking and try again.")


def test_pdf_with_text_but_no_transactions():
    e = err(text_pdf([["Statement of account", "There are no transactions for this period. " * 3]]), "x.pdf")
    assert e.code == "no_transactions"


def _patch_sizes(zip_bytes: bytes, size: int) -> bytes:
    """Rewrite each central-directory entry's uncompressed size, like a zip bomb would claim."""
    b = bytearray(zip_bytes)
    i = b.find(b"PK\x01\x02")
    while i != -1:
        b[i + 24:i + 28] = struct.pack("<I", size)
        i = b.find(b"PK\x01\x02", i + 4)
    return bytes(b)


def test_zip_bomb_guard():
    bomb = _patch_sizes(icici_xlsx(), 100 * 1024 * 1024)
    assert sum(i.file_size for i in zipfile.ZipFile(io.BytesIO(bomb)).infolist()) > bs.MAX_UNZIPPED
    e = err(bomb, "bomb.xlsx")
    assert e.code == "too_big"


def test_unsupported_and_unreadable_files():
    assert err(b"\x00\x01\x02\x03binary\x00\x00" * 10, "x.bin").code == "unsupported"
    docx = io.BytesIO()
    with zipfile.ZipFile(docx, "w") as z:
        z.writestr("word/document.xml", "<w:document/>")
    assert err(docx.getvalue(), "x.xlsx").code == "unsupported"
    assert err(b"\x89PNG\r\n\x1a\n" + b"\x00" * 50, "photo.png").code == "unreadable"
    assert err(b"", "empty.csv").code == "unreadable"
    assert err(b"%PDF-1.4 not really a pdf", "broken.pdf").code == "unreadable"
    assert err(b"PK\x03\x04 broken zip", "broken.xlsx").code == "unreadable"


def test_error_messages_are_plain_and_codes_known():
    codes = {"password_needed", "wrong_password", "unsupported", "unreadable", "no_transactions", "too_big"}
    for data, name in [(b"", "a.csv"), (b"\x00\x01" * 20, "a.bin"), (b"Date,Narration,Amount\n", "a.csv")]:
        e = err(data, name)
        assert e.code in codes and e.message and "Traceback" not in e.message and "Error" not in e.message


def test_result_shape_matches_the_interface():
    r = bs.parse(SBI_CSV.encode(), "s.csv")
    assert set(r) == {"format", "bank", "account_last4", "from", "to", "opening_balance", "closing_balance",
                      "transactions", "warnings"}
    for t in r["transactions"]:
        assert set(t) == {"date", "narration", "payee", "amount", "kind", "balance", "ref", "guessed"}
        assert t["guessed"] is False
        assert t["amount"] > 0 and round(t["amount"], 2) == t["amount"] and t["kind"] in ("debit", "credit")
        assert len(t["date"]) == 10 and isinstance(t["ref"], str) and len(t["ref"]) >= 16


# ---------------------------------------------------------------------------------------------------------------------
# Hostile files: slow patterns, bombs, nesting (review fixes)

import threading  # noqa: E402
import time  # noqa: E402
import zlib  # noqa: E402

HOSTILE = ["A" * 60000, "a." * 30000, "1-" * 30000, "1 " * 30000, "UPI/" * 15000, "NACH " * 12000, "@" * 60000,
           "a@" * 30000, "1,1" * 20000, "X" * 60000 + "&", " 111" * 15000 + " Z", "(" * 60000, "0" * 60000,
           "1" + " " * 30000 + "x", "01/04/2026" + " " * 30000 + "x"]


def test_hostile_10mb_csv_parses_fast():
    rows, size, k = ["Date,Narration,Withdrawal,Deposit,Balance"], 0, 0
    while size < 10_000_000:
        h = HOSTILE[k % len(HOSTILE)]
        for r in (f'01/04/2026,"{h}","{h[:30000]}",,"{h[:30000]}"', f'"{h[:30000]}",NEFT DR-GUPTA ESTATES,100.00,,9900.00',
                  f"02/04/2026,UPI/DR/{k:012d}/SHARMA STEEL TRADERS/HDFC,100.00,,9900.00"):
            rows.append(r)
            size += len(r)
        k += 1
    data = ("\n".join(rows) + "\n").encode()
    assert len(data) >= 10_000_000
    t = time.perf_counter()
    r = bs.parse(data, "hostile.csv")
    assert time.perf_counter() - t < 2.0
    assert all(len(x["narration"]) <= 200 for x in r["transactions"]) and len(r["transactions"]) >= 100


def test_hostile_xlsx_parses_fast():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date", "Narration", "Withdrawal", "Deposit", "Balance"])
    size, k = 0, 0
    while size < 10_000_000:
        h = HOSTILE[k % len(HOSTILE)][:32000]           # Excel's own cell limit is 32,767 characters
        ws.append(["01/04/2026", h, h, None, h])
        ws.append(["02/04/2026", f"UPI/DR/{k:012d}/SHARMA STEEL TRADERS/HDFC", 100.0, None, 9900.0])
        size += 3 * len(h)
        k += 1
    buf = io.BytesIO()
    wb.save(buf)
    t = time.perf_counter()
    r = bs.parse(buf.getvalue(), "hostile.xlsx")
    assert time.perf_counter() - t < 2.0
    assert len(r["transactions"]) == k


def test_helpers_are_linear_on_huge_strings():
    big = "a." * 300_000 + "1-" * 300_000 + "A" * 600_000 + "UPI/" * 100_000
    t = time.perf_counter()
    bs.mask(big)
    bs.payee(big)
    assert bs._amount(big) == (None, None) and bs._to_date(big) is None
    bs.categorize([{"narration": h, "kind": k, "amount": 1.0, "date": "2026-04-01"} for h in HOSTILE + [big]
                   for k in ("debit", "credit")])
    bs._label(big)
    assert time.perf_counter() - t < 2.0


def test_nested_html_tables_are_refused_quickly():
    nested = "<table><tr><td>" * 20000 + "x" + "</td></tr></table>" * 20000
    t = time.perf_counter()
    e = err(nested.encode(), "x.xls")
    assert time.perf_counter() - t < 1.0
    assert e.code == "unreadable" and e.message == bs.MSG_TOO_COMPLEX
    wide = ("<table><tr>" + "<td>1</td>" * (bs.MAX_HTML_CELLS + 1) + "</tr></table>").encode()
    assert err(wide, "x.html").message == bs.MSG_TOO_COMPLEX


def test_html_layout_table_entities_and_scripts():
    html = """<html><script>var x = "<table><tr><td>01/04/2026</td></tr></table>";</script><table><tr><td>
    <table><tr><td>Account No :000123456789</td></tr></table>
    <table><tr><th>Date</th><th>Narration</th><th>Withdrawal Amt.</th><th>Deposit Amt.</th><th>Closing Balance</th></tr>
    <tr><td>01/04/26</td><td>NEFT DR-GUPTA&nbsp;ESTATES &amp; SONS<br>RENT</td><td>1,000.00</td><td></td><td>9,000.00</td></tr>
    <tr><td colspan="2">01/04/26 total</td><td>1,000.00</td><td></td><td></td></tr>
    </table></td></tr></table></html>"""
    r = bs.parse(html.encode(), "x.xls")
    assert [(t["date"], t["narration"], t["amount"]) for t in r["transactions"]] == [
        ("2026-04-01", "NEFT DR-GUPTA ESTATES & SONS RENT", 1000.0)]
    assert r["account_last4"] == "6789"


def pdf_bomb(n: int) -> bytes:
    """A one-page PDF whose content stream inflates to 14*n bytes of drawing commands."""
    comp = zlib.compress(b"0 0 m 1 1 l S\n" * n, 9)
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << >> >>",
            b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(comp) + comp + b"\nendstream"]
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 5\n0000000000 65535 f \n" + b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size 5 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref
    return bytes(out)


def test_pdf_decompression_bomb_is_refused_without_inflating_it():
    bomb = pdf_bomb(1_500_000)                               # 41 KB that inflates to 21 MB of drawing commands
    assert len(bomb) < 50_000
    t = time.perf_counter()
    e = err(bomb, "bomb.pdf")
    assert e.code == "unreadable" and e.message == bs.MSG_TOO_COMPLEX
    assert time.perf_counter() - t < 2.0
    assert bs.parse(axis_table_pdf(), "a.pdf")["transactions"]      # ordinary PDFs are well under the cap


def test_isolated_gives_the_same_answer_as_parse():
    assert bs.parse_isolated(SBI_CSV.encode(), "s.csv") == bs.parse(SBI_CSV.encode(), "s.csv")
    locked = encrypt(axis_table_pdf(), "SHAR0104")
    assert bs.parse_isolated(locked, "s.pdf", "SHAR0104") == bs.parse(locked, "s.pdf", "SHAR0104")
    with pytest.raises(StatementError) as e:
        bs.parse_isolated(locked, "s.pdf", "WRONG999")
    assert e.value.code == "wrong_password" and "WRONG999" not in e.value.message
    with pytest.raises(StatementError) as e:
        bs.parse_isolated(locked, "s.pdf")
    assert e.value.code == "password_needed"
    with pytest.raises(StatementError) as e:
        bs.parse_isolated(scanned_pdf(), "scan.pdf")
    assert e.value.code == "unreadable" and e.value.message == bs.MSG_SCAN


def test_isolated_keeps_password_and_file_out_of_argv_and_env(monkeypatch):
    seen = []
    real = bs.subprocess.Popen

    def spy(args, **kw):
        seen.append((list(args), dict(kw.get("env") or {})))
        return real(args, **kw)

    monkeypatch.setattr(bs.subprocess, "Popen", spy)
    monkeypatch.setenv("DATABASE_URL", "mongodb://secret-server-credential")
    locked = encrypt(axis_table_pdf(), "SHAR0104")
    bs.parse_isolated(locked, "My Statement.pdf", "SHAR0104")
    (argv, env), = seen
    flat = " ".join(argv) + " " + " ".join(f"{k}={v}" for k, v in env.items())
    assert "SHAR0104" not in flat and "My Statement" not in flat and "secret-server-credential" not in flat
    assert argv[1:3] == ["-I", "-B"] and "--child" in argv and set(env) <= {"PATH"}


def test_isolated_kills_a_pdf_bomb_at_the_timeout():
    t = time.perf_counter()
    with pytest.raises(StatementError) as e:
        bs.parse_isolated(pdf_bomb(70_000), "bomb.pdf", timeout=3)        # 0.98 MB of drawing commands: under the cap, slow
    assert e.value.code == "unreadable" and e.value.message == bs.MSG_TOO_COMPLEX
    assert time.perf_counter() - t < 8


def test_isolated_memory_limit_applies(monkeypatch):
    monkeypatch.setattr(bs, "CHILD_MEMORY", 100 * 1024 * 1024)          # too little to even load the PDF library
    t = time.perf_counter()
    with pytest.raises(StatementError) as e:
        bs.parse_isolated(pdf_bomb(70_000), "bomb.pdf", timeout=30)
    assert e.value.message == bs.MSG_TOO_COMPLEX
    assert time.perf_counter() - t < 15                                 # the memory limit, not the 30 s timeout


def test_isolated_runs_at_most_two_at_once(monkeypatch):
    running, peak, lock = [0], [0], threading.Lock()

    def slow_child(data, filename, password, timeout):
        with lock:
            running[0] += 1
            peak[0] = max(peak[0], running[0])
        time.sleep(0.3)
        with lock:
            running[0] -= 1
        return {"ok": True}

    monkeypatch.setattr(bs, "_run_child", slow_child)
    threads = [threading.Thread(target=bs.parse_isolated, args=(b"x", "x.csv")) for _ in range(6)]
    for th in threads:
        th.start()
    for th in threads:
        th.join()
    assert peak[0] == 2


def test_isolated_says_busy_when_no_slot_frees_up():
    assert bs._SLOTS.acquire(timeout=5) and bs._SLOTS.acquire(timeout=5)
    try:
        with pytest.raises(StatementError) as e:
            bs.parse_isolated(SBI_CSV.encode(), "s.csv", timeout=1)
        assert e.value.code == "unreadable" and e.value.message == bs.MSG_BUSY
    finally:
        bs._SLOTS.release()
        bs._SLOTS.release()


def test_isolated_cheap_checks_and_untrusted_child_answers():
    with pytest.raises(StatementError) as e:
        bs.parse_isolated(b"", "x.csv")
    assert e.value.code == "unreadable"
    with pytest.raises(StatementError) as e:
        bs.parse_isolated(b"x" * (bs.MAX_FILE_BYTES + 1), "x.csv")
    assert e.value.code == "too_big"
    for bad in [None, [], {"ok": True, "result": {"format": "csv"}}, {"ok": False, "code": "rm -rf", "message": "x"},
                {"ok": False, "code": "unreadable", "message": "x" * 600}]:
        with pytest.raises(StatementError) as e:
            bs._checked(bad)
        assert e.value.message == bs.MSG_TOO_COMPLEX
    with pytest.raises(StatementError) as e:
        bs._checked({"ok": False, "code": "no_transactions", "message": bs.MSG_NO_TXNS})
    assert e.value.code == "no_transactions" and e.value.message == bs.MSG_NO_TXNS


# ---------------------------------------------------------------------------------------------------------------------
# Guessed directions, grouped digits, fast recurring (review fixes)

def test_guessed_direction_is_flagged_per_transaction():
    data = simple_csv([("01/04/2026", "NEFT-PATIL INTERIORS INV 77", "1,80,000.00"),
                       ("02/04/2026", "SOMETHING", "5,000.00"),
                       ("03/04/2026", "UPI/CR/412345678901/A ONE/HDFC", "1,000.00")],
                      header="Date,Narration,Amount")
    r = bs.parse(data, "s.csv")
    assert [(t["kind"], t["guessed"]) for t in r["transactions"]] == [("debit", True), ("debit", True), ("credit", True)]
    assert any(w.startswith("2 of 3 transactions") and "flagged" in w for w in r["warnings"])
    assert any(w.startswith("1 of 3 transactions") and "description" in w for w in r["warnings"])


def test_directions_from_the_statement_are_not_flagged():
    for data, name in [(SBI_CSV.encode(), "s.csv"), (text_pdf(KOTAK_PAGES), "k.pdf"), (icici_xlsx(), "i.xlsx")]:
        r = bs.parse(data, name)
        assert not any(t["guessed"] for t in r["transactions"])
        assert not any("flagged" in w for w in r["warnings"])


@pytest.mark.parametrize("raw,masked", [
    ("CARD 4111-1111-1111-1234", "CARD XXXX-XXXX-XXXX-1234"),
    ("CALL 98200-12345", "CALL XXXXX-X2345"),
    ("CALL 98200 12345 TODAY", "CALL XXXXX X2345 TODAY"),
    ("AADHAAR 1234 5678 9012", "AADHAAR XXXX XXXX 9012"),
    ("PHONE +91-98200-12345", "PHONE +XX-XXXXX-X2345"),
    ("ON 01-04-2026", "ON 01-04-2026"),                                  # 8 digits: a date, left alone
    ("INV 1234-5678", "INV 1234-5678"),
    ("AMT 1,23,456.78 PAID", "AMT 1,23,456.78 PAID"),
])
def test_mask_digit_groups(raw, masked):
    assert bs.mask(raw) == masked
    assert bs.mask(masked) == masked


def test_recurring_is_fast_with_many_payments_to_one_payee():
    ts = [{"date": f"2026-{1 + i % 12:02d}-{5 + i % 3:02d}", "narration": "NEFT DR-GUPTA ESTATES-RENT", "payee": "Gupta Estates",
           "kind": "debit", "amount": 45000.0 + (i % 7) * 10, "category": "rent"} for i in range(5000)]
    ts += [{"date": f"2026-{1 + i % 2:02d}-20", "narration": "NEFT DR-GUPTA ESTATES-RENT", "payee": "Gupta Estates",
            "kind": "debit", "amount": 900.0, "category": "rent"} for i in range(6000)]   # a smaller, more frequent amount
    t = time.perf_counter()
    rec = bs.recurring(ts)
    assert time.perf_counter() - t < 0.5
    assert len(rec) == 1 and rec[0]["months"] == 12 and rec[0]["amount"] == 45030.0 and rec[0]["day_of_month"] == 6
