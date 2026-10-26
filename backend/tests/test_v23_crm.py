"""v2.3 Conversion engine: the CRM (contacts, companies, history, follow-ups, import and export) and the one-page Pipeline &
Opportunities (4 questions pre-filled, stages on the Football Field, the drag-and-drop board, won / lost). Every business,
person, number and email here is invented."""
import csv
import io
from datetime import datetime, timedelta

import openpyxl

from app import crm, plans
from app.agents import scheduler
from app.db import IST, db, now
from tests.conftest import new_owner
from tests.v23_helpers import owner_id, setup_v23, v23_member


async def paste(c, name: str, phone: str, message: str = "Need a quote for a modular kitchen") -> str:
    r = await c.post("/api/inquiries/whatsapp", json={"name": name, "phone": phone, "message": message})
    assert r.status_code == 200, r.text
    return r.json()["id"]


async def test_every_new_inquiry_files_a_contact_and_lands_on_the_first_stage(client):
    c = await v23_member("9800100001")
    lead_id = await paste(c, "Ritu Bansal", "98200 11111")
    await paste(c, "", "98200 11111", "Also the wardrobe price please")      # same person, second inquiry
    lead = await db().leads.find_one({"_id": lead_id})
    assert lead["contact_id"] and lead["pstage"] == "identification" and lead["stage"] == "identification"
    contacts = (await c.get("/api/crm/contacts")).json()
    assert contacts["total"] == 1
    ritu = contacts["items"][0]
    assert ritu["name"] == "Ritu Bansal" and ritu["phone"] == "+919820011111" and ritu["source"] == "whatsapp"
    assert ritu["deals"] == 2 and ritu["open_deals"] == 2
    detail = (await c.get(f"/api/crm/contacts/{ritu['id']}")).json()
    assert len(detail["deals_list"]) == 2 and sum(1 for a in detail["activities"] if a["kind"] == "created") == 2
    # the board shows them before any pipeline is built: the 5 Football Field stages
    deals = (await c.get("/api/crm/deals")).json()
    assert [s["id"] for s in deals["stages"]] == ["identification", "logic", "pain", "vision", "close"] and not deals["built"]
    assert deals["stages"][0]["count"] == 2 and deals["totals"]["open"] == 2


async def test_pipeline_questions_prefilled_built_with_ai_and_edited(client):
    c = await v23_member("9800100002")
    await paste(c, "Kunal Mehta", "98200 22222")
    g = (await c.get("/api/crm/pipeline")).json()
    assert g["pipeline"] is None and g["token_cost"] == 1
    pre = g["prefill"]
    assert "Home and office interiors" in pre["sell"] and "Modular kitchen" in pre["sell"] and pre["customers"].startswith("Families")
    assert "whatsapp" in pre["sources"] and pre["path"] and pre["reasons"] and g["where"]["sell"] == "From your Knowledge Hub"
    used = (await c.get("/api/me")).json()["runs"]["used"]
    answers = {**{k: v for k, v in pre.items() if k != "pain_points"}, "path": ["visit", "quote"], "sources": ["website", "referral", "nope"],
               "reasons": ["Went with a local carpenter"]}
    r = await c.post("/api/crm/pipeline/build", json={"answers": answers, "use_ai": True})
    assert r.status_code == 200, r.text
    built = r.json()
    stages = built["stages"]
    assert stages[0]["ff"] == "identification" and stages[-1]["ff"] == "close" and 4 <= len(stages) <= 10
    assert [s["ff"] for s in stages] == sorted((s["ff"] for s in stages), key=lambda f: crm.FF.index(f))
    assert "Went with a local carpenter" in built["pipeline"]["lost_reasons"] and built["pipeline"]["answers"]["sources"] == ["website", "referral"]
    assert (await c.get("/api/me")).json()["runs"]["used"] == used + 1          # one AI token
    # the open deal moved onto the new first stage
    deals = (await c.get("/api/crm/deals")).json()
    assert deals["built"] and deals["deals"][0]["pstage"] == stages[0]["id"]
    # rename, reorder, add, remove; a stage moved to another part of the field takes its deals with it
    quote_stage = next(s for s in stages if s["ff"] == "vision")
    deal_id = deals["deals"][0]["id"]
    assert (await c.post(f"/api/crm/deals/{deal_id}/move", json={"pstage": quote_stage["id"]})).json()["stage"] == "vision"
    edited = [{"id": stages[0]["id"], "name": "Fresh inquiry", "ff": "identification"},
              {"id": quote_stage["id"], "name": "Quote with them", "ff": "pain"},
              {"name": "Advance paid", "ff": "close"}]
    r = await c.put("/api/crm/pipeline/stages", json={"stages": edited, "lost_reasons": ["Price too high", "Chose a competitor"]})
    assert r.status_code == 200, r.text
    new = r.json()["stages"]
    assert [s["name"] for s in new] == ["Fresh inquiry", "Quote with them", "Advance paid"] and new[1]["id"] == quote_stage["id"]
    lead = await db().leads.find_one({"_id": deal_id})
    assert lead["stage"] == "pain" and lead["pstage"] == quote_stage["id"]
    assert (await c.put("/api/crm/pipeline/stages", json={"stages": [edited[0], {**edited[0], "id": None}]})).status_code == 400
    # rebuilt without AI: free, rules only
    r = await c.post("/api/crm/pipeline/build", json={"answers": {**answers, "path": ["direct"]}, "use_ai": False})
    assert r.status_code == 200 and r.json()["pipeline"]["built_by"] == "rules"
    assert (await c.get("/api/me")).json()["runs"]["used"] == used + 1
    assert [s["name"] for s in r.json()["stages"]] == ["New inquiry", "First reply done", "Ready to buy", "Order confirmed"]


async def test_board_drag_won_with_value_lost_with_reason_and_history(client):
    c = await v23_member("9800100003")
    me = (await c.get("/api/me")).json()
    await c.post("/api/crm/pipeline/build", json={"answers": {"sell": "Interiors", "customers": "Families", "path": ["call", "quote"]},
                                                  "use_ai": False})
    stages = (await c.get("/api/crm/deals")).json()["stages"]
    r = await c.post("/api/crm/deals", json={"name": "Asha Kulkarni", "phone": "98200 33333", "company": "Kulkarni Builders",
                                             "value": 185000, "follow_up_on": (now() + timedelta(days=2)).astimezone(IST).date().isoformat(),
                                             "note": "Wants a 3BHK quote"})
    assert r.status_code == 200, r.text
    deal = r.json()
    assert deal["pstage"] == stages[0]["id"] and deal["contact"]["company"] == "Kulkarni Builders"
    assert await db().agent_jobs.find_one({"_id": f"crmfu:deal:{deal['id']}", "status": "scheduled"})
    # drag to the quotation stage
    quote = next(s for s in stages if s["name"] == "Quotation sent")
    moved = (await c.post(f"/api/crm/deals/{deal['id']}/move", json={"pstage": quote["id"]})).json()
    assert moved["pstage"] == quote["id"] and moved["stage"] == "vision"
    assert (await c.post(f"/api/crm/deals/{deal['id']}/move", json={"pstage": "nope"})).status_code == 400
    # won with its value: the customer list gets them, the follow-up reminder is cancelled
    won = (await c.patch(f"/api/crm/deals/{deal['id']}", json={"status": "won", "value": 210000, "note": "Advance received"})).json()
    assert won["status"] == "won" and won["value"] == 210000 and won["won_at"]
    cust = await db().customers.find_one({"ws": me["user"]["id"], "phone": "+919820033333"})
    assert cust and cust["total_value"] == 210000
    assert (await db().agent_jobs.find_one({"_id": f"crmfu:deal:{deal['id']}"}))["status"] == "cancelled"
    kinds = [a["kind"] for a in won["activities"]]
    assert {"created", "stage", "won", "note"} <= set(kinds)
    # another one lost, with a reason from the pipeline's list
    lost = (await c.post("/api/crm/deals", json={"name": "Vivek Rao", "phone": "98200 44444"})).json()
    reasons = (await c.get("/api/crm/deals")).json()["lost_reasons"]
    out = (await c.patch(f"/api/crm/deals/{lost['id']}", json={"status": "lost", "lost_reason": reasons[0]})).json()
    assert out["status"] == "lost" and out["lost_reason"] == reasons[0]
    body = (await c.get("/api/crm/deals")).json()
    assert body["totals"]["open"] == 0 and body["totals"]["won_30"] == 1 and body["totals"]["won_30_value"] == 210000
    assert body["week"]["won"] == 1 and body["week"]["lost"] == 1 and body["week"]["new"] == 2
    # the dashboard's own pipeline numbers still read the Football Field stage
    assert (await c.get("/api/dashboard")).status_code == 200


async def test_contacts_follow_up_reminder_activities_companies_and_delete(client):
    c = await v23_member("9800100004")
    uid = await owner_id(c)
    r = await c.post("/api/crm/contacts", json={"name": "Neha Joshi", "phone": "98200 55555", "email": "Neha@Example.in",
                                                "company": "Joshi Traders", "tags": ["Architect", "architect", "VIP!"]})
    assert r.status_code == 200, r.text
    neha = r.json()
    assert neha["email"] == "neha@example.in" and neha["tags"] == ["Architect", "VIP"]
    assert (await c.post("/api/crm/contacts", json={"name": "Neha J", "email": "neha@example.in"})).status_code == 409
    assert (await c.post("/api/crm/contacts", json={"name": "No Way"})).status_code == 400
    await c.post("/api/crm/contacts", json={"name": "Ravi Joshi", "phone": "98200 66666", "company": "joshi traders"})
    day = (now() + timedelta(days=1)).astimezone(IST).date().isoformat()
    out = (await c.patch(f"/api/crm/contacts/{neha['id']}", json={"follow_up_on": day, "follow_up_note": "Send the catalogue"})).json()
    assert out["follow_up_on"] == day and any(a["kind"] == "followup" for a in out["activities"])
    job = await db().agent_jobs.find_one({"_id": f"crmfu:contact:{neha['id']}"})
    assert job["status"] == "scheduled" and job["run_at"].replace(tzinfo=None) == datetime.fromisoformat(f"{day}T04:30:00")   # 10:00 IST
    # the reminder runs on the day: a notice in the hub, nothing sent to the customer
    await scheduler.run_due_jobs(now() + timedelta(days=2))
    assert (await db().agent_jobs.find_one({"_id": job["_id"]}))["status"] == "done"
    assert await db().notices.find_one({"user_id": uid, "text": "Follow up today: Neha Joshi — Send the catalogue"})
    assert await db().approvals.count_documents({}) == 0
    # follow-ups list, a logged call, companies grouped
    fu = (await c.get("/api/crm/followups")).json()
    assert [i["name"] for i in fu["upcoming"]] == ["Neha Joshi"] and fu["overdue"] == [] and fu["today"] == [] and fu["date"]
    a = (await c.post(f"/api/crm/contacts/{neha['id']}/activities", json={"kind": "call", "text": "Spoke about the showroom visit"})).json()
    assert a["kind"] == "call" and a["text"] == "Call: Spoke about the showroom visit"
    comps = (await c.get("/api/crm/companies")).json()["items"]
    assert [(x["name"], x["contacts"]) for x in comps] == [("Joshi Traders", 2)]
    found = (await c.get("/api/crm/contacts", params={"q": "55555"})).json()["items"]
    assert [x["name"] for x in found] == ["Neha Joshi"]
    assert [x["name"] for x in (await c.get("/api/crm/contacts", params={"tag": "VIP"})).json()["items"]] == ["Neha Joshi"]
    # a new date moves the reminder; clearing it cancels it
    await c.patch(f"/api/crm/contacts/{neha['id']}", json={"follow_up_on": (now() + timedelta(days=5)).astimezone(IST).date().isoformat()})
    assert (await db().agent_jobs.find_one({"_id": job["_id"]}))["status"] == "scheduled"
    await c.patch(f"/api/crm/contacts/{neha['id']}", json={"follow_up_on": ""})
    assert (await db().agent_jobs.find_one({"_id": job["_id"]}))["status"] == "cancelled"
    # phone clash, then delete (deals keep living)
    ravi = next(x for x in (await c.get("/api/crm/contacts")).json()["items"] if x["name"] == "Ravi Joshi")
    assert (await c.patch(f"/api/crm/contacts/{ravi['id']}", json={"phone": "9820055555"})).status_code == 409
    assert (await c.delete(f"/api/crm/contacts/{ravi['id']}")).status_code == 200
    assert (await c.get(f"/api/crm/contacts/{ravi['id']}")).status_code == 404


async def test_import_csv_and_excel_then_export_without_formulas(client):
    c = await v23_member("9800100005")
    await paste(c, "Pooja Shah", "98200 77777")
    text = "Name,Mobile,Email,Company name,Tags,Notes\n" \
           "Pooja Shah,9820077777,,Shah Exports,Repeat,\n" \
           "Arjun Nair,98200 88888,arjun@example.in,,\"Architect, Referral\",=HYPERLINK(\"x\")\n" \
           "Nobody,,,,,\n"
    r = await c.post("/api/crm/import", files={"file": ("list.csv", text.encode(), "text/csv")})
    assert r.status_code == 200, r.text
    assert r.json() == {"added": 1, "matched": 1, "skipped": 1}
    pooja = next(x for x in (await c.get("/api/crm/contacts")).json()["items"] if x["name"] == "Pooja Shah")
    assert pooja["company"] == "Shah Exports" and pooja["tags"] == ["Repeat"]
    wb = openpyxl.Workbook()
    wb.active.append(["Customer", "Phone", "E-mail"])
    wb.active.append(["Sanjay Iyer", "98200 99999", "sanjay@example.in"])
    buf = io.BytesIO()
    wb.save(buf)
    r = await c.post("/api/crm/import", files={"file": ("old.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.json()["added"] == 1
    assert (await c.post("/api/crm/import", files={"file": ("x.pdf", b"%PDF", "application/pdf")})).status_code == 415
    r = await c.get("/api/crm/export", params={"what": "contacts", "fmt": "csv"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][0] == "Name" and len(rows) == 4
    arjun = next(x for x in rows if x[0] == "Arjun Nair")
    assert arjun[6].startswith("'=") and arjun[5] == "Architect, Referral" and arjun[1] == "+919820088888"
    r = await c.get("/api/crm/export", params={"what": "deals", "fmt": "xlsx"})
    sheet = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in sheet[1]][:6] == ["Name", "Phone", "Email", "Company", "Value (₹)", "Stage"] and sheet.max_row == 2


async def test_older_leads_get_contacts_and_a_won_quote_links_its_invoice(client):
    c = await v23_member("9800100006")
    uid = await owner_id(c)
    # a lead saved the old way (the Growth WhatsApp AI writes leads directly)
    await db().leads.insert_one({"_id": "old-lead-1", "owner_id": uid, "name": "Farah Khan", "phone": "+919820012121", "message": "Sofa repair",
                                 "status": "new", "stage": "logic", "source": "whatsapp", "created_at": now() - timedelta(days=3)})
    deals = (await c.get("/api/crm/deals")).json()["deals"]
    old = next(d for d in deals if d["id"] == "old-lead-1")
    assert old["pstage"] == "logic" and old["contact_id"]
    # a quotation for the deal, accepted, then invoiced: the deal is won and shows both documents
    q = (await c.post("/api/quotes", json={"customer": {"name": "Farah Khan", "phone": "9820012121"}, "lead_id": "old-lead-1",
                                           "items": [{"name": "Sofa repair", "qty": 1, "rate": 12000, "gst": 18}]})).json()
    assert (await c.post(f"/api/quotes/{q['id']}/status", json={"status": "accepted"})).status_code == 200
    # Invoice on yes: accepting it made the invoice
    inv = next(x for x in (await c.get("/api/quotes", params={"kind": "invoice"})).json() if x.get("quote_id") == q["id"])
    d = (await c.get("/api/crm/deals/old-lead-1")).json()
    assert d["status"] == "won" and d["value"] == 14160 and d["invoice"]["number"] == inv["number"] and d["quote"]["number"] == q["number"]
    assert {x["kind"] for x in d["documents"]} == {"quote", "invoice"}
    assert any(a["kind"] == "won" for a in d["activities"])
    # the old Leads page still works and writes the history too
    await paste(c, "Imran Shaikh", "98200 13131")
    lead = await db().leads.find_one({"name": "Imran Shaikh"})
    assert (await c.patch(f"/api/leads/{lead['_id']}", json={"stage": "pain"})).status_code == 200
    d = (await c.get(f"/api/crm/deals/{lead['_id']}")).json()
    assert d["stage"] == "pain" and d["pstage"] == "pain" and any("moved to Pain" in a["text"] for a in d["activities"])


async def test_locked_for_free_owners_and_off_before_the_v23_setup(client):
    free = await v23_member("9800100007", paid=False)
    r = await free.get("/api/crm/contacts")
    assert r.status_code == 403 and r.json()["detail"]["code"] == "locked"
    assert (await free.get("/api/crm/deals")).status_code == 403
    # before the v2.3 setup the CRM is switched off for everyone (the code defaults stay as they were)
    plans.reset_cache()
    await db().config.delete_many({})
    await plans.load(force=True)
    assert not plans.engines_on()
    member = await new_owner("9800100008")
    await db().users.update_one({"_id": await owner_id(member)}, {"$set": {"plan": "program"}})
    r = await member.get("/api/crm/contacts")
    assert r.status_code == 403 and r.json()["detail"]["code"] == "off"
    assert (await member.get("/api/crm/deals")).status_code in (200, 400)   # the pipeline is part of the Action Program as before
    await setup_v23()
