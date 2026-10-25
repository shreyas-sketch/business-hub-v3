"""Builds the OwnerDesk v2.2 feature test pack for one invented business (see data.py) into OUT."""
import csv
import io
import os
import sys
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

sys.path.insert(0, os.path.dirname(__file__))
from data import (BIZ, COMPETITORS, CUSTOMERS, EXPENSES, FAMILIES, FIXED_COSTS, FUNNEL, INQUIRIES, LEAD_RESPONSE, RESEARCH,  # noqa: E402
                  REVIEWS, TODAY, VENDORS, WINBACK, bank_rows, sales_rows)

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
FONT = "/usr/share/fonts/truetype/dejavu/"
pdfmetrics.registerFont(TTFont("Body", FONT + "DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("Bold", FONT + "DejaVuSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("Mono", FONT + "DejaVuSansMono.ttf"))

INK, GOLD, MUTED, LINE, WASH = colors.HexColor("#14151A"), colors.HexColor("#A8872F"), colors.HexColor("#5B5E66"), colors.HexColor("#D9D6CE"), colors.HexColor("#F6F3EC")
S = {
    "h1": ParagraphStyle("h1", fontName="Bold", fontSize=20, leading=25, textColor=INK, spaceAfter=4),
    "h2": ParagraphStyle("h2", fontName="Bold", fontSize=12.5, leading=16, textColor=INK, spaceBefore=12, spaceAfter=5),
    "kick": ParagraphStyle("kick", fontName="Bold", fontSize=8.5, leading=11, textColor=GOLD, spaceAfter=3),
    "p": ParagraphStyle("p", fontName="Body", fontSize=9.6, leading=14, textColor=INK, spaceAfter=4),
    "small": ParagraphStyle("small", fontName="Body", fontSize=8.2, leading=11.5, textColor=MUTED),
    "cell": ParagraphStyle("cell", fontName="Body", fontSize=8.6, leading=11.5, textColor=INK),
    "cellb": ParagraphStyle("cellb", fontName="Bold", fontSize=8.6, leading=11.5, textColor=INK),
    "num": ParagraphStyle("num", fontName="Body", fontSize=8.6, leading=11.5, textColor=INK, alignment=TA_RIGHT),
    "mono": ParagraphStyle("mono", fontName="Mono", fontSize=8.4, leading=12, textColor=INK, backColor=WASH, borderPadding=6, spaceBefore=4, spaceAfter=8),
}
W = A4[0] - 36 * mm


def rs(x: float, paise: bool = False) -> str:
    neg, x = x < 0, abs(x)
    whole, frac = f"{x:.2f}".split(".")
    head, tail = whole[:-3], whole[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    s = ",".join(([head] if head else []) + groups + [tail]) if head or groups else tail
    return ("−" if neg else "") + "₹" + s + (f".{frac}" if paise else "")


def dmy(d) -> str:
    return d.strftime("%d %b %Y")


def P(t, s="p"):
    return Paragraph(t, S[s])


def table(rows, widths, head=True, zebra=True, align_right=()):
    data = [[c if not isinstance(c, str) else P(c, "cellb" if head and i == 0 else ("num" if j in align_right and i else "cell"))
             for j, c in enumerate(r)] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1 if head else 0)
    st = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), .4, LINE),
          ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 5)]
    if head:
        st += [("BACKGROUND", (0, 0), (-1, 0), WASH), ("LINEBELOW", (0, 0), (-1, 0), .9, GOLD)]
    t.setStyle(TableStyle(st))
    return t


def doc(path, code, feature, title, story_fn, issuer=None, tag=None):
    def frame(c, d):
        c.saveState()
        c.setFillColor(INK)
        c.rect(0, A4[1] - 16 * mm, A4[0], 16 * mm, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Bold", 10)
        c.drawString(18 * mm, A4[1] - 10 * mm, issuer or BIZ["name"])
        c.setFont("Body", 8)
        c.setFillColor(colors.HexColor("#C9A84C"))
        c.drawRightString(A4[0] - 18 * mm, A4[1] - 10 * mm, tag or f"OwnerDesk test pack · {code}")
        c.setFillColor(MUTED)
        c.setFont("Body", 7.2)
        c.drawString(18 * mm, 10 * mm, "Invented business for testing. Every name, number, GSTIN, phone and account in this document is made up.")
        c.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {d.page}")
        c.restoreState()
    d = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=24 * mm, bottomMargin=18 * mm,
                          title=f"{code} {feature} — {BIZ['name']}", author="OwnerDesk test pack")
    story = [P(feature.upper(), "kick"), P(title, "h1")]
    story += story_fn()
    d.build(story, onFirstPage=frame, onLaterPages=frame)


def steps_block(how, expect):
    out = [P("How to use it in the hub", "h2")]
    out += [P(f"{i}. {s}") for i, s in enumerate(how, 1)]
    out += [P("What you should see", "h2")]
    out += [P(f"• {s}") for s in expect]
    return out


# ───────────────────────── data files ─────────────────────────
SALES = sales_rows()
BANK = bank_rows()


def write_sales():
    head = ["Date", "Particulars", "Mobile", "GSTIN/UIN", "Voucher No.", "Taxable Value", "CGST", "SGST", "IGST", "Gross Total", "Due Date", "Paid On"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Register"
    first, last = SALES[0][0], SALES[-1][0]
    for line in [BIZ["legal"], BIZ["address"], f"GSTIN: {BIZ['gstin']}", "Sales Register",
                 f"{first.strftime('%d-%b-%y')} to {last.strftime('%d-%b-%y')}", ""]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws["A4"].font = Font(bold=True, size=12)
    ws.append(head)
    for c in ws[7]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F6F3EC")
    for r in SALES:
        ws.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8] or None, r[9], r[10], r[11]])
    n = ws.max_row
    for row in ws.iter_rows(min_row=8, max_row=n):
        for i in (0, 10, 11):
            row[i].number_format = "DD-MM-YYYY"
        for i in (5, 6, 7, 8, 9):
            row[i].number_format = "#,##,##0.00"
    ws.append(["Grand Total", None, None, None, None, sum(r[5] for r in SALES), sum(r[6] for r in SALES), sum(r[7] for r in SALES), None,
               sum(r[9] for r in SALES)])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKL", [12, 34, 13, 18, 18, 14, 12, 12, 8, 14, 12, 12]):
        ws.column_dimensions[col].width = w
    wb.save(f"{OUT}/06_Sales_Register_Nov2025-Oct2026.xlsx")
    with open(f"{OUT}/06_Sales_Register_Nov2025-Oct2026.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(head)
        for r in SALES:
            w.writerow([r[0].strftime("%d-%m-%Y"), r[1], r[2], r[3], r[4], f"{r[5]:.2f}", f"{r[6]:.2f}", f"{r[7]:.2f}", "", f"{r[9]:.2f}",
                        r[10].strftime("%d-%m-%Y"), r[11].strftime("%d-%m-%Y") if r[11] else ""])


def write_customers():
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["Name", "Mobile", "Email", "City", "Type", "Last purchase", "Total value (₹)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    by = {}
    for r in SALES:
        x = by.setdefault(r[1], {"mob": r[2], "last": r[0], "value": 0})
        x["last"], x["value"] = max(x["last"], r[0]), x["value"] + r[9]
    kinds = {c[0]: c[3] for c in CUSTOMERS}
    for name, x in sorted(by.items(), key=lambda kv: -kv[1]["value"]):
        slug = "".join(ch for ch in name.lower() if ch.isalnum())[:18]
        ws.append([name, x["mob"], f"{slug}@example.in", "Pune", kinds.get(name, "family").title(), x["last"], round(x["value"], 2)])
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = "DD-MM-YYYY"
        row[6].number_format = "#,##,##0"
    for col, w in zip("ABCDEFG", [34, 13, 28, 10, 11, 14, 15]):
        ws.column_dimensions[col].width = w
    wb.save(f"{OUT}/11_Win-Back_Customer_List.xlsx")


def write_bank():
    head = ["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]
    top = ["ZZ BANK LTD - STATEMENT OF ACCOUNT (INVENTED DATA)", f"Account Name : {BIZ['legal'].upper()}", f"Account No : {BIZ['account']}",
           f"Branch : {BIZ['bank']}", "Statement From : 01/07/2026 To : 30/09/2026", ""]
    with open(f"{OUT}/05_Bank_Statement_Jul-Sep2026.csv", "w", newline="") as f:
        w = csv.writer(f)
        for line in top:
            w.writerow([line] if line else [])
        w.writerow(head)
        for r in BANK:
            w.writerow([r[0].strftime("%d/%m/%y"), r[1], r[2], r[3].strftime("%d/%m/%y"), f"{r[4]:.2f}" if r[4] else "", f"{r[5]:.2f}" if r[5] else "", f"{r[6]:.2f}"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    for line in top:
        ws.append([line])
    ws.append(head)
    for r in BANK:
        ws.append([r[0], r[1], r[2], r[3], r[4] or None, r[5] or None, r[6]])
    for row in ws.iter_rows(min_row=8):
        row[0].number_format = row[3].number_format = "DD/MM/YYYY"
        for i in (4, 5, 6):
            row[i].number_format = "#,##0.00"
    for col, w in zip("ABCDEFG", [12, 62, 15, 12, 15, 15, 16]):
        ws.column_dimensions[col].width = w
    wb.save(f"{OUT}/05_Bank_Statement_Jul-Sep2026.xlsx")
    # A PDF like net banking's, locked with a password (the reader asks for it once and never keeps it)
    buf = io.BytesIO()
    from reportlab.lib.pagesizes import landscape
    d = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                          title="Statement of account")
    story = [P("ZZ BANK LTD — STATEMENT OF ACCOUNT", "h2"), P(f"Account Name : {BIZ['legal'].upper()}<br/>Account No : {BIZ['account']}<br/>"
             f"Branch : {BIZ['bank']}<br/>Statement From : 01/07/2026 To : 30/09/2026", "small"), Spacer(1, 6)]
    rows = [head] + [[r[0].strftime("%d/%m/%y"), r[1], r[2], r[3].strftime("%d/%m/%y"), f"{r[4]:,.2f}" if r[4] else "",
                      f"{r[5]:,.2f}" if r[5] else "", f"{r[6]:,.2f}"] for r in BANK]
    small = ParagraphStyle("s", parent=S["cell"], fontSize=7.4, leading=9.4)
    smallr = ParagraphStyle("sr", parent=small, alignment=TA_RIGHT)
    data = [[Paragraph(str(c), small if j in (0, 1, 2, 3) else smallr) for j, c in enumerate(r)] for r in rows]
    t = Table(data, colWidths=[18 * mm, 122 * mm, 27 * mm, 18 * mm, 26 * mm, 26 * mm, 28 * mm], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#8C8C8C")), ("BACKGROUND", (0, 0), (-1, 0), WASH),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    story.append(t)
    d.build(story)
    r = PdfReader(io.BytesIO(buf.getvalue()))
    w = PdfWriter()
    for pg in r.pages:
        w.add_page(pg)
    w.encrypt(user_password="DMK4417", owner_password="DMK4417-owner", algorithm="AES-128")
    with open(f"{OUT}/05_Bank_Statement_Jul-Sep2026_password_DMK4417.pdf", "wb") as f:
        w.write(f)


def write_bills():
    os.makedirs(f"{OUT}/04_Expense_Bills", exist_ok=True)
    for i, (fname, ago, vkey, cat, taxable, gst, claim, paid_by, note) in enumerate(EXPENSES, 1):
        d = TODAY - timedelta(days=ago)
        path = f"{OUT}/04_Expense_Bills/{fname}"
        if fname.endswith(".jpg"):
            receipt_photo(path, d, taxable, note)
            continue
        if vkey == "MSEDCL (Factory)":
            electricity_bill(path, d, taxable)
            continue
        vname, vaddr, vgst = VENDORS[vkey]
        items = {
            "plywood": [("BWP Marine Plywood 18 mm, 8x4 ft", "4412", 60, 1950.00), ("Laminate 1 mm, High Gloss White", "4823", 90, 720.00),
                        ("Edge band tape 22 mm (50 m roll)", "3920", 40, 412.50), ("MDF Board 6 mm, 8x4 ft", "4411", 0, 0)],
            "hardware": [("Soft-close hinge, 110°, pair", "8302", 240, 165.00), ("Tandem box drawer system 500 mm", "8302", 36, 520.00),
                         ("Telescopic channel 450 mm, pair", "8302", 60, 236.67), ("Profile handle, 160 mm, matte black", "8302", 120, 0)],
            "granite": [("Quartz countertop, 20 mm, Calacatta (sq ft)", "6810", 210, 380.00), ("Edge polishing (running ft)", "9988", 180, 60.00),
                        ("Sink cut-out with polishing", "9988", 6, 0)],
            "ads": [("Search ads management, October 2026", "998365", 1, 18000.00), ("Instagram ads management, October 2026", "998365", 1, 12000.00)],
        }[vkey]
        items = [x for x in items if x[2] and x[3]]
        sub = sum(q * rate for _, _, q, rate in items)
        lines = [[a, h, q, round(q * rate * taxable / sub, 2)] for a, h, q, rate in items]
        lines[-1][3] = round(taxable - sum(x[3] for x in lines[:-1]), 2)
        tax_invoice(path, d, vname, vaddr, vgst, lines, gst, f"{vkey[:3].upper()}/26-27/{1180 + i * 37}", note)


def tax_invoice(path, d, vname, vaddr, vgst, items, gst, inv_no, note):
    taxable = round(sum(amt for *_, amt in items), 2)
    cgst = sgst = round(gst / 2, 2)
    total = round(taxable + cgst + sgst, 2)
    rows = [["#", "Description of goods / services", "HSN/SAC", "Qty", "Rate", "Amount"]]
    rows += [[str(i), a, h, f"{q:g}", rs(amt / q, True), rs(amt, True)] for i, (a, h, q, amt) in enumerate(items, 1)]
    tot = [["", "", "", "", "Taxable value", rs(taxable, True)], ["", "", "", "", "CGST @ 9%", rs(cgst, True)],
           ["", "", "", "", "SGST @ 9%", rs(sgst, True)], ["", "", "", "", "Invoice total", rs(total, True)]]

    def story():
        s = [table([[f"<b>{vname}</b><br/>{vaddr}<br/>GSTIN: {vgst}", f"<b>Invoice no.</b> {inv_no}<br/><b>Date</b> {dmy(d)}<br/>"
                     f"<b>Place of supply</b> Maharashtra (27)"]], [W * .6, W * .4], head=False), Spacer(1, 6),
             table([[f"<b>Bill to</b><br/>{BIZ['legal']}<br/>{BIZ['factory']}<br/>GSTIN: {BIZ['gstin']}",
                     f"<b>Ship to</b><br/>{BIZ['factory']}<br/><br/>Payment: 15 days, bank transfer"]], [W * .6, W * .4], head=False), Spacer(1, 8)]
        t = table(rows + tot, [8 * mm, 72 * mm, 22 * mm, 14 * mm, 28 * mm, W - 144 * mm], align_right=(3, 4, 5))
        s += [t, Spacer(1, 6), P(f"Amount in words: Rupees {words(total)} only.", "small"),
              P(f"Note: {note}. Goods once sold will not be taken back. Subject to Pune jurisdiction. This is a computer-generated invoice.", "small")]
        return s
    doc(path, "Bill", "Tax invoice", "Tax invoice", story, issuer=vname, tag="Original for recipient")


def electricity_bill(path, d, amount):
    def story():
        units = 1712
        rows = [["Item", "Amount"], ["Fixed charges (LT-V industrial, 25 kW)", rs(3150, True)], [f"Energy charges ({units} units)", rs(13234.80, True)],
                ["Wheeling charges", rs(1198.40, True)], ["Electricity duty", rs(1056.80, True)], ["Bill amount", rs(amount, True)]]
        return [P("Maharashtra State Electricity Distribution Co. Ltd. — electricity bill (sample layout)", "small"), Spacer(1, 4),
                table([[f"<b>Consumer</b> {BIZ['legal']}<br/>{BIZ['factory']}<br/><b>Consumer no.</b> 170044556677",
                        f"<b>Bill month</b> {(d - timedelta(days=20)).strftime('%B %Y')}<br/><b>Bill date</b> {dmy(d - timedelta(days=4))}<br/>"
                        f"<b>Due date</b> {dmy(d + timedelta(days=11))}"]], [W * .6, W * .4], head=False), Spacer(1, 8),
                table(rows, [W * .7, W * .3], align_right=(1,)), Spacer(1, 6),
                P("Electricity carries no GST, so there is nothing to claim back on this bill.", "small")]
    doc(path, "Bill", "Electricity bill", "Electricity bill — factory", story, issuer="Maharashtra State Electricity Distribution Co. Ltd. (sample)",
        tag="LT-V industrial")


def receipt_photo(path, d, amount, note):
    """A phone photo of a handwritten-style receipt book page (what most transport and labour bills look like)."""
    img = Image.new("RGB", (1200, 1600), (236, 230, 214))
    g = ImageDraw.Draw(img)
    big = ImageFont.truetype(FONT + "DejaVuSans-Bold.ttf", 54)
    mid = ImageFont.truetype(FONT + "DejaVuSans.ttf", 40)
    hand = ImageFont.truetype(FONT + "DejaVuSerif-Bold.ttf", 40) if os.path.exists(FONT + "DejaVuSerif-Bold.ttf") else mid
    g.rectangle([60, 60, 1140, 1540], outline=(150, 40, 40), width=6)
    g.text((110, 100), "RAMESH TEMPO SERVICES", font=big, fill=(150, 40, 40))
    g.text((110, 180), "Tempo & goods transport · Bhosari, Pune", font=mid, fill=(150, 40, 40))
    g.text((110, 240), "Mob. 98220 77810", font=mid, fill=(150, 40, 40))
    g.line([110, 320, 1090, 320], fill=(150, 40, 40), width=3)
    g.text((110, 350), "Receipt No. 0472", font=mid, fill=(60, 60, 90))
    g.text((700, 350), f"Date: {d.strftime('%d/%m/%Y')}", font=hand, fill=(30, 40, 110))
    lines = [("Received from:", "Deshmukh Modular"), ("", "Kitchens Pvt. Ltd."), ("For:", "3 trips – Bhosari to"), ("", "Wakad & Hinjewadi sites"),
             ("Rupees:", "Twelve thousand five"), ("", "hundred only")]
    y = 470
    for a, b in lines:
        g.text((110, y), a, font=mid, fill=(60, 60, 90))
        g.text((480, y), b, font=hand, fill=(30, 40, 110))
        g.line([470, y + 56, 1090, y + 56], fill=(170, 160, 140), width=2)
        y += 92
    g.rectangle([110, 1060, 560, 1170], outline=(60, 60, 90), width=3)
    g.text((140, 1085), f"₹ {amount:,.0f}/-", font=big, fill=(30, 40, 110))
    g.text((110, 1220), "Paid in cash", font=hand, fill=(30, 40, 110))
    g.text((760, 1380), "R. Jadhav", font=hand, fill=(30, 40, 110))
    g.text((760, 1450), "Signature", font=mid, fill=(60, 60, 90))
    img = img.rotate(2.2, expand=True, fillcolor=(90, 84, 76)).filter(ImageFilter.GaussianBlur(.6)).resize((900, 1200))
    img.save(path, "JPEG", quality=82)


ONES = "zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen".split()
TENS = "_ _ twenty thirty forty fifty sixty seventy eighty ninety".split()


def words(x: float) -> str:
    n = int(round(x))

    def two(k):
        return ONES[k] if k < 20 else TENS[k // 10] + ("-" + ONES[k % 10] if k % 10 else "")

    def three(k):
        return (ONES[k // 100] + " hundred" + (" " + two(k % 100) if k % 100 else "")) if k >= 100 else two(k)
    parts = []
    for size, name in ((10 ** 7, "crore"), (10 ** 5, "lakh"), (1000, "thousand")):
        if n >= size:
            parts.append(f"{three(n // size)} {name}")
            n %= size
    if n or not parts:
        parts.append(three(n))
    return " ".join(parts).capitalize()


# ───────────────────────── the documents ─────────────────────────
def profile_doc():
    def story():
        b = BIZ
        offers = [["What you sell", "Price", "Unit"]] + [[o["name"], o["price"], o["unit"]] for o in b["offers"]]
        costs = [["Fixed cost", "Amount a month", "Goes out on", "Kind"]] + [[n, rs(a), f"{d}th" if d > 3 else {1: "1st", 2: "2nd", 3: "3rd"}[d], k.title()]
                                                                             for n, a, d, k in FIXED_COSTS]
        return [
            P("Everything the AI is allowed to use comes from here. Type these into <b>Knowledge Hub → Business Brain</b> (and the Identity, "
              "Money and Targets sections) before testing anything else. Nothing outside this page is ever invented."),
            P("Basics", "h2"),
            table([["Field", "Type this"], ["Business name", b["name"]], ["City", b["city"]], ["What you do", b["industry"]],
                   ["Ideal customer", b["ideal_customer"]], ["Why customers choose you", b["why_us"].replace("\n", "<br/>")],
                   ["Proof", b["proof"]], ["Language", b["language"]]], [45 * mm, W - 45 * mm]),
            P("What you sell", "h2"), table(offers, [W * .55, W * .25, W * .2]),
            P("Contact", "h2"),
            table([["Field", "Type this"], ["WhatsApp number", f"{b['whatsapp']} <font color='#5B5E66'>(use your own test number)</font>"],
                   ["Email", b["email"]], ["How customers pay", b["payment_note"]], ["Showroom", b["address"]], ["Factory", b["factory"]]],
                  [45 * mm, W - 45 * mm]),
            P("Identity", "h2"),
            table([["Field", "Type this"], ["Legal name", b["legal"]], ["Type of business", "Private limited company"],
                   ["GSTIN", f"<font name='Mono'>{b['gstin']}</font>"], ["State", "Maharashtra"], ["GST returns", "Every month (GSTR-1 and GSTR-3B)"],
                   ["People who work with you", str(b["staff"])]], [45 * mm, W - 45 * mm]),
            P("Money basics", "h2"),
            table([["Field", "Type this"], ["Usual payment terms", f"{b['terms_days']} days (builders and architects; families pay 50% advance)"],
                   ["GST rate", f"{b['gst_rate']}%"], ["Bank balance", f"{rs(1842300)} on {dmy(TODAY - timedelta(days=1))}"],
                   ["Safety buffer", rs(200000)]], [45 * mm, W - 45 * mm]),
            Spacer(1, 4), table(costs, [W * .45, W * .2, W * .15, W * .2], align_right=(1,)),
            P("Targets and brand", "h2"),
            table([["Field", "Type this"], ["Monthly revenue target", rs(b["monthly_target"])], ["Brand message (pick or type)", b["brand_message"]]],
                  [45 * mm, W - 45 * mm]),
        ]
    doc(f"{OUT}/01_Business_Profile_Knowledge_Hub.pdf", "Setup", "Knowledge Hub", "Business profile — type this in first", story)


def funnel_doc():
    f = FUNNEL

    def story():
        pct = lambda a, b: f"{a / b * 100:.1f}%"  # noqa: E731
        rows = [["Step", "September 2026", "Conversion", "Target"],
                ["Website visitors", f"{f['visitors']:,}", "", ""], ["Inquiries (website, WhatsApp, calls)", str(f["inquiries"]), pct(f["inquiries"], f["visitors"]), "3%"],
                ["Site visits / measurements", str(f["meetings"]), pct(f["meetings"], f["inquiries"]), "50%"],
                ["Quotations sent", str(f["proposals"]), pct(f["proposals"], f["meetings"]), "70%"],
                ["Orders won", str(f["won"]), pct(f["won"], f["proposals"]), "30%"]]
        return [
            P("The monthly sales review the owner fills in from the CRM register and Google Analytics. Use it on <b>Funnel leak finder</b>."),
            P("Last month's numbers", "h2"), table(rows, [W * .4, W * .2, W * .2, W * .2], align_right=(1, 2, 3)),
            Spacer(1, 4), P(f"Average order value: <b>{rs(f['avg_sale'])}</b> (kitchen + wardrobes, before GST).", "p"),
            P("Two ways to test it", "h2"),
            P("<b>A. With visitors</b> — type all five numbers and the average sale."),
            P("<b>B. Like most owners</b> — leave website visitors at 0. The funnel then starts at inquiries."),
            *steps_block(
                ["Open <b>Funnel leak finder</b> (press / and type “leak”).",
                 "Type the numbers above (A or B) and the average sale, then press <b>Find my biggest leak</b>.",
                 "Press <b>Get fixes</b> (1 AI run).", "Reload the page: the numbers are kept."],
                [f"<b>Biggest leak: Inquiries → Site visits</b> — {pct(f['meetings'], f['inquiries'])} against a 50% target.",
                 "Quotations (77%) and orders won (35%) are on target.",
                 "“If you fix this one step”: about 2.7 more orders a month, worth about ₹5,72,108.",
                 "Get fixes: three to five actions for booking more site visits, written for a kitchen business in Pune."]),
        ]
    doc(f"{OUT}/02_Funnel_Leak_Finder_Sep2026.pdf", "CV-03", "Funnel leak finder", "September sales review", story)


def inquiries_doc():
    def story():
        rows = [["#", "Channel", "Customer", "Message to type", "Tests"]]
        rows += [[str(i), ch, n, m, t] for i, (ch, n, m, t) in enumerate(INQUIRIES, 1)]
        return [
            P("Six real-sounding inquiries to send to the test business. Send each from <b>a second phone</b> using your own number, so the "
              "WhatsApp messages reach you. Use them for <b>Follow-ups</b> (CV-04) and <b>Lead Response</b> (AG-05)."),
            table(rows, [7 * mm, 26 * mm, 30 * mm, 70 * mm, W - 133 * mm]),
            *steps_block(
                ["In <b>Inquiries</b>, check <b>Follow up on every inquiry</b> is on (event accounts start with it on).",
                 "Open the business's website on another phone and send inquiry 1 with the form.",
                 "Paste inquiry 3 into Inquiries → <b>Add an inquiry</b> (from WhatsApp).",
                 "The next working day (10 am–7 pm), open Inquiries.", "Tap <b>Send on WhatsApp</b> on one; tap <b>Skip this one</b> on another.",
                 "Mark inquiry 3 <b>won</b>."],
                ["Each inquiry shows up in Inquiries at once, with an alert to the owner.",
                 "Day 1: “Follow-up 1 of 3 is ready” on each open inquiry (1 AI run each), and in the Send list.",
                 "Send on WhatsApp opens your WhatsApp with the follow-up typed — nothing goes without your tap.",
                 "Inquiry 4 is told ₹1,450 per sq ft (it's in the Knowledge Hub); no other price is invented.",
                 "After inquiry 3 is marked won, no more follow-ups for it."]),
        ]
    doc(f"{OUT}/03_Follow-ups_Test_Inquiries.pdf", "CV-04", "Follow-ups on day 1, 3 and 7", "Six inquiries to send", story)


def expenses_doc():
    def story():
        rows = [["File", "Date", "Vendor", "Category", "Amount", "GST", "Claim", "Paid by"]]
        for fname, ago, vkey, cat, taxable, gst, claim, paid, note in EXPENSES:
            vname = VENDORS[vkey][0] if vkey in VENDORS else vkey
            rows.append([fname.replace("_", " ").replace(".pdf", "").replace(".jpg", " (photo)")[8:], dmy(TODAY - timedelta(days=ago)), vname, cat,
                         rs(taxable + gst), rs(gst) if gst else "—", "Yes" if claim else "No", paid])
        total = sum(t + g for *_, t, g, c, p, n in [(e[0], e[1], e[2], e[3], e[4], e[5], e[6], e[7], e[8]) for e in EXPENSES])
        claim = sum(e[5] for e in EXPENSES if e[6])
        return [
            P("This month's bills from the accounts file, in the folder <b>04_Expense_Bills</b>. Add each one in <b>Expenses → Add an expense</b> "
              "and attach its file. Amounts include GST."),
            table(rows, [27 * mm, 17 * mm, 31 * mm, 23 * mm, 22 * mm, 19 * mm, 13 * mm, W - 152 * mm], align_right=(4, 5)),
            Spacer(1, 4), P(f"Total: <b>{rs(total)}</b> · GST you can claim back: <b>{rs(claim)}</b>"),
            P("Settings to make first (Expenses → Settings)", "h2"),
            P("The hub picks the categories from the Knowledge Hub's industry — for a factory these start with <b>Raw material</b>, Job work and "
              "Consumables. Add one of your own: <b>Hardware & fittings</b>. Then set these monthly budgets:"),
            table([["Head", "Monthly budget", "This month's bills"], ["Raw material", rs(350000), rs(331462)], ["Hardware & fittings (your own)", rs(60000), rs(74104)],
                   ["Marketing & ads", rs(40000), rs(35400)], ["Freight & courier", rs(15000), rs(12500)]], [W * .45, W * .25, W * .3], align_right=(1, 2)),
            *steps_block(
                ["In Settings, add the category <b>Hardware & fittings</b> and set the budgets above.",
                 "Add each bill: date, amount, vendor, category, how it was paid, the GST on the bill and whether you can claim it. Attach the file.",
                 "Open <b>Where it went</b>.", "Press <b>For your CA → Download for your CA (CSV)</b>.",
                 "With a staff login: add the transport receipt as staff; log in as the owner and approve it."],
                [f"Spent this month includes the {len(EXPENSES)} bills; GST you can claim back is {rs(claim)}.",
                 "Spent: ₹4,72,106. Where it went: Hardware & fittings is over its ₹60,000 budget (₹74,104); the other heads are under.",
                 "The CSV has one row per bill: vendor, GSTIN, category, amount, GST, claimable, bill attached.",
                 "The electricity bill and the cash receipt show no GST to claim."]),
        ]
    doc(f"{OUT}/04_Expense_Tracker_Bills_to_Enter.pdf", "DT-04", "Expense tracker", "This month's bills", story)


def bank_doc():
    def story():
        spent = sum(r[4] for r in BANK if "ATM" not in r[1] and "OWN AC" not in r[1])
        money_in = sum(r[5] for r in BANK)
        return [
            P("The current-account statement for July–September 2026, downloaded from net banking in three formats. Upload <b>one</b> of them on "
              "<b>Bank statement reader</b>."),
            table([["File", "Format", "Password"], ["05_Bank_Statement_Jul-Sep2026.csv", "CSV (Excel can open it)", "—"],
                   ["05_Bank_Statement_Jul-Sep2026.xlsx", "Excel", "—"],
                   ["05_Bank_Statement_Jul-Sep2026_password_DMK4417.pdf", "PDF, locked like most bank PDFs", "<font name='Mono'>DMK4417</font>"]],
                  [W * .55, W * .27, W * .18]),
            P("What is in it", "h2"),
            table([["", "July–September 2026"], ["Money in (customer payments, interest)", rs(money_in)],
                   ["Spent (without cash withdrawn and transfers to own account)", rs(spent)], ["Transactions", str(len(BANK))]],
                  [W * .65, W * .35], align_right=(1,)),
            P("Payments that repeat every month: rent (5th, two of them), the tempo EMI (10th), electricity (18th), software (3rd), broadband (19th), "
              "salaries (1st) and the GST challan (20th).", "small"),
            *steps_block(
                ["Open <b>Bank statement reader</b> and upload one file (for the PDF, type the password once — it isn't kept).",
                 "Check <b>Where it went</b>, then <b>Repeating</b>.",
                 "Change the category of the Swiggy payment to “Staff food & welfare” and tick <b>use it for this payee every time</b>.",
                 "Press <b>Add all to Expenses</b>, then open Expenses.", "Upload the same file again.", "Press <b>Delete it now</b>."],
                ["Account ending 4417, 1 Jul 2026 to 30 Sep 2026, read in seconds.",
                 "Cash withdrawn (₹20,000 a month) and the ₹1,00,000 transfers to the owner's own savings aren't counted as spending.",
                 "Salaries go into Expenses only if you tick them; phone numbers, UPI IDs and account numbers are masked.",
                 "The second upload adds nothing twice.", "Delete it now removes the list (it would go on its own after 24 hours)."]),
        ]
    doc(f"{OUT}/05_Bank_Statement_Reader_Guide.pdf", "DT-05", "Bank statement reader", "Three months of the current account", story)


def sales_summary():
    from collections import defaultdict
    since = TODAY - timedelta(days=365)
    tot, by = 0.0, defaultdict(float)
    for r in SALES:
        if r[0] >= since:
            tot += r[9]
            by[r[1]] += r[9]
    top = sorted(by.items(), key=lambda kv: -kv[1])
    return tot, top


def sales_docs():
    tot, top = sales_summary()
    share = top[0][1] / tot * 100

    def gst_story():
        last_m = (TODAY.replace(day=1) - timedelta(days=1)).replace(day=1)
        month_gst = lambda y, m: sum(r[6] + r[7] for r in SALES if r[0].year == y and r[0].month == m)  # noqa: E731
        sept, octo = month_gst(last_m.year, last_m.month), month_gst(TODAY.year, TODAY.month)
        bills = sum(e[5] for e in EXPENSES if e[6])
        return [
            P("GST set-aside works from what is already in the hub: the sales register (06) for GST on sales, and the approved, claimable expenses (04) "
              "for GST on bills. No new document is needed."),
            table([["Source", "Upload / enter", "Used for"], ["06_Sales_Register_Nov2025-Oct2026.xlsx", "Customer risk → Upload sales", "GST on sales (CGST + SGST columns)"],
                   ["04_Expense_Bills (4 GST bills)", "Expenses → Add an expense", "GST on bills (claimable only)"]], [W * .38, W * .3, W * .32]),
            P("What the numbers should come to", "h2"),
            table([["", last_m.strftime("%B %Y") + " (last month)", TODAY.strftime("%B %Y") + " (so far)"],
                   ["GST on your sales", rs(sept), rs(octo)], ["− GST on your bills (claimable)", rs(0), rs(bills)],
                   ["= about this much to keep aside", rs(sept), rs(octo - bills)],
                   ["GSTR-3B due", "20 " + TODAY.strftime("%b %Y") + " — already passed (Overdue)", "20 " + (TODAY.replace(day=28) + timedelta(days=5)).strftime("%b %Y")]],
                  [W * .34, W * .33, W * .33], align_right=(1, 2)),
            P("This month's bills are all October bills, so last month has none.", "small"),
            *steps_block(
                ["Upload the sales register (06) and add the bills (04) first.", "Open <b>GST set-aside</b>.", "Open <b>How this is worked out</b>.",
                 "On this month's card press <b>Count it in my cash flow</b>, then open Cash flow."],
                ["Both cards marked Estimate, with the amounts above.",
                 "Last month's due date has passed, so it says to mark it done in the compliance calendar if it was paid, and can't go into cash flow.",
                 "This month: “2 bills this month have no GST entered” — the electricity bill and the transport receipt, which really have no GST.",
                 "Cash flow shows the GST payment on 20 November."]),
        ]
    doc(f"{OUT}/07_GST_Set-aside_Guide.pdf", "DT-06", "GST set-aside", "Where the GST numbers come from", gst_story)

    def risk_story():
        rows = [["Customer", "12-month sales", "Share"]] + [[n, rs(v), f"{v / tot * 100:.1f}%"] for n, v in top[:6]]
        return [
            P("The Sales Register exported from Tally (Gateway of Tally → Display → Account Books → Sales Register → Excel), "
              f"{len(SALES)} invoices from November 2025 to October 2026. Upload <b>06_Sales_Register_Nov2025-Oct2026.xlsx</b> "
              "(or the .csv) on <b>Customer risk</b>. Upload it once — Getting paid, GST set-aside and Win-Back use the same sales."),
            table(rows, [W * .55, W * .27, W * .18], align_right=(1, 2)),
            Spacer(1, 4), P(f"Total sales, last 12 months: <b>{rs(tot)}</b> from {len(top)} customers.", "small"),
            *steps_block(
                ["Open <b>Customer risk</b> and upload the sales register.", "Read the top line.", "Open <b>Gone quiet</b>.",
                 "Tap <b>Check in on WhatsApp</b> on Hotel Sahyadri Residency (the number is invented — don't send).", "Press <b>Undo the upload</b> to try it again."],
                [f"Top line: <b>“1 customer is about {share:.0f}% of your sales”</b> — {top[0][0]}.",
                 "The Tally header rows (company name, period) and the Grand Total row are skipped; every invoice is read.",
                 "Gone quiet: Hotel Sahyadri Residency, Gokhale Associates and The Spice Route Cloud Kitchen — regulars well past their usual gap.",
                 "Check in on WhatsApp opens your WhatsApp with a short hello typed."]),
        ]
    doc(f"{OUT}/06_Customer_Risk_Sales_Register_Guide.pdf", "DT-07", "Customer risk", "Twelve months of sales from Tally", risk_story)

    def paid_story():
        late = [r for r in SALES if r[3] and r[11] is None and (TODAY - r[10]).days > 0]
        rows = [["Customer", "Invoice", "Due", "Days late", "Amount"]] + [[r[1], r[4], dmy(r[10]), str((TODAY - r[10]).days), rs(r[9])]
                                                                         for r in sorted(late, key=lambda r: r[10])[:8]]
        return [
            P("Getting paid reads the same sales register (06): its <b>Due Date</b> and <b>Paid On</b> columns show how long each customer really took. "
              f"Your terms ({BIZ['terms_days']} days) come from the Knowledge Hub (01)."),
            P("Unpaid and past due in the register", "h2"),
            table(rows, [W * .34, W * .22, W * .16, W * .1, W * .18], align_right=(3, 4)),
            *steps_block(
                ["Set <b>Usual payment terms</b> to 15 days in Knowledge Hub → Money basics.", "Upload the sales register (if not done already).",
                 "Open <b>Getting paid</b>.", "Press <b>Chase these in Money owed</b>."],
                ["Top line: “You give 15 days. On average you wait 27.” (bigger amounts count more).",
                 "Payers: Nirmiti Homes, Shubham Builders and Vastu Realty are often late; The Spice Route Cloud Kitchen pays on time.",
                 "The late buckets (1–30, 31–60, 61–90, over 90 days) add up to the unpaid invoices above.",
                 "Chase these opens Money owed with the reminders."]),
        ]
    doc(f"{OUT}/08_Getting_Paid_Guide.pdf", "DT-08", "Getting paid", "What you ask for, and what really happens", paid_story)


def competitors_doc():
    def story():
        rows = [["Name", "Website", "Facebook page name", "Why we watch them"]] + [[n, w, f, why] for n, w, f, why in COMPETITORS]
        return [
            P("The three brands the owner loses most quotes to. These are <b>real companies' public websites</b>, used only as competitors to watch: "
              "Competitor Watch reads at most 5 public pages each, follows robots.txt, and never reads social networks. "
              "For your own testing, swap in two or three real local competitors you know."),
            table(rows, [24 * mm, 42 * mm, 30 * mm, W - 96 * mm]),
            *steps_block(
                ["Open <b>Competitor Watch</b> → <b>Add</b>, and add each competitor with its website and Facebook page name.",
                 "Press <b>Check now</b> (it can take up to a minute).", "Press <b>See their running ads</b> on one.",
                 "Press Check now again the same day."],
                ["The first look is saved for each website: pages read, prices and offers seen.",
                 "See their running ads opens Meta's public Ad Library for that page.",
                 "The second Check now says it was already checked today (once a day per competitor).",
                 "From the next Monday 8 am: a short summary of only what changed (1 AI run, only when something changed) and a WhatsApp alert."]),
        ]
    doc(f"{OUT}/09_Competitor_Watch_List.pdf", "AG-01", "Competitor Watch", "Competitors to watch", story)


def reviews_doc():
    def story():
        out = [P("Six reviews from Google, Justdial, Facebook and IndiaMART. Paste each in <b>Reviews → Paste a review</b> with its source, stars, name "
                 "and date. Until the hub's Google key is set, Google reviews are pasted too.")]
        for src, stars, who, ago, text, shows in REVIEWS:
            out.append(KeepTogether([table([[f"<b>{'★' * stars}{'☆' * (5 - stars)}</b>  {who} · {src} · {dmy(TODAY - timedelta(days=ago))}"],
                                            [text], [f"<font color='#5B5E66'>Shows: {shows}</font>"]], [W], head=False), Spacer(1, 4)]))
        out += steps_block(["Paste the 2-star Google review.", "Read the draft (1 AI run); edit it or press Draft again, then <b>Approve</b>.",
                            "Press <b>Copy</b>, then <b>I've posted it</b>.", "Paste a 5-star review and press <b>Ask for a referral</b>."],
                           ["A calm, specific reply in the owner's voice: it owns the delay, offers a call, and never promises a refund, discount or date.",
                            "No price appears in any draft.", "Posted reviews move to the Posted tab.",
                            "Ask for a referral opens WhatsApp with a short message typed to the reviewer (if you have their number)."])
        return out
    doc(f"{OUT}/10_Reviews_to_Paste.pdf", "AG-02", "Reviews & reputation", "Reviews to reply to", story)


def winback_doc():
    def story():
        quiet = [c for c in CUSTOMERS if c[6] >= 90]
        rows = [["Customer", "Purchases (12 months)", "Last purchase"]] + [[c[0], str(c[4]), dmy(TODAY - timedelta(days=c[6]))] for c in quiet]
        return [
            P("Win-Back works from the sales register (06). Optionally import <b>11_Win-Back_Customer_List.xlsx</b> in <b>Customers → Import</b> to add "
              "email addresses, so messages can also go by email."),
            P("Settings to choose", "h2"),
            table([["Setting", "Choose"], ["Quiet for", f"{WINBACK['days']} days"], ["At least", f"{WINBACK['min_purchases']} purchases"],
                   ["Messages a week", str(WINBACK["per_week"])], ["Your offer (word for word)", WINBACK["offer"]]], [W * .35, W * .65]),
            P("Customers who should come up", "h2"), table(rows, [W * .5, W * .25, W * .25]),
            *steps_block(
                ["Upload the sales register on Customer risk first.", "Open <b>Win-Back</b>, choose the settings above and read the sample message.",
                 "Tick <b>I've read this message…</b> and press <b>Switch Win-Back on</b>.", "Press <b>Run now</b> (1 AI run).",
                 "Tap <b>Send on WhatsApp</b> on one (the numbers are invented — don't send).", "Switch it off when you're done."],
                ["This week's messages: the quiet regulars above, biggest spenders first.",
                 "Each message has the customer's first name and your offer exactly as typed — never a price or discount you didn't type.",
                 "Nobody who bought in the last 30 days is messaged.", "Unsent messages expire after 14 days."]),
        ]
    doc(f"{OUT}/11_Win-Back_Setup.pdf", "AG-03", "Win-Back", "Bring back customers who went quiet", story)


def research_doc():
    def story():
        rows = [["Who buys from you", "City", "What you'd sell them"]] + [list(r) for r in RESEARCH]
        return [
            P("Who the owner wants more of. Type one row at a time in <b>Research → Find companies that fit</b>."),
            table(rows, [W * .5, W * .2, W * .3]),
            *steps_block(["Press <b>Write 5 post ideas</b> (1 AI run).", "Type row 1 and press <b>Make my research plan</b> (1 AI run).",
                          "Open one search link.", "Download the <b>CSV template</b>."],
                         ["Five post ideas, each with a hook, what to say and the format. With a Competitor Watch report that has changes, each answers one; "
                          "without one, they come from the business profile and the page says so.",
                          "A plan: kinds of companies, filters, questions to qualify them, and search links built by the hub (not the AI).",
                          "Nothing is scraped and no personal details are collected; you open the searches yourself."]),
        ]
    doc(f"{OUT}/12_Content_and_Company_Research_Brief.pdf", "AG-04", "Content & company research", "Who to find", story)


def lead_doc():
    def story():
        L = LEAD_RESPONSE
        return [
            P("Lead Response sends the first reply from the owner's <b>own WhatsApp number</b>, so it needs that number connected and one approved template."),
            P("1. The template (create it in the owner's AiSensy account)", "h2"),
            table([["Field", "Value"], ["Name", f"<font name='Mono'>{L['template']}</font>"], ["Category", "Utility"], ["Language", "English"],
                   ["Body", f"<font name='Mono'>{L['body']}</font>"],
                   ["Samples", "{{1}} Neha · {{2}} Deshmukh Modular Kitchens · {{3}} Would tomorrow at 11 am or Saturday at 4 pm suit you for a site visit? "
                               "Just reply with the time."]], [30 * mm, W - 30 * mm]),
            P("2. Connections → WhatsApp — your own number", "h2"),
            table([["Field", "Value"], ["Provider", "AiSensy"], ["Number", "The owner's WhatsApp Business number"], ["API key", "From AiSensy → Manage → API key"],
                   ["Template: first reply with a booking", L["template"]], ["Webhook", "Copy the address shown in Connections into AiSensy"]],
                  [55 * mm, W - 55 * mm]),
            P("3. Lead Response settings", "h2"),
            table([["Setting", "Choose"], ["How to book", L["booking"]], ["What to book", L["what"]], ["Hours", "10 am to 7 pm (never Sunday)"],
                   ["Prices and quotes go to a person", "On"]], [55 * mm, W - 55 * mm]),
            *steps_block(["Press <b>Send it to my WhatsApp</b> to test the first reply.", "Press <b>Switch Lead Response on</b>.",
                          "From another phone, send inquiry 5 (from document 03) with the website form.", "Reply to the WhatsApp with one of the times."],
                         ["Within seconds: “Hello Kunal, thank you for contacting Deshmukh Modular Kitchens. Would … or … suit you for a site visit?”",
                          "The two times are inside 10 am–7 pm and never on a Sunday.", "The reply lands in WhatsApp chats for a person to confirm.",
                          "The page shows the last 30 days: replies sent and how fast.",
                          "If the number or template is missing, the page says which, and nothing is sent."]),
        ]
    doc(f"{OUT}/13_Lead_Response_Setup.pdf", "AG-05", "Lead Response", "First reply from your own number", story)


def start_doc():
    def story():
        rows = [["#", "Feature", "Documents", "Plan"],
                ["—", "Knowledge Hub (do this first)", "01_Business_Profile_Knowledge_Hub.pdf", "All"],
                ["CV-03", "Funnel leak finder", "02_Funnel_Leak_Finder_Sep2026.pdf", "Membership"],
                ["CV-04", "Follow-ups on day 1, 3 and 7", "03_Follow-ups_Test_Inquiries.pdf", "Membership"],
                ["DT-04", "Expense tracker", "04_Expense_Tracker_Bills_to_Enter.pdf + folder 04_Expense_Bills (5 PDFs, 1 photo)", "Membership"],
                ["DT-05", "Bank statement reader", "05_Bank_Statement_Reader_Guide.pdf + statement (.csv, .xlsx, locked .pdf)", "Membership"],
                ["DT-07", "Customer risk", "06_Customer_Risk_Sales_Register_Guide.pdf + 06_Sales_Register (.xlsx, .csv)", "Membership"],
                ["DT-06", "GST set-aside", "07_GST_Set-aside_Guide.pdf (uses 04 and 06)", "Membership"],
                ["DT-08", "Getting paid", "08_Getting_Paid_Guide.pdf (uses 06)", "Membership"],
                ["AG-01", "Competitor Watch", "09_Competitor_Watch_List.pdf", "Membership"],
                ["AG-02", "Reviews & reputation", "10_Reviews_to_Paste.pdf", "Membership"],
                ["AG-03", "Win-Back", "11_Win-Back_Setup.pdf + 11_Win-Back_Customer_List.xlsx (optional)", "Action Program"],
                ["AG-04", "Content & company research", "12_Content_and_Company_Research_Brief.pdf", "Action Program"],
                ["AG-05", "Lead Response", "13_Lead_Response_Setup.pdf (uses 03)", "Growth Mentorship"]]
        return [
            P(f"<b>{BIZ['name']}</b> makes modular kitchens, wardrobes and full home interiors in its own factory in Bhosari, Pune — 18 people, "
              f"about {rs(sales_summary()[0])} of sales a year, builders and architects as well as families. Everything in this pack is that one "
              "business's real-looking paperwork, so the twelve features agree with each other: the same sales feed Customer risk, Getting paid, "
              "GST set-aside and Win-Back; the same bills feed Expenses and GST."),
            P("Order to test in", "h2"),
            P("1. Type the business profile (01) into the Knowledge Hub.  2. Upload the sales register (06) once.  3. Add the bills (04).  "
              "4. Then any feature, in any order. The AI features need <b>Admin → Pulse → Test the AI</b> to say Working."),
            P("What's in the pack", "h2"), table(rows, [14 * mm, 42 * mm, 92 * mm, W - 148 * mm]),
            P("Rules this pack respects", "h2"),
            P("• Every person, phone number, GSTIN, account and invoice is invented. Don't send the WhatsApp messages to these numbers; use your own second "
              "phone for anything that should arrive.<br/>• The competitor websites in 09 are real brands' public sites, read the way the feature reads "
              "any competitor.<br/>• Dates are set around October 2026, so “last month” is September 2026."),
        ]
    doc(f"{OUT}/00_START_HERE.pdf", "Index", "OwnerDesk v2.2 — feature test pack", "Twelve features, one real-looking business", story)


write_sales()
write_customers()
write_bank()
write_bills()
for f in (start_doc, profile_doc, funnel_doc, inquiries_doc, expenses_doc, bank_doc, sales_docs, competitors_doc, reviews_doc, winback_doc,
          research_doc, lead_doc):
    f()
print("\n".join(sorted(os.listdir(OUT))))
